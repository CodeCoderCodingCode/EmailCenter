"""Author the reusable editable application documents with the bundled runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'assets/templates'
OUTPUT.mkdir(parents=True, exist_ok=True)


def document():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.8), Inches(.8)
    section.left_margin, section.right_margin = Inches(.9), Inches(.9)
    for key in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
        style = doc.styles[key]
        style.font.name = 'Calibri'
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(9)
    normal.paragraph_format.line_spacing = 1.12
    doc.styles['Title'].font.size = Pt(23)
    doc.styles['Heading 1'].font.size = Pt(13)
    from docx.oxml.ns import qn
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.core_properties.author = 'EmailCenter'
    doc.core_properties.subject = 'Bearbeitbare Vorlage · persönliche Angaben selbst ergänzen'
    return doc


for key, title in (('Ausbildung', 'Bewerbung um einen Ausbildungsplatz als {{job.title}}'), ('Beruf', 'Bewerbung als {{job.title}}')):
    doc = document()
    doc.add_paragraph('{{applicant.display_name}}\n{{applicant.address}}\n{{applicant.email}} · {{applicant.phone}}')
    doc.add_paragraph('{{company.name}}\n{{contact.display_name}}\n{{company.address}}')
    date = doc.add_paragraph('{{application.date}}')
    date.alignment = 2
    doc.add_paragraph(title, style='Heading 1')
    doc.add_paragraph('{{contact.salutation}} {{contact.display_name}},')
    doc.add_paragraph('hiermit bewerbe ich mich bei {{company.name}} als {{job.title}}.')
    doc.add_paragraph('{{application.cover_letter}}')
    doc.add_paragraph('Gerne möchte ich Sie in einem persönlichen Gespräch von meiner Motivation überzeugen. Über eine Einladung freue ich mich.')
    doc.add_paragraph('Freundliche Grüße\n\n{{applicant.display_name}}')
    doc.add_paragraph('Anlagen: Lebenslauf und Nachweise')
    doc.save(OUTPUT / ('Anschreiben_' + key + '.docx'))

doc = document()
doc.add_paragraph('Lebenslauf', style='Title')
doc.add_paragraph('{{applicant.display_name}}', style='Heading 1')
doc.add_paragraph('{{applicant.address}}\n{{applicant.email}} · {{applicant.phone}}')
for title, field in (('Berufliche Erfahrung / Praktika', 'cv.experience'), ('Schulbildung / Ausbildung', 'cv.education'), ('Kenntnisse und Fähigkeiten', 'cv.skills'), ('Sprachen', 'cv.languages'), ('Interessen / Engagement', 'cv.interests')):
    doc.add_paragraph(title, style='Heading 1')
    doc.add_paragraph('{{' + field + '}}')
doc.add_paragraph('{{application.date}} · {{applicant.display_name}}')
doc.save(OUTPUT / 'Lebenslauf.docx')
print('Created 3 editable DOCX templates.')
