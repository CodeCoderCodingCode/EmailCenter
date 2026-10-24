"""One real local Excel/PDF/project/export/restore workflow; fake mail only."""
from copy import deepcopy
import argparse
import csv
from datetime import datetime, timedelta, timezone
from email import message_from_bytes, policy
import hashlib
import json
from pathlib import Path
import sys
import time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / 'src'), str(ROOT / 'tests')]
from openpyxl import Workbook, load_workbook
import pdfplumber
from emailcenter.project import create_project, open_project
from emailcenter.import_service import ImportService
from emailcenter.bulk_sources import BulkSourceService
from emailcenter.data_source_library import DataSourceLibrary
from emailcenter.word_replace import WordReplaceService
from emailcenter.document_service import DocumentService, find_libreoffice
from emailcenter.mail_templates import MailTemplateService
from emailcenter.mail_bulk import BulkMailService
from emailcenter.mail_service import MailService
from emailcenter.bulk_delivery import BulkDeliveryService
from emailcenter.credentials import WindowsTokenStore
from emailcenter.gmail import GmailProvider
from emailcenter.email_builder import EmailBuilderService, starter_graph, node
from emailcenter.builder_workspace import BuilderWorkspaceService
from emailcenter.export_service import ExportService
from emailcenter.administration import AdministrationService
from emailcenter.backups import inspect_backup, restore_backup
from emailcenter.diagnostics import diagnostic_report
from test_mail import MemoryVault, FakeProvider
from run_completion_checks import code_hashes


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def pdf_values(path, fields):
    with pdfplumber.open(path) as pdf:
        text = '\n'.join(page.extract_text() or '' for page in pdf.pages)
        assert len(pdf.pages) == 2
        assert all(fields[key] in text for key in ('contact.display_name', 'contact.email', 'company.name'))
        assert '{{NAME}}' not in text and '{{COMPANY}}' not in text
        return {'pages': len(pdf.pages), 'personal_values': True, 'sha256': sha(path)}


def main(target):
    target = target.resolve()
    if not target.is_relative_to(ROOT):
        raise ValueError('Workflow output must stay within the workspace.')
    target.mkdir(parents=True, exist_ok=False)
    report = {'status': 'Running', 'checks': {}, 'real_gmail_requests': 0, 'native_ui_acceptance': False}
    report_path = target / 'Workflow_Result.json'
    before, started = code_hashes(), time.monotonic()
    project = None
    def save():
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    def checked(name, value=True):
        report['checks'][name] = value
        save()
        print(name + ': OK', flush=True)
    def no_network(*args, **kwargs):
        report['real_gmail_requests'] += 1
        raise RuntimeError('Live Gmail is disabled in this synthetic workflow.')
    vault, provider = MemoryVault(), FakeProvider()
    clock = [datetime(2026, 10, 13, 12, tzinfo=timezone.utc)]
    mail_options = {'token_store': WindowsTokenStore('local-workflow', backend=vault),
                    'provider_factory': lambda account: provider, 'now': lambda: clock[0]}
    try:
        with patch.object(GmailProvider, '_request', no_network), \
                patch.object(WindowsTokenStore, 'backend', new=property(lambda store: vault)):
            template = ROOT / 'docs/abnahme/A1_Layout/Vorlage.docx'
            original_template_hash = sha(template)
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'Empfänger'
            sheet.append(['Name', 'Email', 'Firma', 'Anrede', 'Anschreiben', 'Stelle'])
            people = [('Ada Beispiel', 'ada@example.test', 'Musterfirma A'),
                      ('Grace Beispiel', 'grace@example.test', 'Musterfirma B')]
            for name, email, company in people:
                sheet.append([name, email, company, 'Guten Tag', 'Nur lokales Anschreiben für ' + company, 'Entwicklung'])
            source = target / 'Empfaenger.xlsx'
            workbook.save(source)
            workbook.close()
            source_hash = sha(source)
            project = create_project(target / 'Projekt', name='A7: vollständiger lokaler Ablauf',
                                     initial_profile={'display_name': 'Lokale Testperson', 'email': 'applicant@example.test'},
                                     initial_signature={'name': 'Lokal', 'content_text': 'Lokale Prüfsignatur'})
            db = project.connection
            importer = ImportService(project)
            run = importer.prepare(str(source), 'contacts', {'display_name': 'Name', 'email': 'Email'}, options={'sheet': 'Empfänger'})
            assert run['counts'] == {'Valid': 2}
            rows = importer.list_rows(run['id'])['rows']
            importer.process(run['id'])
            assert db.execute('SELECT COUNT(*) FROM contacts').fetchone()[0] == 2
            checked('real_xlsx_imported_two_contacts')
            mapping = {'contact.display_name': 'Name', 'contact.email': 'Email', 'company.name': 'Firma',
                       'contact.salutation': 'Anrede', 'application.cover_letter': 'Anschreiben', 'job.title': 'Stelle'}
            sources = [{'kind': 'import', 'row_id': row['id'], 'mapping': mapping} for row in rows]
            fields = [BulkSourceService(db).resolve(item)['fields'] for item in sources]
            library = DataSourceLibrary(db)
            saved_list = library.save_list('Gemeinsame Excel-Werte', [{'fields': value} for value in fields])['list']
            list_sources = [{'kind': 'list', 'list_id': saved_list['id'], 'row_id': row['id']} for row in saved_list['rows']]
            assert [BulkSourceService(db).resolve(s)['fields'] for s in list_sources] == fields
            checked('saved_shared_list_preserves_excel_values')
            rules = [{'find': '{{' + token + '}}', 'replacement': '', 'field': field}
                     for token, field in [('NAME', 'contact.display_name'), ('COMPANY', 'company.name'), ('EMAIL', 'contact.email')]]
            word = WordReplaceService(project)
            word_options = {'source_path': str(template), 'rules': rules, 'sources': sources, 'file_format': 'pdf',
                            'libreoffice_path': find_libreoffice()}
            preview = word.preview(**word_options)
            assert preview['ready']
            outputs = word.create(expected_fingerprint=preview['fingerprint'], output_directory=str(project.path / 'exports/personalized'),
                                  prefix='Anschreiben', **word_options)
            report['word_pdfs'] = [pdf_values(path, value) for path, value in zip(outputs['paths'], fields)]
            checked('real_word_replace_personal_two_page_pdfs')
            documents = DocumentService(project)
            version = documents.import_template(template, name='A7-Anschreiben',
                        mapping={'NAME': 'contact.display_name', 'COMPANY': 'company.name', 'EMAIL': 'contact.email'})['version']
            mail = MailService(project, **mail_options)
            account = mail.register_account('sender@example.test', {'refresh_token': 'synthetic-memory-only'})
            mail_version = MailTemplateService(db).save(name='Lokale Bewerbung', subject='LOKAL {{company.name}} – {{contact.display_name}}',
                              body='{{contact.salutation}} {{contact.display_name}},\n{{application.cover_letter}}')
            bulk = BulkMailService(project)
            options = {'version_id': mail_version['id'], 'provider_account_id': account['id'],
                       'items': [{'contact_id': 'import:' + row['id'], 'source': reference, 'create_application': True}
                                 for row, reference in zip(rows, sources)],
                       'document_version_id': version['id'], 'document_format': 'pdf'}
            preview = bulk.preview(**options)
            assert preview['ready'], preview['rows']
            batch = bulk.create(name='Lokale persönliche Zweiergruppe', expected_fingerprint=preview['fingerprint'], **options)['batch']
            assert db.execute('SELECT COUNT(*) FROM applications').fetchone()[0] == 2
            for item, value in zip(batch['items'], fields):
                draft = mail.draft(item['draft_id'])
                assert value['company.name'] in draft['subject'] and value['contact.display_name'] in draft['body']
                entry = draft['configuration']['attachments'][0]
                pdf_values(project.file_path(entry['path']), value)
                assert draft['application_id']
            checked('bulk_applications_personal_mail_and_pdf_captured')
            delivery = BulkDeliveryService(project, mail=mail)
            group = delivery.begin(batch['id'], account['id'], interval_seconds=1, daily_budget=2)
            for _ in batch['items']:
                group = delivery.prepare_next(group['id'])
            assert group['status'] == 'Ready'
            group = delivery.approve(group['id'], group['fingerprint'])
            group = delivery.start(group['id'], group['fingerprint'], background=True)
            group = delivery.next(group['id'])
            assert group['counts']['sent'] == 1 and len(provider.sent) == 1
            delivery.pause(group['id'])
            dispatch_id, project_path = group['id'], project.path
            project.close()
            project = open_project(project_path)
            db = project.connection
            mail = MailService(project, **mail_options)
            mail.recover()
            delivery = BulkDeliveryService(project, mail=mail)
            group = delivery.snapshot(dispatch_id)
            assert group['status'] == 'Paused' and group['counts']['sent'] == 1
            clock[0] += timedelta(seconds=2)
            group = delivery.approve(group['id'], group['fingerprint'], confirm_overdue=True)
            group = delivery.start(group['id'], group['fingerprint'], background=True)
            group = delivery.next(group['id'])
            assert group['status'] == 'Completed' and group['counts']['sent'] == 2 and len(provider.sent) == 2
            delivery.next(group['id'])
            assert len(provider.sent) == 2
            assert db.execute("SELECT COUNT(*) FROM application_mail_receipts WHERE status='Sent'").fetchone()[0] == 2
            for raw, value in zip(provider.sent, fields):
                message = message_from_bytes(raw, policy=policy.default)
                assert str(message['To']) == value['contact.email']
                attachments = list(message.iter_attachments())
                assert len(attachments) == 1 and attachments[0].get_content_type() == 'application/pdf'
            checked('fake_bulk_pause_reopen_resume_skips_sent_and_updates_applications')

            graph = starter_graph()
            data, subject, body, sink = graph['nodes']
            data['config']['sources'] = list_sources
            subject['config']['text'] = 'BUILDER {{company.name}}'
            body['config']['text'] = '{{contact.salutation}} {{contact.display_name}}\n{{application.cover_letter}}'
            sink['config']['body_format'] = 'html'
            doc, replace, convert, rename, image = [node(kind) for kind in ('docx', 'replace', 'convert', 'attachment', 'inline_image')]
            doc['config']['path'] = str(template)
            replace['config']['rules'] = rules
            rename['config']['filename'] = 'Builder-{{company.name}}.pdf'
            image['config'].update(path=str(ROOT / 'assets/Logo256.png'), alt='Lokales Logo', width=80)
            graph['nodes'].extend([doc, replace, convert, rename, image])
            for a, out, b, port in [(doc, 'file', replace, 'docx'), (data, 'values', replace, 'values'),
                                   (replace, 'file', convert, 'docx'), (convert, 'file', rename, 'file'),
                                   (data, 'values', rename, 'values'), (rename, 'file', sink, 'attachments'),
                                   (image, 'text', sink, 'body')]:
                graph['edges'].append({'from': a['id'], 'output': out, 'to': b['id'], 'input': port})
            builder = EmailBuilderService(project)
            preview = builder.preview(graph, libreoffice_path=find_libreoffice())
            assert preview['ready'], preview['rows']
            builder_batch = builder.create('Gemeinsame Liste im Builder', graph, preview['fingerprint'], libreoffice_path=find_libreoffice())['batch']
            for item, value in zip(builder_batch['items'], fields):
                draft = mail.draft(item['draft_id'])
                assert draft['recipients']['to'] == [value['contact.email']]
                attached = draft['configuration']['attachments']
                assert len(attached) == 2 and sum(bool(e.get('cid')) for e in attached) == 1
                pdf = next(entry for entry in attached if not entry.get('cid'))
                assert pdf['filename'] == 'Builder-' + value['company.name'] + '.pdf'
                pdf_values(project.file_path(pdf['path']), value)
            state = {'graph': graph, 'history': [deepcopy(graph)], 'future': [], 'name': 'A7-Arbeitsstand',
                     'builder_id': None, 'version_id': None, 'provider_account_id': account['id'], 'signature_id': None,
                     'last_batch': builder_batch['id'], 'delivery_mode': 'drafts', 'scheduled_local': '', 'timezone_id': 'Europe/Berlin',
                     'interval_seconds': 2, 'view': {'zoom': 1.2, 'x': 20, 'y': 10}}
            BuilderWorkspaceService(project).save(project.project_id, state, 0)
            checked('shared_list_builder_real_pdf_inline_image_and_checkpoint')
            exports = ExportService(project)
            contacts_preview = exports.preview('contacts')
            csv_path = project.path / 'exports/contacts.csv'
            result = exports.write(preview=contacts_preview, columns=['display_name', 'email'], target_path=str(csv_path))
            AdministrationService(project).log_export('records', result, {'entity': 'contacts'})
            with csv_path.open(encoding='utf-8-sig', newline='') as stream:
                assert [row[1] for row in list(csv.reader(stream, delimiter=';'))[1:]] == [p[1] for p in people]
            jobs_preview = exports.preview('email_jobs', batch_id=batch['id'])
            xlsx_path = project.path / 'exports/status.xlsx'
            result = exports.write(preview=jobs_preview, columns=['subject', 'status', 'attempts'], target_path=str(xlsx_path), file_format='xlsx')
            AdministrationService(project).log_export('jobs', result, {'batch_id': batch['id']})
            book = load_workbook(xlsx_path, read_only=True)
            assert [row[1:] for row in list(book.active.values)[1:]] == [('Sent', 1), ('Sent', 1)]
            book.close()
            checked('csv_xlsx_exports_include_personal_values_sent_status_and_history')
            future = mail.save_draft(subject='LOKAL nicht ausführen', body='Wiederherstellung darf diesen Auftrag nicht aktivieren.',
                                     recipients_value={'to': ['future@example.test']}, provider_account_id=account['id'])
            future_job = mail.prepare_job(future['id'], scheduled_local=(clock[0] + timedelta(days=1)).isoformat(), timezone_id='UTC')
            mail.approve(future_job['id'])
            expected_files = {str(path.relative_to(project.path)).replace('\\', '/'): sha(path)
                              for folder in ('documents', 'templates', 'imports', 'exports')
                              for path in (project.path / folder).rglob('*') if path.is_file()}
            archive = project.create_backup()
            inspection = inspect_backup(archive)
            assert inspection['verified_index']
            restored_path = target / 'Wiederhergestellt'
            restore_backup(archive, restored_path, inspection['hash'])
            with open_project(restored_path) as restored:
                rdb = restored.connection
                assert rdb.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
                assert not rdb.execute('PRAGMA foreign_key_check').fetchall()
                assert rdb.execute('SELECT COUNT(*) FROM contacts').fetchone()[0] == 2
                assert rdb.execute('SELECT COUNT(*) FROM applications').fetchone()[0] == 2
                assert rdb.execute("SELECT COUNT(*) FROM email_jobs WHERE status='Sent'").fetchone()[0] == 2
                assert not rdb.execute('SELECT 1 FROM email_provider_accounts WHERE is_enabled=1').fetchone()
                assert not rdb.execute('SELECT 1 FROM email_bulk_dispatches WHERE background=1').fetchone()
                restored_future = dict(rdb.execute('SELECT * FROM email_jobs WHERE id=?', (future_job['id'],)).fetchone())
                assert restored_future['approval_status'] == 'pending' and restored_future['requires_overdue_confirmation']
                for relative, digest in expected_files.items():
                    assert sha(restored.file_path(relative)) == digest, relative
                assert BuilderWorkspaceService(restored).load()['state']['graph'] == graph
                diagnostic = diagnostic_report(restored)
                serialized = json.dumps(diagnostic, ensure_ascii=False)
                assert not any(secret in serialized for secret in ('synthetic-memory-only', 'Ada Beispiel', 'ada@example.test', str(project_path)))
                (target / 'Diagnostic.json').write_text(json.dumps(diagnostic, indent=2) + '\n', encoding='utf-8')
            checked('backup_restore_all_file_hashes_domain_state_checkpoint_and_dispatch_guards')
            assert sha(template) == original_template_hash and sha(source) == source_hash
            assert len(provider.sent) == 2 and report['real_gmail_requests'] == 0
            checked('external_originals_unchanged_no_real_credentials_or_mail')
            report.update(status='Completed', fake_send_attempts=2, fake_sent=2, verified_restored_files=len(expected_files),
                          original_docx_sha256=original_template_hash, original_xlsx_sha256=source_hash,
                          tested_source_hashes=before, elapsed_seconds=round(time.monotonic() - started, 3))
            after = code_hashes()
            report['code_changed_during_run'] = sorted(name for name in set(before) | set(after) if before.get(name) != after.get(name))
            assert not report['code_changed_during_run']
            save()
            print(json.dumps({key: value for key, value in report.items() if key not in ('tested_source_hashes', 'word_pdfs')}), flush=True)
    except Exception as error:
        report.update(status='Failed', error_type=type(error).__name__)
        save()
        raise
    finally:
        if project is not None:
            project.close()


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=ROOT / 'docs/abnahme/A7_Workflow/Full_Workflow')
    main(parser.parse_args().output)
