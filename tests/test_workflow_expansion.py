"""Shared libraries, atomic exports and reviewed Word copies; no live mail."""

import csv
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from emailcenter.application.workspace import WorkspaceService
from emailcenter.bulk_sources import BulkSourceService
from emailcenter.data_source_library import DataSourceLibrary
from emailcenter.documents import DocumentError, extract_docx_text
from emailcenter.export_service import EXPORTS, ExportService
from emailcenter.import_service import ImportService
from emailcenter.mail_bulk import BulkMailService
from emailcenter.mail_templates import MailTemplateService
from emailcenter.word_replace import WordReplaceService, file_hash, replace_docx
from test_documents import W, _make_docx


class ProjectFixture(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ROOT)
        self.root = Path(self.temp.name)
        self.workspace = WorkspaceService()
        self.workspace.create(str(self.root / "project"), "Workflows")
        self.project, self.db = self.workspace.project, self.workspace.project.connection
        self.library = DataSourceLibrary(self.db)
        self.source = self.root / "original.docx"
        _make_docx(self.source)

    def tearDown(self):
        self.workspace.close()
        self.temp.cleanup()

    def saved_list(self):
        return self.library.save_list("Bewerbungen", [{"fields": {"contact.display_name": name,
            "contact.email": name.lower() + "@example.test", "company.name": name + " GmbH",
            "application.cover_letter": "Anschreiben für " + name}} for name in ("Ada", "Grace")])["list"]

    def import_rows(self):
        from openpyxl import Workbook
        path = self.root / "source.xlsx"
        book = Workbook()
        book.active.append(["Name", "Email", "Firma", "Anrede"])
        book.active.append(["Ada", "ada@example.test", "Ada GmbH", ""])
        book.active.append(["Grace", "grace@example.test", "Grace GmbH", "Frau"])
        book.save(path)
        book.close()
        importer = ImportService(self.project)
        run = importer.prepare(str(path), "contacts", {"display_name": "Name", "email": "Email"})
        rows = importer.list_rows(run["id"])["rows"]
        mapping = {"contact.display_name": "Name", "contact.email": "Email", "company.name": "Firma", "contact.salutation": "Anrede"}
        return rows, mapping

    def contact(self, name="Ada", **fields):
        return self.workspace.create_record("contacts", {"display_name": name, "email": name.lower() + "@example.test", **fields})


class SourceLibraryTests(ProjectFixture):
    def test_profile_roundtrip_edit_and_stale_update(self):
        original = self.library.save_profile("Excel", {"contact.email": "Email"}, {"contact.salutation": "Guten Tag"})["profile"]
        edited = self.library.save_profile("Excel", original["mapping"], {"contact.salutation": "Hallo"}, profile_id=original["id"], expected_version=original["version"])["profile"]
        self.assertNotEqual(original["version"], edited["version"])
        with self.assertRaisesRegex(ValueError, "geändert"):
            self.library.save_profile("Excel", {}, profile_id=original["id"], expected_version=original["version"])
        self.workspace.close()
        self.workspace.open(str(self.root / "project"))
        self.assertEqual(DataSourceLibrary(self.workspace.project.connection).profile(edited["id"])["defaults"]["contact.salutation"], "Hallo")

    def test_profile_validation_and_unique_name(self):
        self.library.save_profile("Excel", {})
        for changes in ({"name": "excel"}, {"mapping": {"bad field": "Name"}}, {"defaults": {"bad field": "x"}}, {"empty_policy": "ignore"}):
            with self.subTest(changes=changes), self.assertRaises(ValueError):
                self.library.save_profile(**{"name": "New", "mapping": {}, **changes})
        self.assertEqual(len(self.library.profiles()), 1)

    def test_defaults_empty_policy_and_personal_override(self):
        rows, mapping = self.import_rows()
        source = {"kind": "import", "row_id": rows[0]["id"], "mapping": mapping, "defaults": {"contact.salutation": "Guten Tag", "custom": "Standard"}}
        resolved = BulkSourceService(self.db).resolve(source)
        self.assertEqual(resolved["fields"]["contact.salutation"], "Guten Tag")
        self.assertEqual(resolved["fields"]["custom"], "Standard")
        self.assertEqual(BulkSourceService(self.db).resolve({**source, "empty_policy": "keep"})["fields"]["contact.salutation"], "")
        version = MailTemplateService(self.db).save(name="Brief", subject="Hallo", body="{{contact.salutation}} {{contact.display_name}}")
        preview = BulkMailService(self.project).preview(version_id=version["id"], items=[{"contact_id": "import:" + rows[0]["id"], "source": source, "values": {"contact.salutation": "Persönlich"}}])
        self.assertTrue(preview["ready"])
        self.assertIn("Persönlich Ada", preview["rows"][0]["body"])

    def test_list_edit_keeps_ids_and_reorders_live_sources(self):
        value = self.saved_list()
        refs = [{"kind": "list", "list_id": value["id"], "row_id": row["id"]} for row in value["rows"]]
        changed = [{"id": row["id"], "fields": {**row["fields"], "company.name": "Neu"}} for row in reversed(value["rows"])]
        saved = self.library.save_list("Bewerbungen", changed, list_id=value["id"], expected_version=value["version"])["list"]
        self.assertEqual(saved["rows"][0]["id"], value["rows"][1]["id"])
        self.assertEqual(BulkSourceService(self.db).resolve(refs[0])["fields"]["company.name"], "Neu")
        selected = BulkSourceService(self.db).select_list(value["id"], [r["id"] for r in value["rows"]])["contacts"]
        self.assertEqual([r["display_name"] for r in selected], ["Grace", "Ada"])
        with self.assertRaisesRegex(ValueError, "geändert"):
            self.library.save_list("Bewerbungen", changed, list_id=value["id"], expected_version=value["version"])

    def test_removed_row_and_cross_list_row_reuse_are_rejected(self):
        value = self.saved_list()
        with self.assertRaisesRegex(ValueError, "anderen Liste"):
            self.library.save_list("Copy", value["rows"])
        self.library.save_list(value["name"], [value["rows"][1]], list_id=value["id"], expected_version=value["version"])
        with self.assertRaisesRegex(ValueError, "entfernt"):
            BulkSourceService(self.db).resolve({"kind": "list", "list_id": value["id"], "row_id": value["rows"][0]["id"]})

    def test_list_validation_and_rollback(self):
        value = self.saved_list()
        for rows in ([], [{}], [{"fields": {}}], [value["rows"][0]] * 2, [{"id": "invalid", "fields": {"name": "x"}}], [{"fields": {"bad field": "x"}}]):
            with self.subTest(rows=rows), self.assertRaises(ValueError):
                self.library.save_list("Wrong", rows)
        with patch("emailcenter.data_source_library._write_activity", side_effect=RuntimeError("db failure")):
            with self.assertRaises(RuntimeError):
                self.library.save_list(value["name"], [value["rows"][0]], list_id=value["id"], expected_version=value["version"])
        self.assertEqual(self.library.list(value["id"]), value)

    def test_list_without_email_is_usable_for_documents(self):
        value = self.library.save_list("Nur Dokumente", [{"fields": {"person": "Ada"}}])["list"]
        selected = BulkSourceService(self.db).select_list(value["id"], [value["rows"][0]["id"]])
        self.assertTrue(selected["contacts"][0]["error"])
        resolved = BulkSourceService(self.db).resolve(selected["contacts"][0]["source"])
        self.assertEqual(resolved["fields"]["person"], "Ada")

    def test_source_change_blocks_bulk_create_and_preserves_staged_snapshot(self):
        value = self.saved_list()
        version = MailTemplateService(self.db).save(name="Brief", subject="{{company.name}}", body="Hallo")
        selected = BulkSourceService(self.db).select_list(value["id"], [value["rows"][0]["id"]])["contacts"][0]
        options = {"version_id": version["id"], "items": [{"contact_id": selected["id"], "source": selected["source"]}]}
        preview = self.workspace.mail_bulk_preview(**options)
        changed = [{"id": r["id"], "fields": {**r["fields"], "company.name": "Neu"}} for r in value["rows"]]
        self.library.save_list(value["name"], changed, list_id=value["id"], expected_version=value["version"])
        with self.assertRaisesRegex(ValueError, "geändert"):
            self.workspace.mail_bulk_create(preview_id=preview["preview_id"], name="Group")
        preview = self.workspace.mail_bulk_preview(**options)
        result = self.workspace.mail_bulk_create(preview_id=preview["preview_id"], name="Group")
        self.assertEqual(self.db.execute("SELECT subject FROM email_drafts WHERE id=?", (result["batch"]["items"][0]["draft_id"],)).fetchone()[0], "Neu")
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM email_jobs").fetchone()[0], 0)


class ExportTests(ProjectFixture):
    def test_csv_selection_filter_columns_and_formula_escape(self):
        ada, grace = self.contact(phone="=1+1"), self.contact("Grace")
        self.contact("100%_literal")
        exporter = ExportService(self.project)
        self.assertEqual(exporter.preview(search="%_")["total"], 1)
        self.assertEqual(exporter.preview(search="Ada")["total"], 1)
        preview = self.workspace.export_preview(entity="contacts", selected_ids=[ada["id"]])
        path = self.root / "contacts.csv"
        result = self.workspace.export_write(preview_id=preview["preview_id"], columns=["display_name", "phone"], target_path=str(path))
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, delimiter=";"))
        self.assertEqual(rows, [["Name", "Telefon"], ["Ada", "'=1+1"]])
        self.assertEqual(result["count"], 1)
        self.assertEqual(self.db.execute("SELECT phone FROM contacts WHERE id=?", (ada["id"],)).fetchone()[0], "=1+1")

    def test_all_export_entities_query_and_exclude_oauth_data(self):
        for entity in EXPORTS:
            value = ExportService(self.project).preview(entity)
            self.assertTrue(value["columns"])
            keys = [c["key"] for c in value["columns"]]
            self.assertFalse(any("token" in key or "credential" in key or "config" in key for key in keys))

    def test_xlsx_formula_strings_and_dates_remain_correct_types(self):
        from openpyxl import load_workbook
        self.contact(phone="=1+1", notes="ü\nLine")
        service = ExportService(self.project)
        preview = service.preview()
        path = self.root / "contacts.xlsx"
        service.write(preview=preview, columns=["phone", "notes"], target_path=str(path), file_format="xlsx")
        book = load_workbook(path)
        self.assertEqual(book.active["A2"].value, "=1+1")
        self.assertEqual(book.active["A2"].data_type, "s")
        self.assertEqual(book.active["B2"].value, "ü\nLine")
        self.assertEqual(book.active.freeze_panes, "A2")
        book.close()

    def test_changed_data_requires_new_preview(self):
        contact = self.contact()
        preview = self.workspace.export_preview()
        self.db.execute("UPDATE contacts SET display_name='Changed' WHERE id=?", (contact["id"],))
        self.db.commit()
        path = self.root / "out.csv"
        with self.assertRaisesRegex(ValueError, "geändert"):
            self.workspace.export_write(preview_id=preview["preview_id"], columns=["display_name"], target_path=str(path))
        self.assertFalse(path.exists())

    def test_target_protection_overwrite_and_atomic_failure(self):
        self.contact()
        service = ExportService(self.project)
        options = {"preview": service.preview(), "columns": ["display_name"]}
        with self.assertRaisesRegex(ValueError, "exports"):
            service.write(**options, target_path=str(self.project.path / "project.csv"))
        target = self.root / "out.csv"
        target.write_text("original")
        with self.assertRaisesRegex(ValueError, "existiert"):
            service.write(**options, target_path=str(target))
        with patch("emailcenter.export_service.os.replace", side_effect=PermissionError("locked")), self.assertRaises(PermissionError):
            service.write(**options, target_path=str(target), overwrite=True)
        self.assertEqual(target.read_text(), "original")
        self.assertEqual(list(self.root.glob(".emailcenter-export-*")), [])
        service.write(**options, target_path=str(target), overwrite=True)
        self.assertIn("Ada", target.read_text(encoding="utf-8-sig"))

    def test_cancel_and_target_change_preserve_existing_file(self):
        self.contact()
        service = ExportService(self.project)
        path = self.root / "out.csv"
        path.write_text("original")
        with self.assertRaisesRegex(ValueError, "abgebrochen"):
            service.write(preview=service.preview(), columns=["display_name"], target_path=str(path), overwrite=True, cancelled=lambda: True)
        self.assertEqual(path.read_text(), "original")
        calls = []
        def change():
            calls.append(1)
            if len(calls) == 2:
                path.write_text("external")
            return False
        with self.assertRaisesRegex(ValueError, "zwischenzeitlich"):
            service.write(preview=service.preview(), columns=["display_name"], target_path=str(path), overwrite=True, cancelled=change)
        self.assertEqual(path.read_text(), "external")
        self.assertEqual(list(self.root.glob(".emailcenter-export-*")), [])

    def test_excel_length_encoding_and_invalid_selection_fail_without_output(self):
        self.contact(notes="x" * 32768)
        service = ExportService(self.project)
        with self.assertRaisesRegex(ValueError, "lang"):
            service.write(preview=service.preview(), columns=["notes"], target_path=str(self.root / "out.xlsx"), file_format="xlsx")
        self.assertFalse((self.root / "out.xlsx").exists())
        for options in ({"entity": "activity_log"}, {"selected_ids": []}, {"selected_ids": ["a", "a"]}, {"batch_id": "wrong"}):
            with self.subTest(options=options), self.assertRaises(ValueError):
                service.preview(**options)
        with self.assertRaises(ValueError):
            service.write(preview=service.preview(), columns=["token"], target_path=str(self.root / "out.csv"))

    def test_deleted_contacts_excluded_and_closed_preview_rejected(self):
        contact = self.contact()
        self.db.execute("UPDATE contacts SET deleted_at='now' WHERE id=?", (contact["id"],))
        self.db.commit()
        preview = self.workspace.export_preview()
        self.assertEqual(preview["total"], 0)
        self.workspace.close()
        self.workspace.open(str(self.root / "project"))
        with self.assertRaisesRegex(ValueError, "erneut"):
            self.workspace.export_write(preview_id=preview["preview_id"], columns=["display_name"], target_path=str(self.root / "out.csv"))

    def test_date_export_uses_native_excel_date_in_utc(self):
        from datetime import datetime
        from openpyxl import load_workbook
        company = self.workspace.create_record("companies", {"name": "Example"})
        applicant = self.workspace.create_record("applicants", {"display_name": "Ada"})
        application = self.workspace.create_record("applications", {"title": "Developer", "company_id": company["id"], "applicant_profile_id": applicant["id"]})
        self.db.execute("UPDATE applications SET next_follow_up_at='2026-10-12T12:30:00+02:00' WHERE id=?", (application["id"],))
        self.db.commit()
        service = ExportService(self.project)
        target = self.root / "dates.xlsx"
        service.write(preview=service.preview("applications"), columns=["title", "next_follow_up_at"], target_path=str(target), file_format="xlsx")
        book = load_workbook(target)
        self.assertEqual(book.active["B2"].value, datetime(2026, 10, 12, 10, 30))
        self.assertEqual(book.active["B2"].data_type, "d")
        book.close()

    def test_csv_unencodable_value_preserves_old_output(self):
        self.contact(notes="漢字")
        target = self.root / "out.csv"
        target.write_text("old")
        service = ExportService(self.project)
        with self.assertRaises(UnicodeEncodeError):
            service.write(preview=service.preview(), columns=["notes"], target_path=str(target), encoding="cp1252", overwrite=True)
        self.assertEqual(target.read_text(), "old")
        self.assertEqual(list(self.root.glob(".emailcenter-export-*")), [])


class WordReplaceTests(ProjectFixture):
    def rules(self, field="", replacement="Ada"):
        return [{"find": "{{NAME}}", "replacement": replacement, "field": field}]

    def test_runs_tables_headers_and_original_are_preserved(self):
        original = self.source.read_bytes()
        target = self.root / "copy.docx"
        rules = self.rules() + [{"find": "{{ROLE}}", "replacement": "Engineer", "field": ""}, {"find": "{{DATE}}", "replacement": "2026", "field": ""}]
        result = replace_docx(self.source, target, rules)
        text = extract_docx_text(target)
        self.assertIn("Hello Ada", text)
        self.assertIn("Engineer", text)
        self.assertIn("2026", text)
        self.assertEqual(sum(result["replaced"].values()), 3)
        self.assertEqual(self.source.read_bytes(), original)
        with zipfile.ZipFile(target) as archive:
            xml = archive.read("word/document.xml").decode()
            self.assertIn("<w:b", xml)
            self.assertIn("xmlns:w14", xml)

    def simple_source(self, xml):
        with zipfile.ZipFile(self.source, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", f'<w:document xmlns:w="{W}"><w:body>{xml}</w:body></w:document>')
        return self.source

    def test_case_words_nonrecursive_longest_match_and_newlines(self):
        self.simple_source("<w:p><w:r><w:t>cat CAT scatter cat dog</w:t></w:r></w:p>")
        rules = [{"find": "cat", "replacement": "dog\nLine\tTab", "field": ""}, {"find": "dog", "replacement": "bird", "field": ""}]
        target = self.root / "copy.docx"
        result = replace_docx(self.source, target, rules, case_sensitive=False, whole_words=True)
        text = extract_docx_text(target)
        self.assertIn("scatter", text)
        self.assertIn("dog\nLine\tTab", text)
        self.assertEqual(result["replaced"], {"cat": 3, "dog": 1})

    def test_rules_do_not_cross_breaks_tabs_or_paragraphs(self):
        self.simple_source("<w:p><w:r><w:t>ca</w:t><w:br/><w:t>t</w:t><w:tab/><w:t>ca</w:t></w:r></w:p><w:p><w:r><w:t>t</w:t></w:r></w:p>")
        result = replace_docx(self.source, self.root / "copy.docx", [{"find": "cat", "replacement": "bird", "field": ""}])
        self.assertEqual(result["replaced"]["cat"], 0)

    def test_missing_fields_block_or_keep_and_literal_empty_deletes(self):
        target = self.root / "copy.docx"
        with self.assertRaisesRegex(ValueError, "Fehlende"):
            replace_docx(self.source, target, self.rules("person"))
        self.assertFalse(target.exists())
        result = replace_docx(self.source, target, self.rules("person"), missing_policy="keep")
        self.assertEqual(result["missing"], ["person"])
        self.assertIn("{{NAME}}", extract_docx_text(target))
        replace_docx(self.source, target, self.rules(replacement=""))
        self.assertNotIn("{{NAME}}", extract_docx_text(target))

    def test_invalid_rules_package_and_original_target_are_rejected(self):
        for rules in ([], self.rules() * 2, [{"find": "a\nb", "replacement": "x", "field": ""}], [{"find": "", "replacement": "x", "field": ""}]):
            with self.subTest(rules=rules), self.assertRaises(ValueError):
                replace_docx(self.source, self.root / "copy.docx", rules)
        with self.assertRaises(DocumentError):
            replace_docx(self.source, self.source, self.rules())
        self.simple_source('<!DOCTYPE bad [<!ENTITY name "bad">]><w:p/>')
        with self.assertRaises(DocumentError):
            replace_docx(self.source, self.root / "copy.docx", self.rules())

    def test_single_output_registered_and_preview_consumed(self):
        preview = self.workspace.word_replace_preview(source_path=str(self.source), rules=self.rules())
        self.assertTrue(preview["ready"])
        self.assertEqual(list((self.project.path / "tmp").glob("word-replace-*")), [])
        result = self.workspace.word_replace_create(preview_id=preview["preview_id"], output_directory=str(self.project.path / "exports/word"), prefix="Anschreiben")
        self.assertEqual(result["count"], 1)
        document = result["documents"][0]
        self.assertEqual(document["document_type"], "word_replace")
        self.assertEqual(document["file_hash"], file_hash(result["paths"][0]))
        with self.assertRaisesRegex(ValueError, "erneut"):
            self.workspace.word_replace_create(preview_id=preview["preview_id"], output_directory=str(self.root / "out"))

    def test_saved_list_produces_personal_copies(self):
        value = self.saved_list()
        sources = [{"kind": "list", "list_id": value["id"], "row_id": row["id"]} for row in value["rows"]]
        service = WordReplaceService(self.project)
        options = {"source_path": str(self.source), "rules": self.rules("contact.display_name"), "sources": sources}
        preview = service.preview(**options)
        result = service.create(**options, expected_fingerprint=preview["fingerprint"], output_directory=str(self.root / "out"))
        self.assertEqual(result["count"], 2)
        self.assertIn("Hello Ada", extract_docx_text(result["paths"][0]))
        self.assertIn("Hello Grace", extract_docx_text(result["paths"][1]))
        self.assertEqual(result["documents"], [])

    def test_excel_rows_supply_document_values(self):
        rows, mapping = self.import_rows()
        refs = [{"kind": "import", "row_id": row["id"], "mapping": mapping} for row in rows]
        preview = WordReplaceService(self.project).preview(str(self.source), self.rules("company.name"), sources=refs)
        self.assertTrue(preview["ready"])
        self.assertIn("Hello Ada GmbH", preview["rows"][0]["text"])
        self.assertIn("Hello Grace GmbH", preview["rows"][1]["text"])

    def test_company_import_without_email_supplies_word_data_and_cannot_send(self):
        path = self.root / "companies.csv"
        path.write_text("Firma;Notiz\nExample GmbH;Dokumentdaten\n", encoding="utf-8")
        importer = ImportService(self.project)
        run = importer.prepare(str(path), "companies", {"name": "Firma", "notes": "Notiz"})
        row = importer.list_rows(run["id"])["rows"][0]
        source = BulkSourceService(self.db).select([row["id"]], {"company.name": "Firma"}, require_mail_fields=False)["contacts"][0]
        preview = WordReplaceService(self.project).preview(str(self.source), self.rules("company.name"), sources=[source["source"]])
        self.assertTrue(preview["ready"])
        self.assertIn("Hello Example GmbH", preview["rows"][0]["text"])
        version = MailTemplateService(self.db).save(name="Brief", subject="{{company.name}}", body="Hallo")
        mail = BulkMailService(self.project).preview(version_id=version["id"], items=[{"contact_id": source["id"], "source": source["source"]}])
        self.assertFalse(mail["ready"])
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM email_jobs").fetchone()[0], 0)

    def test_changed_original_blocks_reviewed_output(self):
        service = WordReplaceService(self.project)
        preview = service.preview(str(self.source), self.rules())
        self.simple_source("<w:p><w:r><w:t>New {{NAME}}</w:t></w:r></w:p>")
        with self.assertRaisesRegex(ValueError, "geändert"):
            service.create(source_path=str(self.source), rules=self.rules(), expected_fingerprint=preview["fingerprint"], output_directory=str(self.root / "out"))
        self.assertFalse((self.root / "out").exists())

    def test_changed_list_blocks_reviewed_output(self):
        value = self.saved_list()
        options = {"source_path": str(self.source), "rules": self.rules("contact.display_name"), "sources": [{"kind": "list", "list_id": value["id"], "row_id": value["rows"][0]["id"]}]}
        service = WordReplaceService(self.project)
        preview = service.preview(**options)
        changed = [{"id": r["id"], "fields": {**r["fields"], "contact.display_name": "New"}} for r in value["rows"]]
        self.library.save_list(value["name"], changed, list_id=value["id"], expected_version=value["version"])
        with self.assertRaisesRegex(ValueError, "geändert"):
            service.create(**options, expected_fingerprint=preview["fingerprint"], output_directory=str(self.root / "out"))

    def test_pdf_preflight_and_conversion_adapter(self):
        service = WordReplaceService(self.project)
        with patch("emailcenter.word_replace.find_libreoffice", return_value=None), self.assertRaisesRegex(ValueError, "LibreOffice"):
            service.preview(str(self.source), self.rules(), file_format="pdf")
        executable = self.root / "soffice.exe"
        executable.write_bytes(b"fake")
        options = {"source_path": str(self.source), "rules": self.rules(), "file_format": "pdf", "libreoffice_path": str(executable), "conversion_timeout": 45}
        preview = service.preview(**options)
        def convert(source, output, **kw):
            self.assertIn("Hello Ada", extract_docx_text(source))
            self.assertEqual(kw["timeout_seconds"], 45)
            Path(output).write_bytes(b"%PDF-1.7 fake adapter output")
        with patch("emailcenter.word_replace.convert_docx_to_pdf", side_effect=convert):
            result = service.create(**options, expected_fingerprint=preview["fingerprint"], output_directory=str(self.root / "out"))
        self.assertEqual(result["format"], "pdf")
        self.assertTrue(Path(result["paths"][0]).read_bytes().startswith(b"%PDF"))

    def test_output_registration_failure_restores_all_files_and_database(self):
        value = self.saved_list()
        options = {"source_path": str(self.source), "rules": self.rules("contact.display_name"), "sources": [{"kind": "list", "list_id": value["id"], "row_id": r["id"]} for r in value["rows"]]}
        service = WordReplaceService(self.project)
        preview = service.preview(**options)
        directory = self.project.path / "exports/word"
        directory.mkdir(parents=True)
        first = directory / "Dokument-001-Ada.docx"
        first.write_bytes(b"original output")
        from emailcenter.database import DocumentRepository
        create = DocumentRepository.create
        calls = []
        def fail_second(repo, **kw):
            calls.append(1)
            if len(calls) == 2:
                raise RuntimeError("disk failure")
            return create(repo, **kw)
        with patch.object(DocumentRepository, "create", fail_second), self.assertRaises(RuntimeError):
            service.create(**options, expected_fingerprint=preview["fingerprint"], output_directory=str(directory), overwrite=True)
        self.assertEqual(first.read_bytes(), b"original output")
        self.assertEqual([p.name for p in directory.iterdir()], [first.name])
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM documents").fetchone()[0], 0)

    def test_cancel_publication_rolls_back_new_outputs(self):
        value = self.saved_list()
        options = {"source_path": str(self.source), "rules": self.rules("contact.display_name"), "sources": [{"kind": "list", "list_id": value["id"], "row_id": r["id"]} for r in value["rows"]]}
        service = WordReplaceService(self.project)
        preview = service.preview(**options)
        directory = self.project.path / "exports/word"
        def cancel_after_first():
            return bool(directory.exists() and list(directory.glob("*.docx")))
        with self.assertRaisesRegex(ValueError, "abgebrochen"):
            service.create(**options, expected_fingerprint=preview["fingerprint"], output_directory=str(directory), cancelled=cancel_after_first)
        self.assertEqual(list(directory.iterdir()), [])
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM documents").fetchone()[0], 0)

    def test_invalid_prefix_and_project_output_protection(self):
        service = WordReplaceService(self.project)
        preview = service.preview(str(self.source), self.rules())
        for prefix in ("..", "CON", "a/b"):
            with self.subTest(prefix=prefix), self.assertRaises(ValueError):
                service.create(source_path=str(self.source), rules=self.rules(), expected_fingerprint=preview["fingerprint"], output_directory=str(self.root / "out"), prefix=prefix)
        with self.assertRaisesRegex(ValueError, "exports"):
            service.create(source_path=str(self.source), rules=self.rules(), expected_fingerprint=preview["fingerprint"], output_directory=str(self.project.path / "documents"))

    def test_longest_rule_wins_and_unresolved_tokens_are_reported(self):
        self.simple_source("<w:p><w:r><w:t>ab {{OPEN}}</w:t></w:r></w:p>")
        preview = WordReplaceService(self.project).preview(str(self.source), [
            {"find": "a", "replacement": "first", "field": ""}, {"find": "ab", "replacement": "long", "field": ""}])
        self.assertEqual(preview["rows"][0]["replaced"], {"a": 0, "ab": 1})
        self.assertEqual(preview["rows"][0]["unresolved"], ["OPEN"])
        self.assertIn("long {{OPEN}}", preview["rows"][0]["text"])

    def test_change_during_pdf_conversion_blocks_publication(self):
        value = self.saved_list()
        executable = self.root / "soffice.exe"
        executable.write_bytes(b"fake")
        options = {"source_path": str(self.source), "rules": self.rules("contact.display_name"), "file_format": "pdf", "libreoffice_path": str(executable),
                   "sources": [{"kind": "list", "list_id": value["id"], "row_id": value["rows"][0]["id"]}]}
        service = WordReplaceService(self.project)
        preview = service.preview(**options)
        def convert(source, target, **kwargs):
            Path(target).write_bytes(b"%PDF fake")
            self.db.execute("UPDATE data_source_list_rows SET fields_json=? WHERE id=?", (json.dumps({"contact.display_name": "Changed"}), value["rows"][0]["id"]))
            self.db.commit()
        directory = self.root / "out"
        with patch("emailcenter.word_replace.convert_docx_to_pdf", side_effect=convert), self.assertRaisesRegex(ValueError, "während"):
            service.create(**options, expected_fingerprint=preview["fingerprint"], output_directory=str(directory))
        self.assertEqual(list(directory.iterdir()), [])

    def test_locked_rollback_keeps_recovery_copies(self):
        service = WordReplaceService(self.project)
        preview = service.preview(str(self.source), self.rules())
        directory = self.project.path / "exports/word"
        directory.mkdir(parents=True)
        target = directory / "Dokument.docx"
        target.write_bytes(b"old output")
        original_replace = Path.replace
        def locked_backup(path, destination):
            if path.name.startswith("backup-"):
                raise PermissionError("locked")
            return original_replace(path, destination)
        with patch("emailcenter.word_replace._write_activity", side_effect=RuntimeError("db failure")), patch.object(Path, "replace", locked_backup), self.assertRaisesRegex(DocumentError, "Wiederherstellungsdateien"):
            service.create(source_path=str(self.source), rules=self.rules(), expected_fingerprint=preview["fingerprint"], output_directory=str(directory), overwrite=True)
        recoveries = list(directory.glob(".word-replace-recovery-*"))
        self.assertEqual(len(recoveries), 1)
        self.assertEqual((recoveries[0] / "backup-0").read_bytes(), b"old output")
        self.assertIn("Dokument.docx", (recoveries[0] / "WIEDERHERSTELLUNG.txt").read_text(encoding="utf-8"))
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM documents").fetchone()[0], 0)


if __name__ == "__main__":
    unittest.main()
