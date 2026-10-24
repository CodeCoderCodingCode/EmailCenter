"""Offscreen desktop integration checks; optional when Qt is not installed."""

import os
from pathlib import Path
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

try:
    from PySide6.QtCore import QEvent, QThread
    from PySide6.QtWidgets import QApplication, QMessageBox
    from emailcenter.ui.controller import WorkspaceController
    from emailcenter.ui.dialogs import GenerateDialog, PreviewDialog, TemplateDialog
    from emailcenter.ui.main_window import MainWindow
    from emailcenter.ui.pages import PlannedPage, RecordPage
    from emailcenter.ui.theme import apply_theme
    from emailcenter.preferences import Preferences, PreferencesStore
    QT_AVAILABLE = True
except ImportError:
    QT_AVAILABLE = False

from test_documents import _make_docx


@unittest.skipUnless(QT_AVAILABLE, "PySide6 not installed")
class DesktopTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.application = QApplication.instance() or QApplication([])
        cls.application.setQuitOnLastWindowClosed(False)
        # Windows' offscreen plugin does not enumerate installed system fonts.
        from PySide6.QtGui import QFontDatabase
        if not QFontDatabase.families():
            font = Path(os.environ.get("WINDIR", "C:/Windows")) / "Fonts" / "segoeui.ttf"
            if font.is_file():
                QFontDatabase.addApplicationFont(str(font))

    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(dir=ROOT)
        self.directory = Path(self.temporary.name)
        apply_theme(self.application)
        self.window = MainWindow(settings_store=PreferencesStore(self.directory / "preferences.json"))
        self.window.show()
        self.errors = []
        self.message_patch = patch.object(QMessageBox, "warning", side_effect=lambda *args: self.errors.append(args[-1]))
        self.message_patch.start()

    def tearDown(self):
        self.wait_idle()
        self.window.close()
        self.wait_for(lambda: not self.window.controller.thread.isRunning())
        self.application.processEvents()
        self.window.deleteLater()
        self.application.sendPostedEvents(None, QEvent.Type.DeferredDelete)
        self.application.processEvents()
        self.message_patch.stop()
        self.temporary.cleanup()

    def wait_for(self, condition, seconds=10):
        deadline = time.monotonic() + seconds
        while not condition() and time.monotonic() < deadline:
            self.application.processEvents()
            time.sleep(0.005)
        self.assertTrue(condition(), "Qt operation timed out")
        self.application.processEvents()

    def wait_idle(self):
        self.wait_for(lambda: not self.window.controller.busy)

    def submit(self, action, payload):
        self.assertTrue(self.window.controller.submit(action, payload))
        self.wait_idle()
        self.assertEqual(self.errors, [])

    def open_test_project(self):
        self.submit("create_project", {"path": str(self.directory / "project"), "name": "Meine Bewerbungen"})

    def test_window_starts_without_project_and_pages_can_be_extended(self):
        self.assertFalse(self.window.actions["backup"].isEnabled())
        self.assertFalse(self.window.pages["companies"].isEnabled())
        self.assertTrue(self.window.actions["new"].isEnabled())
        self.window.register_page("future", "Später", PlannedPage("Erweiterung", "Neue Ansicht"))
        self.window.navigate("future")
        self.assertIs(self.window.current_page(), self.window.pages["future"])

    def test_project_navigation_enters_workspace_and_close_returns_to_projects(self):
        self.assertIs(self.window.current_page(), self.window.pages["overview"])
        self.assertTrue(self.window.navigation.item(self.window.page_order.index("workspace")).isHidden())
        self.open_test_project()
        self.assertIs(self.window.current_page(), self.window.pages["workspace"])
        self.assertFalse(self.window.navigation.item(self.window.page_order.index("workspace")).isHidden())
        self.assertTrue(self.window.toast.isVisible())
        self.submit("close_project", {})
        self.assertIs(self.window.current_page(), self.window.pages["overview"])
        self.assertTrue(self.window.navigation.item(self.window.page_order.index("workspace")).isHidden())
        self.assertFalse(self.window.toast.isVisible())

    def test_workspace_mode_is_persisted_and_editors_open_without_navigation(self):
        self.open_test_project()
        page = self.window.pages["workspace"]
        self.assertFalse(page.buttons["application"].isEnabled())
        page.buttons["profile"].click()
        self.wait_idle()
        self.assertEqual(self.window._record_dialog.data["entity"], "applicants")
        self.window._record_dialog.reject()
        self.submit("set_workspace_mode", {"mode": "mail"})
        self.assertEqual(page.mode.currentData(), "mail")
        self.assertTrue(page.buttons["profile"].isHidden())
        self.assertTrue(page.buttons["compose"].isVisible())
        page.buttons["compose"].click()
        self.wait_idle()
        self.assertIs(self.window.current_page(), page)
        self.assertIsNotNone(self.window._mail_dialog)
        self.window._mail_dialog.reject()
        self.submit("close_project", {})
        self.submit("open_project", {"path": str(self.directory / "project")})
        self.assertEqual(page.mode.currentData(), "mail")

    def test_workspace_recent_application_editor_and_responsive_cards(self):
        self.open_test_project()
        company = []
        self.window.controller.completed.connect(lambda action, value: company.append(value) if action == "create_record" else None)
        self.submit("create_record", {"entity": "companies", "fields": {"name": "Example"}})
        self.submit("create_record", {"entity": "applicants", "fields": {"display_name": "Ada"}})
        self.submit("create_record", {"entity": "applications", "fields": {"title": "Engineer", "company_id": company[0]["id"], "applicant_profile_id": company[1]["id"]}})
        page = self.window.pages["workspace"]
        self.assertTrue(page.buttons["application"].isEnabled())
        self.assertIn("Engineer", page.recent.item(0).text())
        page.open_recent(page.recent.item(0))
        self.wait_idle()
        self.assertIs(self.window.current_page(), page)
        self.assertEqual(self.window._record_dialog.fields["title"].text(), "Engineer")
        self.window._record_dialog.reject()
        self.window.resize(1250, 820)
        self.application.processEvents()
        self.assertEqual(page.grid.getItemPosition(1)[1], 1)
        screenshot_path = os.environ.get("EMAILCENTER_WORKSPACE_SCREENSHOT_PATH")
        if screenshot_path:
            self.window.toast.dismiss()
            self.window.grab().save(screenshot_path)
        self.window.resize(820, 600)
        self.application.processEvents()
        self.assertEqual(page.grid.getItemPosition(1)[1], 0)

    def test_toast_plain_text_replacement_timeout_and_dismiss(self):
        toast = self.window.toast
        toast.notify("Erledigt", "<b>Literal</b>")
        self.assertEqual(toast.message.text(), "<b>Literal</b>")
        self.assertTrue(toast.isVisible())
        toast.notify("Neu", "Aktuell", duration=20)
        self.assertEqual(toast.message.text(), "Aktuell")
        self.wait_for(lambda: not toast.isVisible())
        toast.notify("Gespeichert", "Okay")
        toast.close_button.click()
        self.assertFalse(toast.isVisible())
        self.assertFalse(toast.timer.isActive())

    def test_project_setup_dialog_passes_mode_and_goal(self):
        self.window.create_project()
        dialog = self.window._project_dialog
        dialog.fields["name"].setText("Aufträge")
        dialog.fields["folder"].setText("orders")
        dialog.parent_path.setText(str(self.directory))
        dialog.next()
        dialog.fields["workspace_mode"].setCurrentIndex(1)
        dialog.fields["goal"].setPlainText("Angebote an Kunden")
        self.assertFalse(dialog.profile_group.isVisible())
        dialog.next()
        self.assertIn("Angebote an Kunden", dialog.summary.toPlainText())
        self.assertFalse((self.directory / "orders").exists())
        dialog.accept()
        self.wait_idle()
        self.assertIsNone(self.window._project_dialog)
        self.assertEqual(self.window.snapshot["workspace_mode"], "mail")
        self.assertEqual(self.window.pages["workspace"].goal.text(), "Angebote an Kunden")

    def test_project_wizard_profile_signature_back_and_cancel(self):
        self.window.create_project()
        dialog = self.window._project_dialog
        dialog.fields["name"].setText("Bewerbungen 2026")
        dialog.fields["folder"].setText("applications")
        dialog.parent_path.setText(str(self.directory))
        dialog.next()
        dialog.add_profile.setChecked(True)
        dialog.fields["display_name"].setText("Ada Beispiel")
        dialog.fields["email"].setText("ada@example.test")
        dialog.fields["phone"].setText("030 12345")
        dialog.add_signature.setChecked(True)
        dialog.fields["signature_name"].setText("Persönlich")
        dialog.fields["signature_text"].setPlainText("Viele Grüße\nAda Beispiel")
        dialog.next()
        self.assertIn("ada@example.test", dialog.summary.toPlainText())
        dialog.back()
        self.assertEqual(dialog.fields["phone"].text(), "030 12345")
        dialog.fields["goal"].setPlainText("Neue Stelle finden")
        dialog.next()
        self.assertIn("Neue Stelle finden", dialog.summary.toPlainText())
        screenshot = os.environ.get("EMAILCENTER_PROJECT_WIZARD_SCREENSHOT_PATH")
        if screenshot:
            self.application.processEvents()
            dialog.grab().save(screenshot)
        dialog.set_busy(True)
        dialog.reject()
        self.assertIs(self.window._project_dialog, dialog)
        dialog.set_busy(False)
        dialog.reject()
        self.assertIsNone(self.window._project_dialog)
        self.assertFalse((self.directory / "applications").exists())
        self.assertIsNone(self.window.snapshot)

    def test_project_wizard_creates_records_and_new_draft_uses_default_signature(self):
        self.window.create_project()
        dialog = self.window._project_dialog
        dialog.fields["name"].setText("Meine Bewerbungen")
        dialog.fields["folder"].setText("setup")
        dialog.parent_path.setText(str(self.directory))
        dialog.next()
        dialog.add_profile.setChecked(True)
        dialog.fields["display_name"].setText("Ada")
        dialog.fields["email"].setText("ada@example.test")
        dialog.add_signature.setChecked(True)
        dialog.fields["signature_name"].setText("Privat")
        dialog.fields["signature_text"].setPlainText("Viele Grüße\nAda")
        dialog.next()
        dialog.accept()
        self.wait_idle()
        self.assertEqual(self.errors, [])
        self.assertEqual(self.window.snapshot["counts"]["applicants"], 1)
        self.submit("list", {"entity": "applicants"})
        self.assertEqual(self.window.pages["applicants"].rows[0]["display_name"], "Ada")
        self.submit("mail_compose_details", {})
        compose = self.window._mail_dialog
        self.assertEqual(compose.signature.currentText(), "Privat")
        compose.signature.setCurrentIndex(0)
        compose.fields["to"].setText("recipient@example.test")
        compose.fields["subject"].setText("Test")
        compose.body.setPlainText("Hallo")
        compose.save()
        self.wait_idle()
        workflow = self.window.pages["workspace"].mail_workflow
        self.submit("mail_compose_details", {"draft_id": workflow.data["drafts"][0]["id"]})
        self.assertIsNone(self.window._mail_dialog.signature.currentData())
        self.window._mail_dialog.reject()

    def test_project_wizard_validation_and_backend_error_preserve_inputs_and_project(self):
        from emailcenter.project import ProjectError
        self.open_test_project()
        old_path = self.window.snapshot["path"]
        self.window.create_project()
        self.wait_idle()
        dialog = self.window._project_dialog
        dialog.fields["name"].setText("Neues Projekt")
        dialog.parent_path.setText(str(self.directory))
        for folder in ("../outside", "CON.txt", "a:b", ".."):
            dialog.fields["folder"].setText(folder)
            dialog.next()
            self.assertEqual(dialog.currentId(), 0)
            self.assertTrue(self.errors.pop())
        dialog.fields["folder"].setText("valid")
        dialog.next()
        dialog.add_profile.setChecked(True)
        dialog.fields["display_name"].setText("Ada")
        dialog.fields["email"].setText("invalid")
        dialog.next()
        self.assertEqual(dialog.currentId(), 1)
        self.assertTrue(self.errors.pop())
        dialog.fields["email"].setText("ada@example.test")
        dialog.next()
        with patch("emailcenter.application.workspace.create_project", side_effect=ProjectError("Ordner nicht beschreibbar")):
            dialog.accept()
            self.wait_idle()
        self.assertIn("nicht beschreibbar", self.errors.pop())
        self.assertIs(self.window._project_dialog, dialog)
        self.assertTrue(dialog.isEnabled())
        self.assertEqual(dialog.fields["email"].text(), "ada@example.test")
        self.assertEqual(self.window.snapshot["path"], old_path)
        dialog.accept()
        self.wait_idle()
        self.assertIsNone(self.window._project_dialog)
        self.assertEqual(self.window.snapshot["name"], "Neues Projekt")

    def test_workspace_reviews_existing_mail_job_without_sending(self):
        self.open_test_project()
        self.enable_fake_gmail()
        draft = self.create_mail_draft()
        self.submit("mail_prepare_job", {"draft_id": draft["id"]})
        self.window._mail_dialog.reject()
        self.submit("set_workspace_mode", {"mode": "mail"})
        page = self.window.pages["workspace"]
        self.assertEqual(page.recent.count(), 1)
        page.open_recent(page.recent.item(0))
        self.wait_idle()
        self.assertEqual(self.window._mail_dialog.job["status"], "Draft")
        self.assertIs(self.window.current_page(), page)
        self.assertEqual(self.fake_provider.sent, [])
        self.window._mail_dialog.reject()

    def action_result(self, action, payload):
        values = []
        def collect(completed, result):
            if completed == action:
                values.append(result)
        self.window.controller.completed.connect(collect)
        try:
            self.submit(action, payload)
        finally:
            self.window.controller.completed.disconnect(collect)
        self.assertTrue(values)
        return values[0]

    def context_fixture(self, *, documents=False):
        company = self.action_result("create_record", {"entity": "companies", "fields": {"name": "Example GmbH"}})
        profile = self.action_result("create_record", {"entity": "applicants", "fields": {"display_name": "Ada", "email": "ada@example.test"}})
        job = self.action_result("create_record", {"entity": "jobs", "fields": {"title": "Entwicklerin", "reference_number": "DEV-42", "company_id": company["id"]}})
        application = self.action_result("create_record", {"entity": "applications", "fields": {
            "title": "Bewerbung Entwicklerin", "company_id": company["id"],
            "applicant_profile_id": profile["id"], "job_posting_id": job["id"]}})
        contacts = []
        for name in ("Grace", "Linus"):
            contact = self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": name, "email": name.lower()+"@example.test"}})
            contacts.append(contact)
            links = self.action_result("record_links", {"entity": "applications", "record_id": application["id"]})
            self.submit("save_record_link", {"entity": "applications", "record_id": application["id"],
                        "contact_id": contact["id"], "expected_version": links["version"]})
            self.window._record_dialog.reject()
        other = self.action_result("create_record", {"entity": "applications", "fields": {"title": "Andere Bewerbung", "company_id": company["id"]}})
        generated = None
        if documents:
            source = self.directory / "context-template.docx"
            _make_docx(source)
            imported = self.action_result("import_template", {"source_path": str(source), "name": "Anschreiben",
                "mapping": {"NAME": "applicant.display_name", "COMPANY": "company.name"}, "required_placeholders": ["NAME"]})
            for record, filename in ((application, "Anschreiben.docx"), (other, "Andere.docx")):
                result = self.action_result("generate_document", {"template_version_id": imported["version"]["id"],
                    "values": {"NAME": "Ada", "COMPANY": "Example GmbH"}, "output_filename": filename,
                    "application_id": record["id"]})
                if record is application:
                    generated = result
        return application, contacts, other, generated

    def test_workspace_context_selects_recipient_and_copies_only_chosen_document(self):
        self.open_test_project()
        application, contacts, _, generated = self.context_fixture(documents=True)
        page = self.window.pages["workspace"]
        self.assertFalse(page.context_compose.isEnabled())
        page.context_record.setCurrentIndex(page.context_record.findData(application["id"]))
        page.context_compose.click()
        self.wait_idle()
        dialog = self.window._mail_dialog
        self.assertIs(self.window.current_page(), page)
        self.assertEqual(dialog.application.currentData(), application["id"])
        self.assertEqual(dialog.fields["to"].text(), "")
        self.assertIn("Mehrere Ansprechpartner", dialog.context_info.text())
        self.assertIn("DEV-42", dialog.fields["subject"].text())
        self.assertEqual(dialog.attachments, [])
        self.assertEqual(dialog.documents.count(), 2)
        self.assertEqual(dialog.documents.itemData(1)["id"], generated["document"]["id"])
        dialog.contact_choice.setCurrentIndex(2)
        dialog.use_contact.click()
        self.assertEqual(dialog.fields["to"].text(), "linus@example.test")
        self.assertIn("Guten Tag Linus", dialog.body.toPlainText())
        dialog.documents.setCurrentIndex(1)
        dialog.add_document()
        dialog.add_document()
        self.assertEqual(len(dialog.attachments), 1)
        dialog.fields["subject"].setText("Meine angepasste Bewerbung")
        dialog.body.setPlainText("Guten Tag Linus,\nmein fertiger Nachrichtentext.")
        dialog.resize(600, 550)
        self.application.processEvents()
        self.assertLessEqual(dialog.width(), 600)
        from PySide6.QtWidgets import QScrollArea
        self.assertGreater(dialog.findChild(QScrollArea).verticalScrollBar().maximum(), 0)
        screenshot = os.environ.get("EMAILCENTER_CONTEXT_COMPOSE_SCREENSHOT_PATH")
        if screenshot:
            dialog.resize(780, 760)
            self.application.processEvents()
            dialog.grab().save(screenshot)
        original = Path(generated["path"]).read_bytes()
        dialog.save()
        self.wait_idle()
        draft = page.mail_workflow.data["drafts"][0]
        self.assertEqual(draft["application_id"], application["id"])
        self.assertEqual(draft["subject"], "Meine angepasste Bewerbung")
        captured = Path(self.window.snapshot["path"]) / draft["configuration"]["attachments"][0]["path"]
        self.assertEqual(captured.read_bytes(), original)
        Path(generated["path"]).write_bytes(b"source modified later")
        self.assertEqual(captured.read_bytes(), original)
        self.enable_fake_gmail()
        page.mail_workflow.buttons["prepare"].click()
        self.window._mail_dialog.prepare()
        self.wait_idle()
        self.assertEqual(self.window._mail_dialog.job["status"], "Draft")
        self.assertEqual(self.fake_provider.sent, [])
        self.window._mail_dialog.reject()

    def test_contact_record_action_opens_context_and_cancel_leaves_no_draft(self):
        self.open_test_project()
        contact = self.action_result("create_record", {"entity": "contacts", "fields": {
            "display_name": "Grace", "email": "grace@example.test"}})
        self.window.navigate("contacts")
        self.wait_idle()
        page = self.window.pages["contacts"]
        self.assertFalse(page.action_buttons["compose_context"][0].isEnabled())
        page.table.selectRow(0)
        page.action_buttons["compose_context"][0].click()
        self.wait_idle()
        dialog = self.window._mail_dialog
        self.assertIs(self.window.current_page(), page)
        self.assertEqual(dialog.fields["to"].text(), "grace@example.test")
        self.assertIn("Grace", dialog.fields["subject"].text())
        self.assertIsNone(dialog.application.currentData())
        self.assertEqual(dialog.context["selected_contact_id"], contact["id"])
        dialog.reject()
        self.submit("snapshot", {})
        self.assertEqual(self.window.snapshot["counts"]["mail_drafts"], 0)
        self.window.resize(900, 650)
        self.application.processEvents()
        self.assertEqual(page.actions_row.getItemPosition(2)[0], 1)

    def test_context_overwrite_requires_confirmation_and_errors_keep_existing_draft(self):
        self.open_test_project()
        application, _, _, _ = self.context_fixture()
        self.submit("mail_compose_details", {})
        dialog = self.window._mail_dialog
        dialog.fields["to"].setText("manual@example.test")
        dialog.fields["subject"].setText("Manueller Betreff")
        dialog.body.setPlainText("Manueller Text")
        dialog.save()
        self.wait_idle()
        draft = self.window.pages["workspace"].mail_workflow.data["drafts"][0]
        self.submit("mail_compose_details", {"draft_id": draft["id"]})
        dialog = self.window._mail_dialog
        dialog.application.setCurrentIndex(dialog.application.findData(application["id"]))
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            dialog.load_context.click()
        self.assertEqual(dialog.fields["subject"].text(), "Manueller Betreff")
        with patch("emailcenter.application.workspace.compose_context", side_effect=ValueError("Kontext nicht verfügbar")), patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.load_context.click()
            self.wait_idle()
        self.assertIn("nicht verfügbar", self.errors.pop())
        self.assertIs(self.window._mail_dialog, dialog)
        self.assertTrue(dialog.isEnabled())
        self.assertEqual(dialog.body.toPlainText(), "Manueller Text")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.load_context.click()
            self.wait_idle()
        self.assertIs(self.window._mail_dialog, dialog)
        self.assertEqual(dialog.draft_id, draft["id"])
        self.assertIn("DEV-42", dialog.fields["subject"].text())
        self.assertEqual(dialog.fields["to"].text(), "")
        dialog.contact_choice.setCurrentIndex(1)
        dialog.use_contact.click()
        dialog.save()
        self.wait_idle()
        workflow = self.window.pages["workspace"].mail_workflow
        self.assertEqual(workflow.data["draft_total"], 1)
        self.assertEqual(workflow.data["drafts"][0]["id"], draft["id"])

    def test_missing_contact_email_can_be_completed_manually_and_stale_selection_fails_cleanly(self):
        self.open_test_project()
        contact = self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": "Grace"}})
        page = self.window.pages["workspace"]
        page.context_kind.setCurrentIndex(1)
        page.context_record.setCurrentIndex(page.context_record.findData(contact["id"]))
        page.context_compose.click()
        self.wait_idle()
        dialog = self.window._mail_dialog
        self.assertEqual(dialog.fields["to"].text(), "")
        self.assertIn("E-Mail fehlt", dialog.context_info.text())
        dialog.fields["to"].setText("grace@example.test")
        dialog.body.setPlainText("Hallo Grace,\nmeine Nachricht.")
        dialog.save()
        self.wait_idle()
        self.assertEqual(page.mail_workflow.data["drafts"][0]["recipients"]["to"], ["grace@example.test"])
        self.window.workspace_action("compose_context", {"contact_id": "deleted-record"})
        self.wait_idle()
        self.assertIn("nicht gefunden", self.errors.pop())
        self.assertIsNone(self.window._mail_dialog)
        self.assertEqual(page.context_record.currentData(), contact["id"])

    def test_context_contact_choice_preserves_manual_text_and_changed_application_disables_old_proposals(self):
        self.open_test_project()
        application, _, other, _ = self.context_fixture()
        self.window.navigate("applications")
        self.wait_idle()
        page = self.window.pages["applications"]
        page.table.selectRow(next(i for i, row in enumerate(page.rows) if row["id"] == application["id"]))
        page.action_buttons["compose_context"][0].click()
        self.wait_idle()
        dialog = self.window._mail_dialog
        self.assertIs(self.window.current_page(), page)
        dialog.body.setPlainText("Mein eigener Text")
        dialog.contact_choice.setCurrentIndex(1)
        dialog.use_contact.click()
        self.assertEqual(dialog.body.toPlainText(), "Mein eigener Text")
        dialog.application.setCurrentIndex(dialog.application.findData(other["id"]))
        self.assertFalse(dialog.documents.isEnabled())
        self.assertFalse(dialog.context_contacts.isEnabled())
        self.assertIn("Bewerbung geändert", dialog.context_info.text())
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.load_context.click()
            self.wait_idle()
        self.assertEqual(dialog.context["application_id"], other["id"])
        self.assertEqual(dialog.fields["to"].text(), "")
        self.assertEqual(dialog.documents.count(), 1)
        self.assertEqual(dialog.attachments, [])
        dialog.reject()

    def test_workspace_mode_error_restores_selector_and_due_toast_is_not_repeated(self):
        self.open_test_project()
        page = self.window.pages["workspace"]
        page.mode.setCurrentIndex(1)
        self.assertTrue(self.window.controller.submit("set_workspace_mode", {"mode": "invalid"}))
        self.wait_idle()
        self.assertEqual(len(self.errors), 1)
        self.errors.clear()
        self.assertEqual(page.mode.currentData(), "applications")
        self.window.toast.dismiss()
        self.window.apply_followup_summary({"due": 1})
        self.assertTrue(self.window.toast.isVisible())
        self.window.toast.dismiss()
        self.window.apply_followup_summary({"due": 1})
        self.assertFalse(self.window.toast.isVisible())

    def test_workspace_single_mail_flow_edit_prepare_approve_and_send(self):
        self.open_test_project()
        self.enable_fake_gmail()
        self.submit("set_workspace_mode", {"mode": "mail"})
        page = self.window.pages["workspace"]
        workflow = page.mail_workflow
        workflow.new_button.click()
        self.wait_idle()
        dialog = self.window._mail_dialog
        dialog.fields["to"].setText("target@example.test")
        dialog.fields["subject"].setText("Einzelauftrag")
        dialog.body.setPlainText("Guten Tag,\nmeine Nachricht.")
        dialog.save()
        self.wait_idle()
        self.assertIs(self.window.current_page(), page)
        self.assertEqual(workflow.selected()["subject"], "Einzelauftrag")
        self.assertIn("Senden / planen", workflow.details.text())
        workflow.buttons["edit"].click()
        self.wait_idle()
        dialog = self.window._mail_dialog
        dialog.fields["subject"].setText("Geprüfter Einzelauftrag")
        dialog.save()
        self.wait_idle()
        workflow.buttons["prepare"].click()
        dialog = self.window._mail_dialog
        self.assertEqual(dialog.zone.currentText(), self.window.preferences.timezone_id)
        dialog.prepare()
        self.wait_idle()
        self.assertEqual(workflow.kind, "jobs")
        self.assertEqual(workflow.selected()["status"], "Draft")
        self.assertFalse(workflow.buttons["send"].isEnabled())
        self.assertEqual(self.fake_provider.sent, [])
        review = self.window._mail_dialog
        self.assertEqual(review.job["subject_snapshot"], "Geprüfter Einzelauftrag")
        review.confirm.setChecked(True)
        review.approve()
        self.wait_idle()
        self.assertTrue(workflow.buttons["send"].isEnabled())
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            workflow.buttons["send"].click()
        self.assertEqual(self.fake_provider.sent, [])
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            workflow.buttons["send"].click()
        self.wait_idle()
        self.assertEqual(len(self.fake_provider.sent), 1)
        self.assertEqual(workflow.selected()["status"], "Sent")
        self.assertFalse(workflow.buttons["send"].isEnabled())
        self.assertIn("Versand bestätigt", workflow.details.text())
        self.assertIs(self.window.current_page(), page)
        screenshot_path = os.environ.get("EMAILCENTER_MAIL_WORKSPACE_SCREENSHOT_PATH")
        if screenshot_path:
            self.window.toast.dismiss()
            self.window.grab().save(screenshot_path)

    def test_workspace_preparation_errors_preserve_dialog_and_prevent_busy_close(self):
        self.open_test_project()
        self.enable_fake_gmail()
        self.create_mail_draft()
        workflow = self.window.pages["workspace"].mail_workflow
        workflow.buttons["prepare"].click()
        dialog = self.window._mail_dialog
        from PySide6.QtWidgets import QDialogButtonBox
        send_button = dialog.buttons.button(QDialogButtonBox.StandardButton.Save)
        self.assertEqual(send_button.text(), "Senden")
        dialog.zone.setCurrentText("Unknown/Zone")
        dialog.prepare()
        self.wait_idle()
        self.assertEqual(len(self.errors), 1)
        self.errors.clear()
        self.assertIs(self.window._mail_dialog, dialog)
        self.assertEqual(dialog.zone.currentText(), "Unknown/Zone")
        self.assertTrue(dialog.isEnabled())
        self.assertEqual(workflow.data["job_total"], 0)
        dialog.set_busy(True)
        dialog.reject()
        self.assertIs(self.window._mail_dialog, dialog)
        dialog.set_busy(False)
        dialog.zone.setCurrentText("UTC")
        dialog.scheduled.setChecked(True)
        self.assertEqual(send_button.text(), "Versand planen")
        from PySide6.QtCore import QDate, QDateTime, QTime, QTimeZone
        dialog.when.setDateTime(QDateTime(QDate(2020, 1, 1), QTime(9, 0), QTimeZone(b"UTC")))
        dialog.prepare()
        self.wait_idle()
        self.assertIn("zukünftigen", self.errors.pop())
        self.assertIs(self.window._mail_dialog, dialog)
        dialog.zone.setCurrentText("Europe/Berlin")
        dialog.when.setDateTime(QDateTime(QDate(2030, 3, 31), QTime(2, 30), QTimeZone(b"UTC")))
        dialog.prepare()
        self.wait_idle()
        self.assertIn("existiert", self.errors.pop())
        self.assertEqual(dialog.when.dateTime().toString("HH:mm"), "02:30")
        self.assertEqual(workflow.data["job_total"], 0)
        dialog.zone.setCurrentText("UTC")
        dialog.scheduled.setChecked(False)
        self.assertEqual(send_button.text(), "Senden")
        dialog.prepare()
        self.wait_idle()
        self.assertEqual(self.window._mail_dialog.job["timezone_id"], "UTC")
        self.window._mail_dialog.reject()

    def test_workspace_mail_pagination_edit_older_draft_and_reset(self):
        self.open_test_project()
        self.enable_fake_gmail()
        workflow = self.window.pages["workspace"].mail_workflow
        workflow.limit = 2
        account_id = self.window.pages["mail"].data["accounts"][0]["id"]
        for index in range(3):
            self.submit("mail_save_draft", {"subject": f"Nachricht {index}", "body": "Text",
                "recipients_value": {"to": ["target@example.test"]}, "provider_account_id": account_id})
        self.assertEqual(workflow.records.count(), 2)
        self.assertTrue(workflow.next.isEnabled())
        workflow.next.click()
        self.wait_idle()
        identifier = workflow.selected()["id"]
        self.assertEqual(workflow.records.count(), 1)
        self.assertIn("3–3 von 3", workflow.summary.text())
        self.assertFalse(workflow.next.isEnabled())
        workflow.buttons["edit"].click()
        self.wait_idle()
        self.assertEqual(self.window._mail_dialog.draft_id, identifier)
        self.window._mail_dialog.reject()
        workflow.request_load()
        self.wait_idle()
        self.assertEqual(workflow.selected()["id"], identifier)
        workflow.previous.click()
        self.wait_idle()
        self.assertEqual(workflow.offsets["drafts"], 0)
        self.submit("close_project", {})
        self.assertIsNone(workflow.data)
        self.assertEqual(workflow.selection, {"drafts": None, "jobs": None})

    def test_workspace_uncertain_job_only_allows_reconciliation(self):
        from emailcenter.gmail import ProviderError
        self.open_test_project()
        self.enable_fake_gmail()
        draft = self.create_mail_draft()
        self.submit("mail_prepare_job", {"draft_id": draft["id"]})
        review = self.window._mail_dialog
        review.confirm.setChecked(True)
        review.approve()
        self.wait_idle()
        workflow = self.window.pages["workspace"].mail_workflow
        self.fake_provider.failure = ProviderError("timeout", uncertain=True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            workflow.buttons["send"].click()
        self.wait_idle()
        self.assertEqual(workflow.selected()["status"], "Uncertain")
        self.assertFalse(workflow.buttons["send"].isEnabled())
        self.assertFalse(workflow.buttons["cancel"].isEnabled())
        self.assertTrue(workflow.buttons["reconcile"].isEnabled())
        self.window.workspace_mail_action("send", workflow.selected())
        self.assertEqual(len(self.fake_provider.sent), 1)
        self.fake_provider.matches = [{"id": "gmail123", "message_id": workflow.selected()["message_id_snapshot"]}]
        workflow.buttons["reconcile"].click()
        self.wait_idle()
        self.assertEqual(workflow.selected()["status"], "Sent")
        self.assertEqual(len(self.fake_provider.sent), 1)

    def test_workspace_calendar_schedule_and_cancel_remain_explicit(self):
        from PySide6.QtCore import QDate, QDateTime, QTime, QTimeZone
        self.open_test_project()
        self.enable_fake_gmail()
        self.create_mail_draft()
        workflow = self.window.pages["workspace"].mail_workflow
        workflow.buttons["prepare"].click()
        dialog = self.window._mail_dialog
        dialog.scheduled.setChecked(True)
        dialog.zone.setCurrentText("Europe/Berlin")
        dialog.when.setDateTime(QDateTime(QDate(2030, 1, 20), QTime(9, 30), QTimeZone(b"UTC")))
        dialog.prepare()
        self.wait_idle()
        review = self.window._mail_dialog
        self.assertEqual(review.job["scheduled_at_utc"], "2030-01-20T08:30:00Z")
        review.confirm.setChecked(True)
        review.approve()
        self.wait_idle()
        self.assertEqual(workflow.selected()["status"], "Scheduled")
        self.assertIn("Scheduler", workflow.details.text())
        self.assertEqual(self.fake_provider.sent, [])
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            workflow.buttons["cancel"].click()
        self.assertEqual(workflow.selected()["status"], "Scheduled")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            workflow.buttons["cancel"].click()
        self.wait_idle()
        self.assertEqual(workflow.selected()["status"], "Cancelled")
        self.assertFalse(workflow.buttons["send"].isEnabled())
        self.assertEqual(self.fake_provider.sent, [])

    def test_draft_before_login_prepares_and_sends_immediately_with_selected_account(self):
        from PySide6.QtCore import QDate, QDateTime, QTime, QTimeZone
        self.open_test_project()
        self.submit("mail_save_draft", {"subject": "Vor der Anmeldung", "body": "Mein Text",
            "recipients_value": {"to": ["target@example.test"]}})
        self.enable_fake_gmail()
        workflow = self.window.pages["workspace"].mail_workflow
        self.assertIn("sender@example.test", workflow.account_status.text())
        self.assertIsNone(workflow.selected()["provider_account_id"])
        workflow.buttons["prepare"].click()
        prepare = self.window._mail_dialog
        self.assertEqual(prepare.account.currentData(), workflow.data["accounts"][0]["id"])
        prepare.scheduled.setChecked(True)
        prepare.when.setDateTime(QDateTime(QDate(2030, 1, 20), QTime(9, 30), QTimeZone(b"UTC")))
        prepare.prepare()
        self.wait_idle()
        review = self.window._mail_dialog
        self.assertFalse(review.send_now_button.isEnabled())
        review.confirm.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            review.send_now_button.click()
        self.assertEqual(self.fake_provider.sent, [])
        self.assertEqual(workflow.selected()["status"], "Draft")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            review.send_now_button.click()
        self.wait_idle()
        self.assertIsNone(self.window._mail_dialog)
        self.assertEqual(workflow.selected()["status"], "Sent")
        self.assertEqual(len(self.fake_provider.sent), 1)
        self.assertIs(self.window.current_page(), self.window.pages["workspace"])

    def test_old_unassigned_draft_editor_selects_only_connected_account(self):
        self.open_test_project()
        self.submit("mail_save_draft", {"subject": "Vor Login", "body": "Unverändert",
            "recipients_value": {"to": ["target@example.test"]}})
        self.enable_fake_gmail()
        workflow = self.window.pages["workspace"].mail_workflow
        workflow.buttons["edit"].click()
        self.wait_idle()
        editor = self.window._mail_dialog
        self.assertIsNotNone(editor.account.currentData())
        self.assertEqual(editor.body.toPlainText(), "Unverändert")
        editor.save()
        self.wait_idle()
        self.assertIsNotNone(workflow.selected()["provider_account_id"])

    def test_workspace_connect_and_scheduler_are_available_without_navigation(self):
        from PySide6.QtWidgets import QFileDialog
        self.open_test_project()
        self.enable_fake_gmail()
        workflow = self.window.pages["workspace"].mail_workflow
        with patch.object(QFileDialog, "getOpenFileName", return_value=("fake-only-no-real-file", "")):
            workflow.connect_button.click()
        self.wait_idle()
        self.assertIn("sender@example.test", workflow.account_status.text())
        workflow.scheduler.setChecked(True)
        self.assertTrue(self.window.pages["mail"].scheduler.isChecked())
        self.window.pages["mail"].scheduler.setChecked(False)
        self.assertFalse(workflow.scheduler.isChecked())
        self.assertIs(self.window.current_page(), self.window.pages["workspace"])
        self.submit("close_project", {})
        self.assertFalse(workflow.scheduler.isChecked())

    def test_multiple_connected_accounts_require_sender_selection_for_old_draft(self):
        self.open_test_project()
        self.submit("mail_save_draft", {"subject": "Absender wählen", "body": "Text",
            "recipients_value": {"to": ["target@example.test"]}})
        self.enable_fake_gmail()
        with patch("emailcenter.application.workspace.MailService.connect_gmail", new=lambda service, client_path: service.register_account("second@example.test", {"refresh_token": "fake-second"})):
            self.submit("mail_connect", {"client_path": "fake-only"})
        workflow = self.window.pages["workspace"].mail_workflow
        workflow.buttons["prepare"].click()
        prepare = self.window._mail_dialog
        self.assertIsNone(prepare.account.currentData())
        prepare.prepare()
        self.wait_idle()
        self.assertIn("Gmail-Konto", self.errors.pop())
        self.assertIs(self.window._mail_dialog, prepare)
        self.assertEqual(workflow.data["job_total"], 0)
        prepare.account.setCurrentIndex(1)
        prepare.prepare()
        self.wait_idle()
        self.assertIsNotNone(self.window._mail_dialog.job["provider_account_id"])
        self.window._mail_dialog.reject()

    def test_immediate_send_failure_refreshes_workspace_and_does_not_leave_approval(self):
        self.open_test_project()
        self.enable_fake_gmail()
        draft = self.create_mail_draft()
        self.submit("mail_prepare_job", {"draft_id": draft["id"]})
        review = self.window._mail_dialog
        review.confirm.setChecked(True)
        self.fake_provider.email = "different@example.test"
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            review.send_now_button.click()
        self.wait_idle()
        self.assertIn("Absender", self.errors.pop())
        self.assertIsNone(self.window._mail_dialog)
        workflow = self.window.pages["workspace"].mail_workflow
        self.assertEqual(workflow.selected()["approval_status"], "pending")
        self.assertFalse(workflow.buttons["send"].isEnabled())
        self.assertEqual(self.fake_provider.sent, [])

    def test_mail_page_prepares_pre_login_draft_with_connected_sender(self):
        from emailcenter.ui.dialogs import FormDialog
        from PySide6.QtWidgets import QDialog
        self.open_test_project()
        self.submit("mail_save_draft", {"subject": "Vor Login", "body": "Text",
            "recipients_value": {"to": ["target@example.test"]}})
        self.enable_fake_gmail()
        self.window.navigate("mail")
        self.wait_idle()
        page = self.window.pages["mail"]
        page.tables["drafts"].selectRow(0)
        def accept(dialog):
            self.assertIsNotNone(dialog.fields["provider_account_id"].currentData())
            return QDialog.DialogCode.Accepted
        with patch.object(FormDialog, "exec", new=accept):
            page.action("prepare")
        self.wait_idle()
        self.assertEqual(self.window._mail_dialog.job["sender_snapshot"], "sender@example.test")
        self.window._mail_dialog.reject()

    def test_settings_save_and_reload_on_new_window_without_project(self):
        self.window.navigate("settings")
        page = self.window.pages["settings"]
        self.assertTrue(page.isEnabled())
        self.assertTrue(self.window.actions["settings"].isEnabled())
        page.theme.setCurrentIndex(page.theme.findData("dark"))
        page.accent.setCurrentIndex(page.accent.findData("#087f5b"))
        page.zone.setCurrentText("UTC")
        self.assertEqual(page.state.text(), "Ungespeicherte Änderungen")
        page.save_button.click()
        self.assertTrue(self.window.dark)
        self.assertEqual(self.window.preferences, Preferences("dark", "#087f5b", "UTC"))
        self.assertIn("#087f5b", self.application.styleSheet())
        secondary = MainWindow(settings_store=PreferencesStore(self.directory / "preferences.json"))
        try:
            secondary.show()
            self.application.processEvents()
            self.assertEqual(secondary.preferences, self.window.preferences)
            self.assertTrue(secondary.dark)
            self.assertEqual(secondary.actions["theme"].text(), "Helles Design")
            secondary.navigate("settings")
            self.assertEqual(secondary.pages["settings"].values(), page.values())
        finally:
            secondary.close()
            self.wait_for(lambda: not secondary.controller.thread.isRunning())
            secondary.deleteLater()
            self.application.sendPostedEvents(None, QEvent.Type.DeferredDelete)
            self.application.processEvents()
        self.application.processEvents()
        screenshot_path = os.environ.get("EMAILCENTER_SETTINGS_SCREENSHOT_PATH")
        if screenshot_path:
            self.window.grab().save(screenshot_path)
        self.assertEqual(self.errors, [])

    def test_settings_save_failure_preserves_saved_preferences_and_form(self):
        self.open_test_project()
        self.window.navigate("settings")
        page = self.window.pages["settings"]
        page.theme.setCurrentIndex(page.theme.findData("dark"))
        page.zone.setCurrentText("UTC")
        original = self.window.preferences
        with patch.object(self.window.settings_store, "save", side_effect=PermissionError("Write denied")):
            page.save_button.click()
        self.assertIn("Write denied", self.errors.pop())
        self.assertEqual(self.window.preferences, original)
        self.assertFalse(self.window.dark)
        self.assertEqual(page.zone.currentText(), "UTC")
        self.assertEqual(page.theme.currentData(), "dark")
        self.assertEqual(page.state.text(), "Ungespeicherte Änderungen")
        page.zone.setCurrentText("Unknown/Place")
        page.save_button.click()
        self.assertIn("Zeitzone", self.errors.pop())
        self.assertEqual(self.window.preferences, original)

    def test_settings_toolbar_theme_retains_unsaved_color_and_timezone(self):
        self.window.navigate("settings")
        page = self.window.pages["settings"]
        page.accent.setCurrentIndex(page.accent.findData("#7045b8"))
        page.zone.setCurrentText("Europe/London")
        self.window.toggle_theme()
        self.assertTrue(self.window.dark)
        self.assertEqual(page.theme.currentData(), "dark")
        self.assertEqual(page.accent.currentData(), "#7045b8")
        self.assertEqual(page.zone.currentText(), "Europe/London")
        persisted, _ = PreferencesStore(self.directory / "preferences.json").load()
        self.assertEqual(persisted, Preferences(theme="dark"))
        page.save_button.click()
        persisted, _ = PreferencesStore(self.directory / "preferences.json").load()
        self.assertEqual(persisted, Preferences("dark", "#7045b8", "Europe/London"))
        self.assertEqual(self.errors, [])

    def test_settings_defaults_and_discard_are_not_saved_until_applied(self):
        self.window.navigate("settings")
        page = self.window.pages["settings"]
        page.theme.setCurrentIndex(page.theme.findData("dark"))
        page.accent.setCurrentIndex(page.accent.findData("#b45309"))
        page.save_button.click()
        original = self.window.preferences
        original_file = self.window.settings_store.path.read_bytes()
        page.defaults_button.click()
        self.assertEqual(page.values(), Preferences().values())
        self.assertEqual(self.window.preferences, original)
        self.assertEqual(self.window.settings_store.path.read_bytes(), original_file)
        page.discard_button.click()
        self.assertEqual(page.values(), original.values())
        page.defaults_button.click()
        page.save_button.click()
        self.assertEqual(self.window.preferences, Preferences())

    def test_settings_timezone_updates_followups_and_new_reminder_defaults(self):
        self.open_test_project()
        self.submit("create_record", {"entity": "applications", "fields": {"title": "Engineer", "next_follow_up_at": "2000-01-01T09:00:00Z"}})
        self.window.navigate("settings")
        settings = self.window.pages["settings"]
        settings.zone.setCurrentText("Asia/Tokyo")
        settings.save_button.click()
        self.window.navigate("followups")
        self.wait_idle()
        page = self.window.pages["followups"]
        self.assertEqual(page.rows[0]["due_label"], "01.01.2000 18:00 JST")
        self.assertEqual(page.rows[0]["due_at_utc"], "2000-01-01T09:00:00Z")
        page.table.selectRow(0)
        page.action_buttons["edit_reminder"][0].click()
        self.wait_idle()
        dialog = self.window._record_dialog
        self.assertEqual(dialog.zone.currentText(), "Asia/Tokyo")
        self.assertEqual(dialog.time.time().hour(), 18)
        dialog.reject()
        page.action_buttons["new_reminder"][0].click()
        self.wait_idle()
        self.assertEqual(self.window._record_dialog.zone.currentText(), "Asia/Tokyo")
        self.window._record_dialog.reject()
        self.assertEqual(self.errors, [])

    def test_settings_custom_colors_are_validated_and_use_readable_text(self):
        from PySide6.QtGui import QColor, QPalette
        from emailcenter.ui.theme import accent_text_color
        self.window.navigate("settings")
        page = self.window.pages["settings"]
        with patch("emailcenter.ui.settings_page.QColorDialog.getColor", return_value=QColor("#fde047")):
            page.choose_color()
        self.assertEqual(page.accent.currentData(), "#fde047")
        page.save_button.click()
        self.assertEqual(accent_text_color("#fde047"), "#000000")
        self.assertEqual(self.application.palette().color(QPalette.ColorRole.HighlightedText).name(), "#000000")
        self.assertEqual(accent_text_color("#087f5b"), "#ffffff")
        with patch("emailcenter.ui.settings_page.QColorDialog.getColor", return_value=QColor()):
            page.choose_color()
        self.assertEqual(page.accent.currentData(), "#fde047")
        with self.assertRaises(ValueError):
            apply_theme(self.application, accent="red; color:white")
        self.assertEqual(self.errors, [])

    def test_settings_reload_reports_damaged_file_without_automatic_overwrite(self):
        self.window.navigate("settings")
        page = self.window.pages["settings"]
        content = b'{"theme": '
        self.window.settings_store.path.write_bytes(content)
        page.reload_button.click()
        self.assertTrue(page.notice.isVisibleTo(page))
        self.assertIn("geladen", page.notice.text())
        self.assertEqual(self.window.settings_store.path.read_bytes(), content)
        self.assertEqual(self.window.preferences, Preferences())
        page.save_button.click()
        self.assertEqual(PreferencesStore(self.directory / "preferences.json").load(), (Preferences(), None))
        self.assertFalse(page.notice.isVisibleTo(page))
        self.assertEqual(self.errors, [])

    def test_create_project_add_records_load_pages_and_close(self):
        self.open_test_project()
        self.assertEqual(self.window.snapshot["name"], "Meine Bewerbungen")
        self.submit("create_record", {"entity": "companies", "fields": {"name": "Beispiel GmbH"}})
        self.submit("create_record", {"entity": "contacts", "fields": {"display_name": "Anna Beispiel", "email": "anna@example.test"}})
        self.submit("create_record", {"entity": "applications", "fields": {"title": "Softwareentwickler"}})
        self.assertEqual(self.window.snapshot["counts"]["companies"], 1)
        self.window.navigate("applications")
        self.wait_idle()
        page = self.window.pages["applications"]
        self.assertEqual(page.table.item(0, 0).text(), "Softwareentwickler")
        page.table.selectRow(0)
        self.assertTrue(page.action_buttons["status"][0].isEnabled())
        self.submit("change_status", {"application_id": page.rows[0]["id"], "status_code": "accepted"})
        page.table.selectRow(0)
        self.assertFalse(page.action_buttons["status"][0].isEnabled())
        self.window.navigate("overview")
        self.wait_idle()
        screenshot_path = os.environ.get("EMAILCENTER_SCREENSHOT_PATH")
        if screenshot_path:
            self.window.grab().save(screenshot_path)
        self.submit("close_project", {})
        self.assertIsNone(self.window.snapshot)
        self.assertEqual(page.table.rowCount(), 0)

    def test_list_search_and_pagination_controls(self):
        self.open_test_project()
        self.window.navigate("contacts")
        self.wait_idle()
        page = self.window.pages["contacts"]
        page.display({"rows": [{"id": "1", "display_name": "Ada"}], "offset": 50, "total": 101})
        self.assertTrue(page.previous.isEnabled())
        self.assertTrue(page.next.isEnabled())
        requests = []
        page.load_requested.disconnect()
        page.load_requested.connect(lambda action, payload: requests.append(payload))
        page.search.setText("Ada")
        page.search_records()
        self.assertEqual(requests[-1]["offset"], 0)
        self.assertEqual(requests[-1]["search"], "Ada")

    def test_followup_manual_reminder_creation_completion_and_history(self):
        from PySide6.QtCore import QDate, QTime
        self.open_test_project()
        self.window.navigate("followups")
        self.wait_idle()
        page = self.window.pages["followups"]
        page.action_buttons["new_reminder"][0].click()
        self.wait_idle()
        dialog = self.window._record_dialog
        dialog.title.setText("Personal nachfragen")
        dialog.description.setPlainText("Unterlagen bereithalten")
        dialog.date.setDate(QDate.currentDate().addDays(-1))
        dialog.time.setTime(QTime(9, 0))
        dialog.save()
        self.wait_idle()
        self.assertIsNone(self.window._record_dialog)
        self.assertEqual(page.total, 1)
        self.assertEqual(page.rows[0]["status_label"], "Überfällig")
        self.assertEqual(self.window.pages["overview"].metrics["followups"].text(), "1")
        page.table.selectRow(0)
        page.action_buttons["finish_reminder"][0].click()
        self.wait_idle()
        self.assertEqual(page.total, 0)
        self.assertEqual(self.window.snapshot["counts"]["followups"], 0)
        page.view.setCurrentIndex(page.view.findData("history"))
        self.wait_idle()
        self.assertEqual(page.rows[0]["status_label"], "Erledigt")
        page.table.selectRow(0)
        self.assertFalse(page.action_buttons["edit_reminder"][0].isEnabled())
        self.assertFalse(page.action_buttons["finish_reminder"][0].isEnabled())
        self.assertEqual(self.errors, [])

    def test_followup_application_calendar_reschedule_and_open_application(self):
        from PySide6.QtCore import QDate
        self.open_test_project()
        self.submit("create_record", {"entity": "applications", "fields": {
            "title": "Softwareentwickler", "notes": "Bewerbung verschickt", "next_follow_up_at": "2000-01-01T09:00:00Z"}})
        self.window.navigate("followups")
        self.wait_idle()
        page = self.window.pages["followups"]
        self.assertEqual(page.rows[0]["source"], "application")
        page.table.selectRow(0)
        page.action_buttons["edit_reminder"][0].click()
        self.wait_idle()
        dialog = self.window._record_dialog
        self.assertTrue(dialog.title.isReadOnly())
        self.assertTrue(dialog.description.isReadOnly())
        self.assertFalse(dialog.application.isEnabled())
        dialog.date.setDate(QDate.currentDate().addDays(2))
        dialog.set_busy(True)
        dialog.reject()
        self.assertIs(self.window._record_dialog, dialog)
        self.assertFalse(dialog.buttons.isEnabled())
        dialog.set_busy(False)
        dialog.save()
        self.wait_idle()
        self.assertEqual(page.total, 0)
        page.view.setCurrentIndex(page.view.findData("upcoming"))
        self.wait_idle()
        self.assertEqual(page.total, 1)
        self.assertEqual(page.rows[0]["description"], "Bewerbung verschickt")
        page.table.selectRow(0)
        page.action_buttons["followup_application"][0].click()
        self.wait_idle()
        self.assertEqual(self.window._record_dialog.fields["title"].text(), "Softwareentwickler")
        self.window._record_dialog.reject()
        page.table.selectRow(0)
        page.action_buttons["finish_reminder"][0].click()
        self.wait_idle()
        self.assertEqual(page.total, 0)
        self.assertEqual(self.errors, [])

    def test_followup_calendar_dst_validation_and_failed_save_keeps_inputs(self):
        from PySide6.QtCore import QDate, QTime
        self.open_test_project()
        self.window.navigate("followups")
        self.wait_idle()
        page = self.window.pages["followups"]
        page.action_buttons["new_reminder"][0].click()
        self.wait_idle()
        dialog = self.window._record_dialog
        dialog.title.setText("Zeitumstellung prüfen")
        dialog.date.setDate(QDate(2026, 3, 29))
        dialog.time.setTime(QTime(2, 30))
        dialog.save()
        self.assertIn("existiert", self.errors.pop())
        dialog.date.setDate(QDate(2026, 10, 25))
        self.assertTrue(dialog.occurrence.isVisibleTo(dialog))
        dialog.save()
        self.assertIn("doppeldeutig", self.errors.pop())
        dialog.occurrence.setCurrentIndex(dialog.occurrence.findData(1))
        emitted = []
        dialog.operation_requested.disconnect()
        dialog.operation_requested.connect(lambda action, payload: emitted.append(payload))
        dialog.save()
        self.assertEqual(emitted[0]["due_at_utc"], "2026-10-25T01:30:00Z")
        dialog.reject()
        page.action_buttons["new_reminder"][0].click()
        self.wait_idle()
        dialog = self.window._record_dialog
        dialog.title.setText("Inputs preserved")
        dialog.operation_requested.disconnect()
        dialog.operation_requested.connect(lambda action, payload: self.window.controller.submit(action, {**payload, "application_id": "missing"}))
        dialog.save()
        self.wait_idle()
        self.assertTrue(self.errors.pop())
        self.assertIs(self.window._record_dialog, dialog)
        self.assertEqual(dialog.title.text(), "Inputs preserved")
        self.assertTrue(dialog.buttons.isEnabled())
        dialog.reject()

    def test_followup_tick_guards_modal_busy_and_project_changes(self):
        self.open_test_project()
        self.submit("create_record", {"entity": "applications", "fields": {"title": "Due", "next_follow_up_at": "2000-01-01T09:00:00Z"}})
        with patch.object(self.window.controller, "submit", return_value=True) as submit:
            self.window.followup_tick()
            submit.assert_called_once_with("followup_summary")
            submit.reset_mock()
            self.window.controller.busy = True
            self.window.followup_tick()
            self.window.controller.busy = False
            with patch.object(QApplication, "activeModalWidget", return_value=self.window):
                self.window.followup_tick()
            self.window.pages["import"].auto_running = True
            self.window.followup_tick()
            self.window.pages["import"].auto_running = False
            submit.assert_not_called()
        self.submit("close_project", {})
        self.assertFalse(self.window.pages["overview"].followup_button.isEnabled())
        self.assertEqual(self.window.navigation.item(self.window.page_order.index("followups")).text(), "Nachfassen")
        self.submit("create_project", {"path": str(self.directory / "other"), "name": "Other"})
        self.assertEqual(self.window.snapshot["counts"]["followups"], 0)

    def test_followup_search_pagination_selection_refresh_and_screenshot(self):
        self.open_test_project()
        self.submit("create_record", {"entity": "applications", "fields": {"title": "Softwareentwickler", "next_follow_up_at": "2000-01-01T09:00:00Z"}})
        for title in ("HR nach Unterlagen fragen", "Interview vorbereiten", "Kontakt 100%"):
            self.submit("reminder_save", {"title": title, "due_at_utc": "2000-01-02T10:00:00Z"})
        self.window.navigate("followups")
        self.wait_idle()
        page = self.window.pages["followups"]
        page.limit = 2
        page.request_load()
        self.wait_idle()
        self.assertTrue(page.next.isEnabled())
        page.move_page(1)
        self.wait_idle()
        self.assertEqual(page.offset, 2)
        page.table.selectRow(0)
        selected = page.selected_record()["record_id"]
        self.window.followup_tick()
        self.wait_idle()
        self.assertEqual(page.selected_record()["record_id"], selected)
        page.search.setText("%")
        page.search_records()
        self.wait_idle()
        self.assertEqual(page.total, 1)
        self.assertEqual(page.rows[0]["title"], "Kontakt 100%")
        page.search.clear()
        self.wait_idle()
        page.limit = 50
        page.request_load()
        self.wait_idle()
        screenshot_path = os.environ.get("EMAILCENTER_FOLLOWUP_SCREENSHOT_PATH")
        if screenshot_path:
            self.window.grab().save(screenshot_path)
        self.assertEqual(self.errors, [])

    def test_record_editor_creates_edits_and_keeps_values_after_failed_save(self):
        self.open_test_project()
        self.window.navigate("companies")
        self.wait_idle()
        self.window.create_record("companies")
        self.wait_idle()
        dialog = self.window._record_dialog
        dialog.fields["name"].setText("Example GmbH")
        dialog.fields["notes"].setPlainText("Line one\nLine two")
        dialog.validate()
        self.wait_idle()
        self.assertIsNone(self.window._record_dialog)
        page = self.window.pages["companies"]
        page.table.selectRow(0)
        page.action_buttons["edit"][0].click()
        self.wait_idle()
        dialog = self.window._record_dialog
        self.assertEqual(dialog.fields["notes"].toPlainText(), "Line one\nLine two")
        dialog.fields["name"].setText("Updated GmbH")
        dialog.set_busy(True)
        dialog.reject()
        self.assertIs(self.window._record_dialog, dialog)
        self.assertFalse(dialog.buttons.isEnabled())
        dialog.set_busy(False)
        dialog.validate()
        self.wait_idle()
        self.assertEqual(page.rows[0]["name"], "Updated GmbH")
        self.assertEqual(self.errors, [])

    def test_profile_jobs_and_application_job_assignment_in_editor(self):
        self.open_test_project()
        self.submit("create_record", {"entity": "companies", "fields": {"name": "Example GmbH"}})
        company_id = self.window.snapshot["choices"]["companies"][0]["id"]
        self.window.navigate("applicants")
        self.wait_idle()
        self.window.create_record("applicants")
        self.wait_idle()
        self.window._record_dialog.fields["display_name"].setText("Ada")
        self.window._record_dialog.validate()
        self.wait_idle()
        self.assertEqual(self.window.pages["applicants"].rows[0]["display_name"], "Ada")
        self.window.navigate("jobs")
        self.wait_idle()
        self.window.create_record("jobs")
        self.wait_idle()
        dialog = self.window._record_dialog
        dialog.fields["title"].setText("Softwareentwickler")
        dialog.fields["company_id"].setCurrentIndex(dialog.fields["company_id"].findData(company_id))
        dialog.fields["location"].setText("Berlin")
        dialog.fields["deadline"].setText("2026-02-30")
        dialog.validate()
        self.wait_idle()
        self.assertIn("Datum", self.errors.pop())
        self.assertIs(self.window._record_dialog, dialog)
        self.assertTrue(dialog.buttons.isEnabled())
        self.assertEqual(dialog.fields["title"].text(), "Softwareentwickler")
        dialog.fields["deadline"].setText("2026-12-01")
        dialog.validate()
        self.wait_idle()
        job = self.window.pages["jobs"].rows[0]
        self.assertEqual(job["company_name"], "Example GmbH")
        screenshot_path = os.environ.get("EMAILCENTER_RECORDS_SCREENSHOT_PATH")
        if screenshot_path:
            self.window.grab().save(screenshot_path)
        self.window.navigate("applications")
        self.wait_idle()
        self.window.create_record("applications")
        self.wait_idle()
        dialog = self.window._record_dialog
        dialog.fields["title"].setText("Bewerbung Softwareentwickler")
        dialog.fields["job_posting_id"].setCurrentIndex(dialog.fields["job_posting_id"].findData(job["id"]))
        dialog.fields["applicant_profile_id"].setCurrentIndex(1)
        dialog.validate()
        self.wait_idle()
        application = self.window.pages["applications"].rows[0]
        self.assertEqual(application["job_posting_id"], job["id"])
        self.assertEqual(application["company_id"], company_id)
        self.assertEqual(self.errors, [])

    def test_contact_relationship_dialog_upserts_primary_and_removes(self):
        self.open_test_project()
        self.submit("create_record", {"entity": "companies", "fields": {"name": "Example GmbH"}})
        for name in ("Ada", "Grace"):
            self.submit("create_record", {"entity": "contacts", "fields": {"display_name": name}})
        self.window.navigate("companies")
        self.wait_idle()
        page = self.window.pages["companies"]
        page.table.selectRow(0)
        page.action_buttons["links"][0].click()
        self.wait_idle()
        dialog = self.window._record_dialog
        self.assertFalse(dialog.save_button.isEnabled())
        dialog.contact.setCurrentIndex(1)
        dialog.role.setText("Recruiting")
        dialog.primary.setChecked(True)
        dialog.save_button.click()
        self.wait_idle()
        self.assertEqual(len(dialog.data["rows"]), 1)
        dialog.table.selectRow(0)
        self.assertEqual(dialog.role.text(), "Recruiting")
        dialog.role.setText("HR")
        dialog.save_button.click()
        self.wait_idle()
        self.assertEqual(len(dialog.data["rows"]), 1)
        self.assertEqual(dialog.data["rows"][0]["relationship_type"], "HR")
        dialog.contact.setCurrentIndex(2)
        dialog.primary.setChecked(True)
        dialog.save_button.click()
        self.wait_idle()
        self.assertEqual(sum(r["is_primary"] for r in dialog.data["rows"]), 1)
        dialog.table.selectRow(0)
        dialog.remove_button.click()
        self.wait_idle()
        self.assertEqual(len(dialog.data["rows"]), 1)
        dialog.reject()
        self.assertEqual(self.errors, [])

    def test_record_editor_preserves_long_imported_values_and_archived_choices(self):
        from emailcenter.ui.record_dialogs import RecordDialog
        notes = "  " + "Long text\n" * 500 + "  "
        website = "x" * 40000
        data = {"entity": "companies", "record": {"id": "id", "name": "Example", "notes": notes, "website": website, "industry": None},
                "version": "fingerprint", "choices": {"companies": [], "applicants": [], "jobs": []}, "statuses": []}
        dialog = RecordDialog(data, self.window)
        emitted = []
        dialog.operation_requested.connect(lambda action, payload: emitted.append(payload))
        self.assertEqual(dialog.fields["website"].text(), website)
        dialog.fields["name"].setText("Renamed")
        dialog.validate()
        self.assertEqual(emitted[0]["fields"], {"name": "Renamed"})
        dialog.reject()
        data.update(entity="applications", record={"id": "id", "title": "Title", "company_id": "archived", "job_posting_id": None,
                                                   "applicant_profile_id": None}, version="fingerprint")
        dialog = RecordDialog(data, self.window)
        self.assertEqual(dialog.fields["company_id"].currentData(), "archived")
        self.assertIn("Nicht verfügbar", dialog.fields["company_id"].currentText())
        dialog.reject()

    def test_mapping_dialog_and_document_generation_inputs(self):
        record = {"id": "version-id", "template_name": "Bewerbung", "version_number": 1,
                  "configuration_json": '{"required_placeholders": ["NAME"]}',
                  "placeholders_json": '["NAME", "COMPANY"]'}
        mapping = TemplateDialog("letter.docx", ["NAME", "COMPANY"], parent=self.window)
        mapping.mapping_widgets["NAME"].setCurrentIndex(mapping.mapping_widgets["NAME"].findData("applicant.display_name"))
        mapping.required_widgets["NAME"].setChecked(True)
        self.assertEqual(mapping.values()["mapping"], {"NAME": "applicant.display_name"})
        snapshot = {"choices": {"applications": []}, "libreoffice": None}
        dialog = GenerateDialog(record, snapshot, self.window)
        emitted = []
        dialog.operation_requested.connect(lambda action, payload: emitted.append((action, payload)))
        dialog.request_operation("generate_document")
        self.assertFalse(emitted)
        self.assertIn("NAME", self.errors.pop())
        dialog.populate({"NAME": "Ada", "COMPANY": "Example"})
        dialog.request_operation("generate_document")
        self.assertEqual(emitted[0][0], "generate_document")
        self.assertEqual(emitted[0][1]["values"]["NAME"], "Ada")
        self.assertFalse(dialog.with_pdf.isEnabled())
        dialog.set_busy(True)
        self.assertFalse(dialog.generate_button.isEnabled())
        dialog.set_busy(False)
        dialog.reject()
        mapping.reject()

    def test_preview_dialog_displays_text_without_interpreting_html(self):
        dialog = PreviewDialog({"before": "{{NAME}}", "after": "<script>Ada</script>", "unresolved": []}, self.window)
        dialog.show()
        from PySide6.QtWidgets import QPlainTextEdit
        editor = dialog.findChild(QPlainTextEdit, "after")
        self.assertEqual(editor.toPlainText(), "<script>Ada</script>")
        dialog.accept()

    def test_document_workflow_uses_background_worker_and_releases_preview(self):
        self.open_test_project()
        source = self.directory / "template.docx"
        _make_docx(source)
        self.submit("import_template", {"source_path": str(source), "name": "Bewerbung",
                                        "mapping": {}, "required_placeholders": ["NAME"]})
        self.window.navigate("templates")
        self.wait_idle()
        record = self.window.pages["templates"].rows[0]
        self.window.record_action("generate", record)
        dialog = self.window._generation_dialog
        dialog.populate({"NAME": "Ada", "COMPANY": "Example GmbH"})
        dialog.request_operation("preview_document")
        self.wait_idle()
        preview = self.window._preview_dialog
        self.assertIsNotNone(preview)
        temporary_folder = Path(self.window.snapshot["path"]) / "tmp"
        self.assertTrue(list(temporary_folder.iterdir()))
        preview.accept()
        self.wait_idle()
        self.assertEqual(list(temporary_folder.iterdir()), [])
        with patch.object(QMessageBox, "information", return_value=QMessageBox.StandardButton.Ok):
            dialog.request_operation("generate_document")
            self.wait_idle()
        self.assertIsNone(self.window._generation_dialog)
        self.assertEqual(self.window.snapshot["counts"]["documents"], 1)
        self.window.navigate("documents")
        self.wait_idle()
        self.assertEqual(self.window.pages["documents"].rows[0]["original_filename"], "Bewerbung.docx")

    def test_pdf_preview_loads_a_real_pdf(self):
        from PySide6.QtGui import QPainter, QPdfWriter
        from PySide6.QtPdf import QPdfDocument
        path = self.directory / "preview.pdf"
        writer = QPdfWriter(str(path))
        painter = QPainter(writer)
        painter.drawText(100, 200, "EmailCenter PDF preview")
        painter.end()
        del painter, writer
        dialog = PreviewDialog({"pdf_path": str(path)}, self.window)
        self.assertEqual(dialog.pdf_widget.document.status(), QPdfDocument.Status.Ready)
        self.assertEqual(dialog.pdf_widget.document.pageCount(), 1)
        dialog.dispose()
        self.assertEqual(dialog.pdf_widget.document.status(), QPdfDocument.Status.Null)
        dialog.accept()

    def test_worker_owns_database_and_gui_receives_main_thread_signals(self):
        threads = []
        self.window.controller.completed.connect(lambda action, value: threads.append((action, QThread.currentThread())))
        self.open_test_project()
        self.assertEqual(threads, [("create_project", self.application.thread()),
                                  ("workspace_mail_snapshot", self.application.thread())])
        # A second request is rejected while the first is queued.
        self.assertTrue(self.window.controller.submit("snapshot"))
        self.assertFalse(self.window.controller.submit("close_project"))
        self.assertFalse(self.window.controller.shutdown())
        self.wait_idle()
        self.assertEqual(threads[2:], [("snapshot", self.application.thread()),
                                      ("workspace_mail_snapshot", self.application.thread())])

    def test_project_error_keeps_interface_usable(self):
        self.assertTrue(self.window.controller.submit("open_project", {"path": str(self.directory / "missing")}))
        self.wait_idle()
        self.assertEqual(len(self.errors), 1)
        self.errors.clear()
        self.assertFalse(self.window.controller.busy)
        self.assertTrue(self.window.actions["open"].isEnabled())
        self.open_test_project()

    def test_theme_can_be_changed_without_recreating_pages(self):
        page = self.window.current_page()
        self.window.toggle_theme()
        self.assertTrue(self.window.dark)
        self.assertIs(self.window.current_page(), page)
        self.window.toggle_theme()
        self.assertFalse(self.window.dark)

    def test_import_mapping_review_correction_and_resume_in_gui(self):
        self.open_test_project()
        self.window.navigate("import")
        self.wait_idle()
        source = self.directory / "contacts.csv"
        source.write_text("Name,E-Mail\nAda,ada@example.test\nGrace,bad-mail\nAda Copy,ada@example.test\n", encoding="utf-8")
        self.submit("inspect_import", {"source_path": str(source)})
        dialog = self.window._import_mapping_dialog
        self.assertEqual(dialog.mapping_widgets["display_name"].currentData(), "Name")
        self.assertEqual(dialog.mapping_widgets["email"].currentData(), "E-Mail")
        dialog.profile_name.setText("Kontakte")
        dialog.prepare()
        self.wait_idle()
        self.assertIsNone(self.window._import_mapping_dialog)
        page = self.window.pages["import"]
        self.assertEqual(page.run["counts"], {"Valid": 1, "Invalid": 1, "DuplicateReview": 1})
        page.start_import()
        self.wait_for(lambda: not page.auto_running and not self.window.controller.busy)
        self.assertEqual(page.run["counts"], {"Imported": 1, "Invalid": 1, "DuplicateReview": 1})
        self.assertEqual(self.window.snapshot["counts"]["contacts"], 1)
        page.table.selectRow(2)
        self.assertTrue(page.force_button.isEnabled())
        screenshot_path = os.environ.get("EMAILCENTER_IMPORT_SCREENSHOT_PATH")
        if screenshot_path:
            self.window.grab().save(screenshot_path)
        page.review_selected("import_new")
        self.wait_idle()
        page.start_import()
        self.wait_for(lambda: not page.auto_running and not self.window.controller.busy)
        invalid = next(row for row in page.rows if row["status"] == "Invalid")
        self.submit("edit_import_row", {"row_id": invalid["id"], "fields": {"display_name": "Grace", "email": "grace@example.test"}})
        page.start_import()
        self.wait_for(lambda: not page.auto_running and not self.window.controller.busy)
        self.assertEqual(page.run["counts"], {"Imported": 3})
        self.assertFalse(page.process_button.isEnabled())
        run_id = page.run_id
        project_path = self.window.snapshot["path"]
        self.submit("close_project", {})
        self.submit("open_project", {"path": project_path})
        self.window.navigate("import")
        self.wait_idle()
        self.assertEqual(page.run_id, run_id)
        self.assertEqual(page.run["counts"], {"Imported": 3})
        self.submit("process_import", {"run_id": run_id})
        self.assertEqual(self.window.snapshot["counts"]["contacts"], 3)

    def test_import_pause_between_batches_keeps_progress_and_continues(self):
        self.open_test_project()
        self.window.navigate("import")
        self.wait_idle()
        source = self.directory / "many.csv"
        source.write_text("Name,E-Mail\n" + "".join(f"Person {index},p{index}@example.test\n" for index in range(150)), encoding="utf-8")
        self.submit("prepare_import", {"source_path": str(source), "entity": "contacts", "mapping": {"display_name": "Name", "email": "E-Mail"}})
        page = self.window.pages["import"]
        pauses = []
        def pause_once(action, value):
            if action == "process_import" and not pauses:
                pauses.append(True)
                self.window.pause_current_import()
        self.window.controller.completed.connect(pause_once)
        page.start_import()
        self.wait_for(lambda: not page.auto_running and not self.window.controller.busy)
        self.assertEqual(page.run["counts"], {"Imported": 100, "Valid": 50})
        self.assertTrue(self.window.navigation.isEnabled())
        self.assertTrue(page.process_button.isEnabled())
        page.start_import()
        self.wait_for(lambda: not page.auto_running and not self.window.controller.busy)
        self.assertEqual(page.run["counts"], {"Imported": 150})
        self.assertEqual(self.window.snapshot["counts"]["contacts"], 150)

    def test_import_missing_required_mapping_does_not_create_a_run(self):
        self.open_test_project()
        source = self.directory / "contacts.csv"
        source.write_text("Other\nValue\n", encoding="utf-8")
        self.submit("inspect_import", {"source_path": str(source)})
        dialog = self.window._import_mapping_dialog
        dialog.prepare()
        self.assertEqual(len(self.errors), 1)
        self.errors.clear()
        self.assertFalse(self.window.controller.busy)
        self.assertIsNotNone(self.window._import_mapping_dialog)
        dialog.reject()

    def test_saved_import_profile_is_reused_without_altering_the_run(self):
        self.open_test_project()
        source = self.directory / "companies.csv"
        source.write_text("Organisation,Site\nExample GmbH,https://example.test\n", encoding="utf-8")
        self.submit("prepare_import", {"source_path": str(source), "entity": "companies",
                                       "mapping": {"name": "Organisation", "website": "Site"}, "profile_name": "Firmen"})
        original_run = self.window.pages["import"].run
        self.submit("inspect_import", {"source_path": str(source)})
        dialog = self.window._import_mapping_dialog
        self.assertEqual(dialog.profile.count(), 2)
        dialog.profile.setCurrentIndex(1)
        self.wait_idle()
        self.assertEqual(dialog.target.currentData(), "companies")
        self.assertEqual(dialog.mapping_widgets["name"].currentData(), "Organisation")
        self.assertEqual(dialog.mapping_widgets["website"].currentData(), "Site")
        self.assertEqual(self.window.pages["import"].run["configuration"], original_run["configuration"])
        dialog.reject()

    def enable_fake_gmail(self):
        from test_mail import FakeProvider, MemoryVault
        backend = MemoryVault()
        self.fake_provider = FakeProvider()
        self.mail_patches = [patch("emailcenter.application.workspace.MailService.connect_gmail", new=lambda service, client_path: service.register_account("sender@example.test", {"refresh_token": "fake"})),
                             patch("emailcenter.credentials.WindowsTokenStore.backend", new=property(lambda store: backend)),
                             patch("emailcenter.mail_service.MailService._provider", new=lambda service, account: self.fake_provider)]
        for value in self.mail_patches:
            value.start()
            self.addCleanup(value.stop)
        self.submit("mail_connect", {"client_path": "fake-only-no-real-file"})

    def create_mail_draft(self):
        page = self.window.pages["mail"]
        page.action("new")
        dialog = self.window._mail_dialog
        dialog.fields["to"].setText("target@example.test")
        dialog.fields["subject"].setText("Meine Bewerbung")
        dialog.body.setPlainText("Guten Tag,\nmeine Bewerbung.")
        dialog.save()
        self.wait_idle()
        self.assertIsNone(self.window._mail_dialog)
        self.assertEqual(len(page.data["drafts"]), 1)
        return page.data["drafts"][0]

    def test_mail_compose_review_send_history_and_uncertain_lock(self):
        self.open_test_project()
        self.window.navigate("mail")
        self.wait_idle()
        self.enable_fake_gmail()
        draft = self.create_mail_draft()
        page = self.window.pages["mail"]
        self.submit("mail_prepare_job", {"draft_id": draft["id"]})
        dialog = self.window._mail_dialog
        self.assertFalse(dialog.approve_button.isEnabled())
        self.assertEqual(dialog.body.toPlainText(), draft["body"])
        self.assertEqual(self.fake_provider.sent, [])
        dialog.confirm.setChecked(True)
        dialog.approve()
        self.wait_idle()
        job = page.data["jobs"][0]
        page.tables["jobs"].selectRow(0)
        self.assertTrue(page.buttons["send"].isEnabled())
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            page.action("send")
        self.wait_idle()
        self.assertEqual(page.data["jobs"][0]["status"], "Sent")
        self.assertEqual(len(self.fake_provider.sent), 1)
        page.tables["jobs"].selectRow(0)
        self.assertFalse(page.buttons["send"].isEnabled())
        self.submit("mail_prepare_job", {"draft_id": draft["id"]})
        self.window._mail_dialog.reject()
        job = page.data["jobs"][0]
        self.submit("mail_approve", {"job_id": job["id"]})
        from emailcenter.gmail import ProviderError
        self.fake_provider.failure = ProviderError("timeout", uncertain=True)
        self.submit("mail_send", {"job_id": job["id"]})
        page.tables["jobs"].selectRow(0)
        self.assertFalse(page.buttons["send"].isEnabled())
        self.assertTrue(page.buttons["reconcile"].isEnabled())
        self.assertEqual(page.data["jobs"][0]["status"], "Uncertain")
        screenshot_path = os.environ.get("EMAILCENTER_MAIL_SCREENSHOT_PATH")
        if screenshot_path:
            self.window.grab().save(screenshot_path)

    def test_mail_composer_rejects_invalid_addresses_without_saving(self):
        self.open_test_project()
        self.window.navigate("mail")
        self.wait_idle()
        self.window.pages["mail"].action("new")
        dialog = self.window._mail_dialog
        dialog.fields["to"].setText("invalid")
        dialog.fields["subject"].setText("Subject")
        dialog.body.setPlainText("Body")
        dialog.save()
        self.assertEqual(len(self.errors), 1)
        self.errors.clear()
        self.assertEqual(self.window.pages["mail"].data["drafts"], [])
        dialog.reject()

    def test_mail_scheduler_requires_opt_in_and_resets_on_project_close(self):
        self.open_test_project()
        self.window.navigate("mail")
        self.wait_idle()
        page = self.window.pages["mail"]
        actions = []
        self.window.controller.completed.connect(lambda action, value: actions.append(action))
        self.window.mail_tick()
        self.assertEqual(actions, [])
        page.scheduler.setChecked(True)
        self.window.mail_tick()
        self.wait_idle()
        self.assertIn("mail_mark_overdue", actions)
        self.window.mail_tick()
        self.wait_idle()
        self.assertIn("mail_due", actions)
        self.window._last_mail_tick -= 120
        self.window.mail_tick()
        self.wait_idle()
        self.assertEqual(actions.count("mail_mark_overdue"), 2)
        self.submit("close_project", {})
        self.assertFalse(page.scheduler.isChecked())

    def test_mail_edit_preserves_multiline_body_and_review_is_plain_text(self):
        self.open_test_project()
        self.window.navigate("mail")
        self.wait_idle()
        self.enable_fake_gmail()
        draft = self.create_mail_draft()
        page = self.window.pages["mail"]
        page.tables["drafts"].selectRow(0)
        page.action("edit")
        dialog = self.window._mail_dialog
        self.assertEqual(dialog.body.toPlainText(), draft["body"])
        dialog.body.setPlainText("<b>HTML als Text</b>\nZweite Zeile")
        dialog.save()
        self.wait_idle()
        self.submit("mail_prepare_job", {"draft_id": draft["id"]})
        self.assertEqual(self.window._mail_dialog.body.toPlainText(), "<b>HTML als Text</b>\nZweite Zeile")
        self.window._mail_dialog.reject()

    def test_import_row_correction_keeps_long_multiline_notes(self):
        from emailcenter.ui.dialogs import FormDialog
        from PySide6.QtWidgets import QDialog, QPlainTextEdit
        self.open_test_project()
        source = self.directory / "contacts.csv"
        notes = "First line\n" + "Detail " * 500
        source.write_text('Name,E-Mail,Notizen\nAda,bad-mail,"' + notes + '"\n', encoding="utf-8")
        self.submit("prepare_import", {"source_path": str(source), "entity": "contacts",
                                       "mapping": {"display_name": "Name", "email": "E-Mail", "notes": "Notizen"}})
        page = self.window.pages["import"]
        page.table.selectRow(0)
        def correct(dialog):
            self.assertIsInstance(dialog.fields["notes"], QPlainTextEdit)
            self.assertEqual(dialog.fields["notes"].toPlainText(), notes.strip())
            dialog.fields["email"].setText("ada@example.test")
            return QDialog.DialogCode.Accepted
        with patch.object(FormDialog, "exec", new=correct):
            page.edit_selected()
        self.wait_idle()
        self.assertEqual(page.run["counts"], {"Valid": 1})
        self.assertEqual(page.rows[0]["mapped_data"]["fields"]["notes"], notes.strip())

    @unittest.skipUnless(__import__("importlib.util", fromlist=["find_spec"]).find_spec("openpyxl"), "openpyxl not installed")
    def test_excel_sheet_can_be_changed_before_preparing_import(self):
        from openpyxl import Workbook
        self.open_test_project()
        source = self.directory / "contacts.xlsx"
        workbook = Workbook()
        workbook.active.title = "Cover"
        workbook.active.append(["Description"])
        sheet = workbook.create_sheet("Contacts")
        sheet.append(["Name", "E-Mail"])
        sheet.append(["Ada", "ada@example.test"])
        workbook.save(source)
        workbook.close()
        self.submit("inspect_import", {"source_path": str(source)})
        dialog = self.window._import_mapping_dialog
        dialog.sheet.setCurrentText("Contacts")
        self.wait_idle()
        self.assertEqual(dialog.inspection["headers"], ["Name", "E-Mail"])
        dialog.prepare()
        self.wait_idle()
        page = self.window.pages["import"]
        self.assertEqual(page.run["configuration"]["options"]["sheet"], "Contacts")
        self.assertEqual(page.run["counts"], {"Valid": 1})


    def test_mail_template_library_versions_duplicate_search_and_pagination(self):
        self.open_test_project()
        self.window.workspace_action("mail_templates", None)
        self.wait_idle()
        library = self.window._template_library
        self.assertEqual(library.data["total"], 0)
        library.new_button.click()
        editor = self.window._template_editor
        editor.name.setText("100% Bewerbung")
        editor.subject.setText("Bewerbung {{company.name}}")
        editor.body.setPlainText("Guten Tag {{contact.display_name}},")
        editor.save()
        self.wait_idle()
        self.assertIsNone(self.window._template_editor)
        self.assertEqual(library.data["total"], 1)
        library.records.setCurrentRow(0)
        library.edit_button.click()
        self.wait_idle()
        editor = self.window._template_editor
        editor.body.setPlainText("Neue Fassung")
        editor.save()
        self.wait_idle()
        self.assertEqual(library.data["rows"][0]["version_number"], 2)
        library.records.setCurrentRow(0)
        library.edit_button.click()
        self.wait_idle()
        editor = self.window._template_editor
        editor.versions.setCurrentIndex(1)
        self.assertIn("contact.display_name", editor.body.toPlainText())
        editor.duplicate_button.click()
        self.wait_idle()
        self.assertEqual(library.data["total"], 2)
        copy = next(row for row in library.data["rows"] if "Kopie" in row["name"])
        self.assertEqual(copy["version_number"], 1)
        library.limit = 1
        library.request_load()
        self.wait_idle()
        library.next.click()
        self.wait_idle()
        self.assertEqual(library.offset, 1)
        library.search.setText("%")
        library.search_records()
        self.wait_idle()
        self.assertEqual(library.offset, 0)
        self.assertEqual(library.data["total"], 2)
        library.search.setText("' OR 1=1 --")
        library.search_records()
        self.wait_idle()
        self.assertEqual(library.data["total"], 0)
        library.reject()

    def test_mail_template_editor_validation_conflict_reload_and_placeholder_insertion(self):
        self.open_test_project()
        self.window.open_template_library()
        self.wait_idle()
        library = self.window._template_library
        library.new_button.click()
        editor = self.window._template_editor
        editor.subject.setText("Betreff")
        editor.body.setPlainText("Text ")
        editor.placeholder.setCurrentText("company.name")
        editor.insert_placeholder()
        self.assertIn("{{company.name}}", editor.body.toPlainText())
        editor.set_busy(True)
        editor.reject()
        self.assertIs(self.window._template_editor, editor)
        editor.set_busy(False)
        editor.save()
        self.wait_idle()
        self.assertIn("Vorlagennamen", self.errors.pop())
        self.assertIs(self.window._template_editor, editor)
        editor.name.setText("Meine Vorlage")
        editor.save()
        self.wait_idle()
        library.records.setCurrentRow(0)
        library.edit_button.click()
        self.wait_idle()
        editor = self.window._template_editor
        editor.body.setPlainText("Nicht verlieren")
        with patch("emailcenter.application.workspace.MailTemplateService.save", side_effect=ValueError("Vorlage wurde geändert")):
            editor.save()
            self.wait_idle()
        self.assertIn("geändert", self.errors.pop())
        self.assertEqual(editor.body.toPlainText(), "Nicht verlieren")
        self.assertTrue(editor.isEnabled())
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            editor.duplicate_button.click()
        self.assertIs(self.window._template_editor, editor)
        self.assertEqual(editor.body.toPlainText(), "Nicht verlieren")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            editor.reload_button.click()
        self.assertEqual(editor.body.toPlainText(), "Nicht verlieren")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            editor.reload_button.click()
            self.wait_idle()
        self.assertIn("{{company.name}}", editor.body.toPlainText())
        editor.reject()
        self.submit("close_project", {})
        self.assertIsNone(self.window._template_library)

    def test_mail_template_context_preview_missing_values_and_apply_preserves_envelope(self):
        self.open_test_project()
        application, _, _, generated = self.context_fixture(documents=True)
        self.enable_fake_gmail()
        version = self.action_result("mail_template_save", {"name": "Bewerbung", "subject": "Bewerbung an {{company.name}}",
            "body": "Guten Tag {{contact.display_name}},\n{{NOTE}}\n{{applicant.display_name}}", "body_format": "plain"})
        signature = self.action_result("mail_save_signature", {"name": "Privat", "content_text": "Viele Grüße"})
        self.submit("mail_compose_details", {"application_id": application["id"]})
        compose = self.window._mail_dialog
        compose.contact_choice.setCurrentIndex(1)
        compose.use_contact.click()
        compose.fields["cc"].setText("cc@example.test")
        compose.fields["bcc"].setText("bcc@example.test")
        compose.signature.setCurrentIndex(compose.signature.findData(signature))
        compose.documents.setCurrentIndex(1)
        compose.add_document()
        before = compose.values()
        compose.template_choice.setCurrentIndex(compose.template_choice.findData(version["id"]))
        compose.preview_template.click()
        self.wait_idle()
        preview = self.window._template_preview
        self.assertEqual(preview.result["missing"], ["NOTE"])
        self.assertEqual(preview.fields["contact.display_name"].toPlainText(), "Grace")
        preview.confirm.setChecked(True)
        self.assertFalse(preview.apply_button.isEnabled())
        preview.fields["NOTE"].setPlainText("Ich freue mich auf Ihre Rückmeldung.")
        self.assertFalse(preview.apply_button.isEnabled())
        preview.refresh_button.click()
        self.wait_idle()
        self.assertEqual(preview.result["missing"], [])
        self.assertTrue(preview.apply_button.isEnabled())
        screenshot = os.environ.get("EMAILCENTER_MAIL_TEMPLATE_PREVIEW_SCREENSHOT_PATH")
        if screenshot:
            self.application.processEvents()
            preview.grab().save(screenshot)
        preview.apply_button.click()
        self.assertIsNone(self.window._template_preview)
        self.assertIn("Ich freue mich", compose.body.toPlainText())
        after = compose.values()
        for key in ("recipients_value", "provider_account_id", "signature_id", "application_id", "attachments"):
            self.assertEqual(after[key], before[key])
        compose.save()
        self.wait_idle()
        draft = self.window.pages["workspace"].mail_workflow.data["drafts"][0]
        self.assertEqual(draft["configuration"]["attachments"][0]["document_id"], generated["document"]["id"])
        self.assertEqual(self.fake_provider.sent, [])
        self.assertEqual(draft["subject"], "Bewerbung an Example GmbH")

    def test_mail_template_render_error_preserves_preview_values_and_cancel_preserves_compose(self):
        self.open_test_project()
        version = self.action_result("mail_template_save", {"name": "Manuelle Werte", "subject": "{{NAME}}", "body": "Hallo {{MESSAGE}}"})
        self.submit("mail_compose_details", {})
        compose = self.window._mail_dialog
        compose.fields["subject"].setText("Manueller Betreff")
        compose.body.setPlainText("Manueller Text")
        compose.template_choice.setCurrentIndex(compose.template_choice.findData(version["id"]))
        compose.request_template()
        self.wait_idle()
        preview = self.window._template_preview
        preview.fields["NAME"].setPlainText("Injected\nHeader")
        preview.fields["MESSAGE"].setPlainText("Nicht verlieren")
        preview.confirm.setChecked(True)
        preview.render()
        self.wait_idle()
        self.assertTrue(self.errors.pop())
        self.assertIs(self.window._template_preview, preview)
        self.assertEqual(preview.fields["MESSAGE"].toPlainText(), "Nicht verlieren")
        self.assertFalse(preview.apply_button.isEnabled())
        preview.fields["NAME"].setPlainText("Ada")
        preview.render()
        self.wait_idle()
        self.assertTrue(preview.apply_button.isEnabled())
        preview.set_busy(True)
        preview.reject()
        self.assertIs(self.window._template_preview, preview)
        preview.set_busy(False)
        preview.reject()
        self.assertEqual(compose.fields["subject"].text(), "Manueller Betreff")
        self.assertEqual(compose.body.toPlainText(), "Manueller Text")
        compose.reject()
        self.submit("snapshot", {})
        self.assertEqual(self.window.snapshot["counts"]["mail_drafts"], 0)

    def test_mail_template_library_from_compose_refreshes_versions_and_preserves_input(self):
        self.open_test_project()
        self.submit("mail_compose_details", {})
        compose = self.window._mail_dialog
        compose.fields["subject"].setText("Meine Eingabe")
        compose.template_library_requested.emit()
        self.wait_idle()
        library = self.window._template_library
        library.new_button.click()
        editor = self.window._template_editor
        editor.name.setText("Antwort")
        editor.subject.setText("Antwort auf Anfrage")
        editor.body.setPlainText("Guten Tag")
        editor.save()
        self.wait_idle()
        self.assertEqual(compose.template_choice.count(), 2)
        self.assertEqual(compose.fields["subject"].text(), "Meine Eingabe")
        self.assertIsNone(compose.template_choice.currentData())
        library.records.setCurrentRow(0)
        library.edit()
        self.wait_idle()
        editor = self.window._template_editor
        editor.body.setPlainText("Version 2")
        editor.save()
        self.wait_idle()
        self.assertEqual(compose.template_choice.count(), 3)
        library.reject()
        compose.template_choice.setCurrentIndex(2)
        compose.request_template()
        self.wait_idle()
        preview = self.window._template_preview
        self.assertEqual(preview.result["version_number"], 1)
        self.assertEqual(preview.body.toPlainText(), "Guten Tag")
        self.assertFalse(preview.apply_button.isEnabled())
        preview.confirm.setChecked(True)
        preview.apply()
        self.assertEqual(compose.body.toPlainText(), "Guten Tag")
        compose.reject()

    def test_mail_template_html_and_changed_recipient_require_explicit_contact_value(self):
        self.open_test_project()
        contact = self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": "Grace", "email": "grace@example.test"}})
        version = self.action_result("mail_template_save", {"name": "HTML", "subject": "Nachricht", "body_format": "html",
            "body": "<p>Hallo {{contact.display_name}}</p>"})
        self.submit("mail_compose_details", {"contact_id": contact["id"]})
        compose = self.window._mail_dialog
        compose.fields["to"].setText("other@example.test")
        compose.template_choice.setCurrentIndex(compose.template_choice.findData(version["id"]))
        compose.request_template()
        self.wait_idle()
        preview = self.window._template_preview
        self.assertEqual(preview.result["missing"], ["contact.display_name"])
        preview.fields["contact.display_name"].setPlainText("Other <HR>")
        preview.render()
        self.wait_idle()
        self.assertEqual(preview.body.toPlainText(), "<p>Hallo Other &lt;HR&gt;</p>")
        self.assertEqual(preview.format_label.text(), "HTML-Quelltext")
        preview.confirm.setChecked(True)
        preview.apply()
        self.assertEqual(compose.format.currentData(), "html")
        self.assertEqual(compose.fields["to"].text(), "other@example.test")
        compose.reject()


    def bulk_fixture(self, *, custom=False):
        from PySide6.QtCore import Qt
        contacts = [self.action_result("create_record", {"entity": "contacts", "fields": {
            "display_name": name, "email": email}}) for name, email in (("Ada", "ada@example.test"), ("Grace", "grace@example.test"))]
        version = self.action_result("mail_template_save", {"name": "Bulk persönlich", "subject": "An {{contact.display_name}}",
            "body": "Hallo {{contact.display_name}}" + (" · {{NOTE}}" if custom else "")})
        self.window.pages["workspace"].mail_workflow.bulk_button.click()
        self.wait_idle()
        dialog = self.window._bulk_dialog
        dialog.template.setCurrentIndex(dialog.template.findData(version["id"]))
        dialog.select_visible()
        self.assertEqual(dialog.selected, {c["id"] for c in contacts})
        return dialog, contacts, version

    def test_bulk_preview_creation_history_reopen_and_edit_single_draft(self):
        self.open_test_project()
        dialog, contacts, version = self.bulk_fixture()
        self.assertFalse(dialog.create_button.isEnabled())
        dialog.request_preview()
        self.wait_idle()
        self.assertEqual(dialog.preview["rows"][0]["name"], "Ada")
        self.assertTrue(dialog.preview["ready"])
        self.assertIn("Hallo Ada", dialog.message.toPlainText())
        dialog.recipient.setCurrentIndex(1)
        self.assertIn("Hallo Grace", dialog.message.toPlainText())
        self.assertFalse(dialog.create_button.isEnabled())
        dialog.confirm.setChecked(True)
        self.assertTrue(dialog.create_button.isEnabled())
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.create()
        self.wait_idle()
        self.assertEqual(self.window.snapshot["counts"]["mail_drafts"], 2)
        self.assertEqual(self.window.snapshot["counts"]["mail_jobs"], 0)
        self.assertEqual(dialog.tabs.currentIndex(), 1)
        self.assertFalse(dialog.create_button.isEnabled())
        batch_id = dialog.batch["id"]
        dialog.reject()
        self.submit("close_project", {})
        self.submit("open_project", {"path": str(self.directory / "project")})
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        dialog = self.window._bulk_dialog
        self.assertEqual(dialog.data["batches"][0]["id"], batch_id)
        dialog.tabs.setCurrentIndex(1)
        dialog.batches.setCurrentRow(0)
        self.wait_idle()
        self.assertEqual(dialog.drafts.count(), 2)
        dialog.drafts.setCurrentRow(1)
        dialog.open_draft()
        self.wait_idle()
        self.assertIsNone(self.window._bulk_dialog)
        self.assertEqual(self.window._mail_dialog.fields["to"].text(), "grace@example.test")
        self.assertEqual(self.window._mail_dialog.fields["subject"].text(), "An Grace")
        self.window._mail_dialog.reject()

    def test_bulk_search_pagination_preserves_checkbox_selection(self):
        from PySide6.QtCore import Qt
        self.open_test_project()
        for i in range(28):
            self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": f"Kontakt {i:02}", "email": f"c{i}@example.test"}})
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        dialog = self.window._bulk_dialog
        self.assertEqual(dialog.contacts.count(), 25)
        identifier = dialog.contacts.item(0).data(Qt.ItemDataRole.UserRole)
        dialog.contacts.item(0).setCheckState(Qt.CheckState.Checked)
        dialog.move_page(1)
        self.assertEqual(dialog.contacts.count(), 3)
        dialog.select_visible()
        self.assertEqual(len(dialog.selected), 4)
        dialog.search.setText("Kontakt 00")
        self.assertEqual(dialog.contacts.count(), 1)
        self.assertEqual(dialog.contacts.item(0).checkState(), Qt.CheckState.Checked)
        self.assertIn(identifier, dialog.selected)
        dialog.search.setText("' OR 1=1 --")
        self.assertEqual(dialog.contacts.count(), 0)
        self.assertEqual(len(dialog.selected), 4)
        dialog.search.clear()
        self.assertEqual(dialog.contacts.count(), 25)
        dialog.clear_selection()
        self.assertEqual(dialog.selected, set())
        dialog.reject()

    def test_bulk_missing_custom_values_and_edits_invalidate_confirmation(self):
        self.open_test_project()
        dialog, contacts, _ = self.bulk_fixture(custom=True)
        dialog.request_preview()
        self.wait_idle()
        self.assertFalse(dialog.preview["ready"])
        self.assertIn("NOTE", dialog.message.toPlainText())
        for index in range(dialog.recipient.count()):
            dialog.recipient.setCurrentIndex(index)
            row = next(i for i in range(dialog.values.rowCount()) if dialog.values.item(i, 0).text() == "NOTE")
            dialog.values.item(row, 1).setText(f"Wert {index}")
        self.assertFalse(dialog.fresh)
        dialog.request_preview()
        self.wait_idle()
        self.assertTrue(dialog.preview["ready"])
        self.assertIn("Wert 1", dialog.message.toPlainText())
        dialog.recipient.setCurrentIndex(0)
        self.assertIn("Wert 0", dialog.message.toPlainText())
        dialog.confirm.setChecked(True)
        self.assertTrue(dialog.create_button.isEnabled())
        row = next(i for i in range(dialog.values.rowCount()) if dialog.values.item(i, 0).text() == "NOTE")
        dialog.values.item(row, 1).setText("Änderung")
        self.assertFalse(dialog.create_button.isEnabled())
        self.assertFalse(dialog.confirm.isChecked())
        dialog.reject()
        self.submit("snapshot", {})
        self.assertEqual(self.window.snapshot["counts"]["mail_drafts"], 0)

    def test_bulk_duplicate_invalid_recipients_and_cancel_never_create_drafts(self):
        from PySide6.QtCore import Qt
        self.open_test_project()
        self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": "Duplicate", "email": "ADA@example.test"}})
        self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": "No address"}})
        version = self.action_result("mail_template_save", {"name": "Bulk", "subject": "Betreff", "body": "Hallo"})
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        dialog = self.window._bulk_dialog
        self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": "Ada", "email": "ada@example.test"}})
        dialog.reject()
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        dialog = self.window._bulk_dialog
        dialog.template.setCurrentIndex(dialog.template.findData(version["id"]))
        dialog.select_visible()
        self.assertIn("mehrfach", dialog.contacts.item(0).text())
        dialog.request_preview()
        self.wait_idle()
        self.assertFalse(dialog.preview["ready"])
        self.assertFalse(dialog.create_button.isEnabled())
        self.assertTrue(dialog.preview["rows"][-1]["error"])
        dialog.set_busy(True)
        dialog.reject()
        self.assertIs(self.window._bulk_dialog, dialog)
        dialog.set_busy(False)
        dialog.reject()
        self.submit("snapshot", {})
        self.assertEqual(self.window.snapshot["counts"]["mail_drafts"], 0)

    def test_bulk_application_documents_and_common_attachments_are_explicit(self):
        from PySide6.QtCore import Qt
        self.open_test_project()
        application, contacts, _, generated = self.context_fixture(documents=True)
        version = self.action_result("mail_template_save", {"name": "Bewerbung Bulk", "subject": "{{job.title}}", "body": "Hallo {{contact.display_name}}"})
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        dialog = self.window._bulk_dialog
        dialog.template.setCurrentIndex(dialog.template.findData(version["id"]))
        chosen = next(i for i in range(dialog.contacts.count()) if dialog.contacts.item(i).data(Qt.ItemDataRole.UserRole) == contacts[0]["id"])
        dialog.contacts.item(chosen).setCheckState(Qt.CheckState.Checked)
        dialog.application.setCurrentIndex(dialog.application.findData(application["id"]))
        self.wait_idle()
        self.assertEqual(dialog.documents.count(), 1)
        self.assertEqual(dialog.documents.item(0).checkState(), Qt.CheckState.Unchecked)
        dialog.request_preview()
        self.wait_idle()
        self.assertEqual(dialog.preview["rows"][0]["attachments"], [])
        dialog.documents.item(0).setCheckState(Qt.CheckState.Checked)
        attachment = self.directory / "common.txt"
        attachment.write_text("Für alle", encoding="utf-8")
        with patch.object(__import__("PySide6.QtWidgets", fromlist=["QFileDialog"]).QFileDialog, "getOpenFileNames", return_value=([str(attachment)], "")):
            dialog.add_files()
        dialog.request_preview()
        self.wait_idle()
        self.assertEqual(len(dialog.preview["rows"][0]["attachments"]), 2)
        screenshot = os.environ.get("EMAILCENTER_BULK_SCREENSHOT_PATH")
        if screenshot:
            dialog.grab().save(screenshot)
        dialog.confirm.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.create()
        self.wait_idle()
        draft = self.action_result("mail_bulk_batch", {"batch_id": dialog.batch["id"]})["items"][0]
        dialog.drafts.setCurrentRow(0)
        dialog.open_draft()
        self.wait_idle()
        self.assertEqual(self.window._mail_dialog.application.currentData(), application["id"])
        self.assertEqual(len(self.window._mail_dialog.values()["attachments"]), 2)
        self.window._mail_dialog.reject()

    def test_bulk_stale_file_error_preserves_input_and_project_change_closes_dialog(self):
        self.open_test_project()
        dialog, _, _ = self.bulk_fixture()
        attachment = self.directory / "file.txt"
        attachment.write_text("First", encoding="utf-8")
        dialog.files = [str(attachment)]
        dialog.update_files()
        dialog.name.setText("Nicht verlieren")
        dialog.request_preview()
        self.wait_idle()
        dialog.confirm.setChecked(True)
        attachment.write_text("Changed", encoding="utf-8")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.create()
        self.wait_idle()
        self.assertIn("geändert", self.errors.pop())
        self.assertEqual(dialog.name.text(), "Nicht verlieren")
        self.assertEqual(len(dialog.selected), 2)
        self.assertFalse(dialog.create_button.isEnabled())
        dialog.request_preview()
        self.wait_idle()
        self.assertTrue(dialog.preview["ready"])
        self.submit("close_project", {})
        self.assertIsNone(self.window._bulk_dialog)


    def source_bulk_fixture(self, *, excel=False, count=2):
        if excel:
            from openpyxl import Workbook
            path = self.directory / "Bewerbungen.xlsx"
            book = Workbook()
            book.active.append(["Name", "Email", "Anrede", "Firma", "Anschreiben", "Stelle"])
            for i in range(count):
                book.active.append([f"Person {i}", f"p{i}@example.test", "Guten Tag", f"Firma {i}", f"Text {i}", "Entwicklung"])
            book.save(path)
            book.close()
        else:
            path = self.directory / "Bewerbungen.csv"
            path.write_text("Name;Email;Anrede;Firma;Anschreiben;Stelle\n" + "".join(
                f"Person {i};p{i}@example.test;Guten Tag;Firma {i};Text {i};Entwicklung\n" for i in range(count)), encoding="utf-8")
        run = self.action_result("prepare_import", {"source_path": str(path), "entity": "contacts", "mapping": {"display_name": "Name", "email": "Email"}})
        version = self.action_result("mail_template_save", {"name": "Bewerbungen", "subject": "An {{company.name}}",
            "body": "{{contact.salutation}} {{contact.display_name}}: {{application.cover_letter}}"})
        self.window.navigate("workspace")
        self.wait_idle()
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        dialog = self.window._bulk_dialog
        dialog.template.setCurrentIndex(dialog.template.findData(version["id"]))
        return dialog, run

    def test_bulk_source_excel_mapping_personal_preview_and_snapshot_history(self):
        self.open_test_project()
        dialog, run = self.source_bulk_fixture(excel=True)
        self.window.open_bulk_source("import")
        self.wait_idle()
        source = self.window._bulk_source_dialog
        self.assertEqual(source.data["run"]["id"], run["id"])
        self.assertEqual(source.mapping["company.name"].currentData(), "Firma")
        source.select_visible()
        source.apply()
        self.wait_idle()
        self.assertIsNone(self.window._bulk_source_dialog)
        self.assertEqual(len(dialog.selected), 2)
        dialog.request_preview()
        self.wait_idle()
        self.assertTrue(dialog.preview["ready"], dialog.preview["rows"])
        self.assertIn("Guten Tag Person 0: Text 0", dialog.message.toPlainText())
        dialog.recipient.setCurrentIndex(1)
        self.assertIn("An Firma 1", dialog.message.toPlainText())
        dialog.confirm.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.create()
        self.wait_idle()
        self.assertEqual(self.window.snapshot["counts"]["contacts"], 0)
        self.assertEqual(self.window.snapshot["counts"]["mail_drafts"], 2)
        self.assertEqual(dialog.batch["items"][0]["source_snapshot"]["filename"], "Bewerbungen.xlsx")
        dialog.reject()

    def test_bulk_import_source_page_selection_and_missing_mapping_preserve_dialog(self):
        self.open_test_project()
        dialog, _ = self.source_bulk_fixture(count=28)
        self.window.open_bulk_source("import")
        self.wait_idle()
        source = self.window._bulk_source_dialog
        source.select_visible()
        self.assertEqual(len(source.selected), 25)
        source.move_page(1)
        self.wait_idle()
        self.assertEqual(source.rows.rowCount(), 3)
        source.select_visible()
        self.assertEqual(len(source.selected), 28)
        source.mapping["contact.email"].setCurrentIndex(0)
        self.assertFalse(source.apply_button.isEnabled())
        source.mapping["contact.email"].setCurrentIndex(source.mapping["contact.email"].findData("Email"))
        self.assertTrue(source.apply_button.isEnabled())
        source.set_busy(True)
        source.reject()
        self.assertIs(self.window._bulk_source_dialog, source)
        source.set_busy(False)
        source.reject()
        self.assertEqual(dialog.selected, set())
        dialog.reject()

    def test_bulk_manual_table_paste_and_personal_file_selection(self):
        self.open_test_project()
        version = self.action_result("mail_template_save", {"name": "Liste", "subject": "{{company.name}}", "body": "{{application.cover_letter}}"})
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        dialog = self.window._bulk_dialog
        dialog.template.setCurrentIndex(dialog.template.findData(version["id"]))
        self.window.open_bulk_source("manual")
        manual = self.window._bulk_source_dialog
        manual.rows.setCurrentCell(0, 0)
        manual.rows.paste_text('Ada\tada@example.test\tGuten Tag\tFirma A\tEntwicklung\t"Mehrzeilig\nAnschreiben"\nGrace\tgrace@example.test\tHallo\tFirma B\tDesign\tZweiter Text')
        self.assertEqual(manual.rows.rowCount(), 2)
        self.assertEqual(manual.rows.item(0, 5).text(), "Mehrzeilig\nAnschreiben")
        with self.assertRaises(ValueError):
            manual.rows.paste_text("\t".join(["x"] * (manual.rows.columnCount() + 1)))
        self.assertEqual(manual.rows.item(0, 0).text(), "Ada")
        manual.apply()
        self.assertIsNone(self.window._bulk_source_dialog)
        self.assertEqual(len(dialog.selected), 2)
        path = self.directory / "AdaCV.txt"
        path.write_text("Ada CV", encoding="utf-8")
        from PySide6.QtWidgets import QFileDialog
        with patch.object(QFileDialog, "getOpenFileNames", return_value=([str(path)], "")):
            dialog.add_personal_files()
        self.assertEqual(dialog.personal_file_status.text(), "1 persönliche Dateien")
        dialog.request_preview()
        self.wait_idle()
        self.assertTrue(dialog.preview["ready"])
        self.assertEqual(len(dialog.preview["rows"][0]["attachments"]), 1)
        self.assertEqual(dialog.preview["rows"][1]["attachments"], [])
        self.assertIn("Mehrzeilig\nAnschreiben", dialog.message.toPlainText())
        dialog.recipient.setCurrentIndex(1)
        self.assertEqual(dialog.personal_file_status.text(), "Keine persönlichen Dateien")
        dialog.reject()

    def test_bulk_manual_docx_text_into_selected_cell_and_empty_list(self):
        from PySide6.QtWidgets import QFileDialog
        self.open_test_project()
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        self.window.open_bulk_source("manual")
        manual = self.window._bulk_source_dialog
        manual.apply()
        self.assertIn("Mindestens", self.errors.pop())
        path = self.directory / "Text.docx"
        _make_docx(path)
        manual.rows.setCurrentCell(0, manual.fields.index("application.cover_letter"))
        with patch.object(QFileDialog, "getOpenFileName", return_value=(str(path), "")):
            manual.load_docx()
        self.wait_idle()
        self.assertIn("Hello", manual.rows.item(0, manual.fields.index("application.cover_letter")).text())
        self.assertIsNone(manual.pending_cell)
        manual.reject()
        self.assertEqual(self.window._bulk_dialog.selected, set())
        self.window._bulk_dialog.reject()

    def test_bulk_excel_document_template_generates_matching_application_attachments(self):
        import zipfile
        from test_documents import W
        self.open_test_project()
        path = self.directory / "Letter.docx"
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/document.xml", f'<w:document xmlns:w="{W}"><w:body><w:p><w:r><w:t>{{{{ANREDE}}}} {{{{NAME}}}}: {{{{TEXT}}}} bei {{{{FIRMA}}}}</w:t></w:r></w:p></w:body></w:document>')
        document = self.action_result("import_template", {"source_path": str(path), "name": "Anschreiben", "required_placeholders": [],
            "mapping": {"ANREDE": "contact.salutation", "NAME": "contact.display_name", "TEXT": "application.cover_letter", "FIRMA": "company.name"}})
        dialog, _ = self.source_bulk_fixture(excel=True)
        dialog.document_choice.setCurrentIndex(dialog.document_choice.findData(document["version"]["id"]))
        self.assertEqual(dialog.document_bindings.rowCount(), 4)
        self.window.open_bulk_source("import")
        self.wait_idle()
        source = self.window._bulk_source_dialog
        source.select_visible()
        source.apply()
        self.wait_idle()
        dialog.request_preview()
        self.wait_idle()
        self.assertTrue(dialog.preview["ready"], dialog.preview["rows"])
        self.assertIn("Person 0: Text 0 bei Firma 0", dialog.message.toPlainText())
        self.assertIn("Anschreiben-Person 0.docx", dialog.message.toPlainText())
        screenshot = os.environ.get("EMAILCENTER_BULK_SOURCE_SCREENSHOT_PATH")
        if screenshot:
            dialog.resize(1100, 900)
            self.application.processEvents()
            scrolls = dialog.findChildren(__import__("PySide6.QtWidgets", fromlist=["QScrollArea"]).QScrollArea)
            scrolls[0].verticalScrollBar().setValue(scrolls[0].verticalScrollBar().maximum())
            self.application.processEvents()
            dialog.grab().save(screenshot)
        dialog.confirm.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            dialog.create()
        self.wait_idle()
        self.assertEqual(self.window.snapshot["counts"]["documents"], 2)
        dialog.drafts.setCurrentRow(0)
        dialog.open_draft()
        self.wait_idle()
        compose = self.window._mail_dialog
        self.assertIn("Person 0: Text 0", compose.body.toPlainText())
        self.assertEqual(len(compose.values()["attachments"]), 1)
        compose.reject()


    def test_bulk_source_remapping_keeps_manual_overrides_and_cancelled_document_values_are_filtered(self):
        self.open_test_project()
        dialog, _ = self.source_bulk_fixture()
        self.window.open_bulk_source("import")
        self.wait_idle()
        source = self.window._bulk_source_dialog
        source.select_visible()
        source.apply()
        self.wait_idle()
        identifier = dialog.recipient.currentData()
        row = next(i for i in range(dialog.values.rowCount()) if dialog.values.item(i, 0).text() == "company.name")
        dialog.values.item(row, 1).setText("Manuelle Firma")
        self.window.open_bulk_source("import")
        self.wait_idle()
        source = self.window._bulk_source_dialog
        source.select_visible()
        source.apply()
        self.wait_idle()
        self.assertEqual(dialog.settings_for(identifier)["values"]["company.name"], "Manuelle Firma")
        # Values of a previously used DOCX field are retained locally, but not sent
        # while that document template is no longer selected.
        dialog.settings_for(identifier)["values"]["DOCX_ONLY"] = "Alte Eingabe"
        dialog.request_preview()
        self.wait_idle()
        self.assertTrue(dialog.preview["ready"], dialog.preview["rows"])
        self.assertEqual(dialog.preview["rows"][0]["subject"], "An Manuelle Firma")
        self.assertEqual(dialog.settings_for(identifier)["values"]["DOCX_ONLY"], "Alte Eingabe")
        dialog.reject()


    def delivery_fixture(self, *, scheduled=False):
        self.open_test_project()
        self.enable_fake_gmail()
        bulk, _, _ = self.bulk_fixture()
        bulk.request_preview()
        self.wait_idle()
        bulk.confirm.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            bulk.create()
        self.wait_idle()
        bulk.open_delivery()
        self.wait_idle()
        delivery = self.window._bulk_delivery_dialog
        self.assertIsNotNone(delivery.account.currentData())
        delivery.interval.setValue(1)
        delivery.scheduled.setChecked(scheduled)
        delivery.prepare()
        self.wait_for(lambda: delivery.dispatch and delivery.dispatch["status"] == "Ready", seconds=10)
        self.wait_idle()
        self.assertEqual(self.fake_provider.sent, [])
        return bulk, delivery

    def approve_delivery(self, delivery):
        self.assertFalse(delivery.start_button.isEnabled())
        delivery.confirm.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            delivery.approve()
        self.wait_idle()
        self.assertTrue(delivery.start_button.isEnabled())
        self.assertEqual(self.fake_provider.sent, [])

    def test_bulk_delivery_common_review_approval_and_personalized_send(self):
        bulk, delivery = self.delivery_fixture()
        self.assertEqual(delivery.jobs.count(), 2)
        self.assertIn("Hallo Ada", delivery.review.toPlainText())
        delivery.jobs.setCurrentRow(1)
        self.assertIn("Hallo Grace", delivery.review.toPlainText())
        self.assertFalse(delivery.approve_button.isEnabled())
        screenshot = os.environ.get("EMAILCENTER_BULK_DELIVERY_SCREENSHOT_PATH")
        if screenshot:
            delivery.grab().save(screenshot)
        self.approve_delivery(delivery)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            delivery.start()
        self.wait_for(lambda: delivery.dispatch["status"] == "Completed", seconds=10)
        self.assertEqual(len(self.fake_provider.sent), 2)
        self.assertEqual(delivery.dispatch["counts"]["sent"], 2)
        self.assertEqual(delivery.progress.value(), 2)
        self.assertFalse(delivery.start_button.isEnabled())
        delivery.reject()
        self.wait_idle()
        self.assertIsNone(self.window._bulk_delivery_dialog)
        self.assertEqual(self.window.snapshot["counts"]["mail_jobs"], 2)
        bulk.open_delivery()
        self.wait_idle()
        self.assertEqual(self.window._bulk_delivery_dialog.dispatch["status"], "Completed")
        self.window._bulk_delivery_dialog.reject()
        self.wait_idle()
        bulk.reject()

    def test_bulk_delivery_pause_stays_clickable_during_send_and_resume_skips_sent(self):
        from threading import Event
        bulk, delivery = self.delivery_fixture()
        self.approve_delivery(delivery)
        entered, release = Event(), Event()
        send = self.fake_provider.send
        def held(raw):
            entered.set()
            release.wait(5)
            return send(raw)
        with patch.object(self.fake_provider, "send", side_effect=held):
            try:
                with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                    delivery.start()
                self.wait_for(entered.is_set)
                self.assertTrue(delivery.busy)
                self.assertTrue(delivery.pause_button.isEnabled())
                self.assertTrue(delivery.cancel_button.isEnabled())
                delivery.reject()
                self.assertIs(self.window._bulk_delivery_dialog, delivery)
                delivery.pause_button.click()
                self.assertEqual(delivery.pending, "pause")
            finally:
                release.set()
            self.wait_for(lambda: delivery.dispatch["status"] == "Paused", seconds=10)
            self.wait_idle()
        self.assertEqual(len(self.fake_provider.sent), 1)
        self.assertTrue(delivery.start_button.isEnabled())
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            delivery.start()
        self.wait_for(lambda: delivery.dispatch["status"] == "Completed", seconds=10)
        self.assertEqual(len(self.fake_provider.sent), 2)
        delivery.reject()
        self.wait_idle()
        bulk.reject()

    def test_bulk_delivery_cancel_after_inflight_send_preserves_its_receipt(self):
        from threading import Event
        bulk, delivery = self.delivery_fixture()
        self.approve_delivery(delivery)
        entered, release = Event(), Event()
        send = self.fake_provider.send
        def held(raw):
            entered.set()
            release.wait(5)
            return send(raw)
        with patch.object(self.fake_provider, "send", side_effect=held):
            try:
                with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                    delivery.start()
                self.wait_for(entered.is_set)
                with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                    delivery.cancel_button.click()
                self.assertEqual(delivery.pending, "cancel")
            finally:
                release.set()
            self.wait_for(lambda: delivery.dispatch["status"] == "Cancelled", seconds=10)
            self.wait_idle()
        self.assertEqual(len(self.fake_provider.sent), 1)
        self.assertEqual([item["job"]["status"] for item in delivery.dispatch["items"]], ["Sent", "Cancelled"])
        self.assertFalse(delivery.start_button.isEnabled())
        delivery.reject()
        self.wait_idle()
        bulk.reject()

    def test_bulk_delivery_planned_close_pauses_and_reopen_keeps_jobs(self):
        bulk, delivery = self.delivery_fixture(scheduled=True)
        self.approve_delivery(delivery)
        self.assertEqual(delivery.start_button.text(), "Geplant starten")
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            delivery.start()
        self.wait_for(lambda: bool(delivery.dispatch.get("wait_until")))
        self.assertTrue(delivery.timer.isActive())
        self.assertEqual(self.fake_provider.sent, [])
        identifier = delivery.dispatch["id"]
        job_ids = [item["job_id"] for item in delivery.dispatch["items"]]
        delivery.reject()
        self.wait_for(lambda: self.window._bulk_delivery_dialog is None)
        self.wait_idle()
        bulk.open_delivery()
        self.wait_idle()
        delivery = self.window._bulk_delivery_dialog
        self.assertEqual(delivery.dispatch["id"], identifier)
        self.assertEqual(delivery.dispatch["status"], "Paused")
        self.assertEqual([item["job_id"] for item in delivery.dispatch["items"]], job_ids)
        self.assertFalse(delivery.auto)
        delivery.reject()
        self.wait_idle()
        bulk.reject()

    def test_bulk_delivery_uncertain_result_stops_and_exposes_reconciliation(self):
        from emailcenter.gmail import ProviderError
        bulk, delivery = self.delivery_fixture()
        self.approve_delivery(delivery)
        self.fake_provider.failure = ProviderError("Timeout beim Versand", uncertain=True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            delivery.start()
        self.wait_for(lambda: delivery.dispatch["counts"]["blocked"] == 1)
        self.wait_idle()
        self.assertEqual(delivery.dispatch["status"], "Paused")
        self.assertFalse(delivery.start_button.isEnabled())
        self.assertTrue(delivery.reconcile_button.isEnabled())
        self.assertFalse(delivery.retry_button.isEnabled())
        self.assertIn("Timeout", delivery.review.toPlainText())
        job = delivery.selected_job()
        self.fake_provider.matches = [{"id": "confirmed", "message_id": job["message_id_snapshot"]}]
        delivery.reconcile_button.click()
        self.wait_idle()
        self.assertEqual(delivery.dispatch["counts"]["sent"], 1)
        self.assertTrue(delivery.start_button.isEnabled())
        self.fake_provider.failure = None
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            delivery.start()
        self.wait_for(lambda: delivery.dispatch["status"] == "Completed", seconds=10)
        self.assertEqual(len(self.fake_provider.sent), 2)
        delivery.reject()
        self.wait_idle()
        bulk.reject()

    def test_bulk_delivery_long_ui_gap_pauses_planned_group(self):
        bulk, delivery = self.delivery_fixture(scheduled=True)
        self.approve_delivery(delivery)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            delivery.start()
        self.wait_for(lambda: bool(delivery.dispatch.get("wait_until")))
        delivery.last_tick = time.monotonic() - 100
        delivery.tick()
        self.wait_idle()
        self.assertEqual(delivery.dispatch["status"], "Paused")
        self.assertEqual(self.fake_provider.sent, [])
        delivery.reject()
        self.wait_idle()
        bulk.reject()

    def test_bulk_delivery_preparation_error_retains_completed_jobs_and_resumes(self):
        from emailcenter.mail_service import MailService
        bulk, delivery = self.delivery_fixture()
        delivery.reject()
        self.wait_idle()
        bulk.reject()
        # A second group starts with the same source contacts, and fails only on its second capture.
        self.window.workspace_mail_action("bulk", None)
        self.wait_idle()
        bulk = self.window._bulk_dialog
        bulk.template.setCurrentIndex(1)
        bulk.select_visible()
        bulk.request_preview()
        self.wait_idle()
        bulk.confirm.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            bulk.create()
        self.wait_idle()
        bulk.open_delivery()
        self.wait_idle()
        delivery = self.window._bulk_delivery_dialog
        original, calls = MailService.prepare_job, []
        def fail_second(service, *args, **kwargs):
            calls.append(args[0])
            if len(calls) == 2:
                raise ValueError("Testfehler bei zweitem Auftrag")
            return original(service, *args, **kwargs)
        with patch.object(MailService, "prepare_job", new=fail_second):
            delivery.prepare()
            self.wait_for(lambda: delivery.dispatch and bool(delivery.dispatch["last_error"]), seconds=10)
            self.wait_idle()
        first = delivery.dispatch["items"][0]["job_id"]
        self.assertEqual(delivery.dispatch["counts"]["prepared"], 1)
        self.assertTrue(delivery.prepare_button.isEnabled())
        self.assertFalse(delivery.start_button.isEnabled())
        delivery.prepare()
        self.wait_for(lambda: delivery.dispatch["status"] == "Ready", seconds=10)
        self.assertEqual(delivery.dispatch["items"][0]["job_id"], first)
        self.assertEqual(self.fake_provider.sent, [])
        delivery.reject()
        self.wait_idle()
        bulk.reject()

    def test_bulk_delivery_gap_during_inflight_send_requires_new_approval_before_next(self):
        from threading import Event
        bulk, delivery = self.delivery_fixture()
        self.approve_delivery(delivery)
        entered, release = Event(), Event()
        send = self.fake_provider.send
        def held(raw):
            entered.set()
            release.wait(5)
            return send(raw)
        with patch.object(self.fake_provider, "send", side_effect=held):
            try:
                with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
                    delivery.start()
                self.wait_for(entered.is_set)
                delivery.last_tick = time.monotonic() - 100
            finally:
                release.set()
            self.wait_for(lambda: delivery.dispatch["status"] == "Paused", seconds=10)
            self.wait_idle()
        self.assertEqual(len(self.fake_provider.sent), 1)
        self.assertFalse(delivery.start_button.isEnabled())
        self.assertTrue(delivery.overdue.isVisible())
        self.assertEqual(delivery.dispatch["items"][1]["job"]["requires_overdue_confirmation"], 1)
        delivery.confirm.setChecked(True)
        delivery.approve()
        self.assertIn("Überfällige", self.errors.pop())
        delivery.overdue.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.Yes):
            delivery.approve()
        self.wait_idle()
        self.assertTrue(delivery.start_button.isEnabled())
        delivery.reject()
        self.wait_idle()
        bulk.reject()

    def test_bulk_delivery_single_job_dialog_does_not_offer_group_bypass(self):
        from emailcenter.ui.mail_page import MailJobDialog, job_actions
        bulk, delivery = self.delivery_fixture()
        self.approve_delivery(delivery)
        job = delivery.dispatch["items"][0]["job"]
        self.assertFalse(job_actions(job)["send"])
        single = MailJobDialog(job, delivery)
        self.assertTrue(single.approve_button.isHidden())
        self.assertTrue(single.send_now_button.isHidden())
        single.deleteLater()
        delivery.reject()
        self.wait_idle()
        bulk.reject()


    def test_workflow_ui_word_replace_review_create_and_project_reset(self):
        self.assertTrue(self.window.navigation.item(self.window.page_order.index("word_replace")).isHidden())
        self.open_test_project()
        self.window.workspace_action("word_replace", None)
        self.wait_idle()
        page = self.window.pages["word_replace"]
        path = self.directory / "original.docx"
        _make_docx(path)
        page.source_path.setText(str(path))
        page.add_rule(find="{{NAME}}", replacement="Ada")
        page.request_preview()
        self.wait_idle()
        self.assertIn("Hello Ada", page.after.toPlainText())
        self.assertFalse(page.create_button.isEnabled())
        page.confirm.setChecked(True)
        self.assertTrue(page.create_button.isEnabled())
        preview_id = page.preview["preview_id"]
        self.window.toggle_theme()
        self.assertEqual(page.preview["preview_id"], preview_id)
        self.assertTrue(page.confirm.isChecked())
        self.window.toggle_theme()
        self.window.resize(820, 600)
        self.application.processEvents()
        self.assertEqual(page.footer_actions.getItemPosition(2)[0], 1)
        self.assertTrue(page.create_button.isVisible())
        self.assertLess(page.create_button.mapTo(page, page.create_button.rect().bottomRight()).y(), page.height())
        self.window.resize(1250, 820)
        self.application.processEvents()
        screenshot = os.environ.get("EMAILCENTER_WORD_REPLACE_SCREENSHOT_PATH")
        if screenshot:
            self.window.toast.dismiss()
            self.window.grab().save(screenshot)
        page.rules.item(0, 1).setText("Grace")
        self.assertIsNone(page.preview)
        self.assertFalse(page.create_button.isEnabled())
        page.request_preview()
        self.wait_idle()
        page.confirm.setChecked(True)
        with patch.object(QMessageBox, "question", return_value=QMessageBox.StandardButton.No):
            page.create()
        self.wait_for(lambda: page.run is not None and page.run["status"] == "Completed")
        self.wait_idle()
        self.assertEqual(self.errors, [])
        self.assertEqual(self.window.snapshot["counts"]["documents"], 1)
        output = Path(page.output_directory.text()) / "Dokument.docx"
        from emailcenter.documents import extract_docx_text
        self.assertIn("Hello Grace", extract_docx_text(output))
        self.submit("close_project", {})
        self.assertEqual(page.rules.rowCount(), 0)
        self.assertEqual(page.source_path.text(), "")

    def test_workflow_ui_saved_list_reused_in_word_and_new_copy_has_own_ids(self):
        from PySide6.QtWidgets import QTableWidgetItem
        self.open_test_project()
        self.window.navigate("word_replace")
        self.wait_idle()
        page = self.window.pages["word_replace"]
        self.window.open_word_replace_source("manual")
        dialog = self.window._bulk_source_dialog
        dialog.list_name.setText("Dokumentdaten")
        dialog.rows.setItem(0, dialog.fields.index("contact.display_name"), QTableWidgetItem("Ada"))
        dialog.save_list()
        self.wait_idle()
        original = dialog.saved_list
        self.assertIsNotNone(original)
        dialog.apply()
        self.wait_idle()
        self.assertEqual(page.sources[0]["kind"], "list")
        self.assertIsNone(self.window._bulk_source_dialog)
        self.window.open_word_replace_source("manual")
        dialog = self.window._bulk_source_dialog
        dialog.list_choice.setCurrentIndex(dialog.list_choice.findData(original["id"]))
        self.wait_idle()
        self.assertEqual(dialog.rows.item(0, dialog.fields.index("contact.display_name")).text(), "Ada")
        dialog.list_choice.setCurrentIndex(0)
        dialog.save_list()
        self.wait_idle()
        self.assertNotEqual(original["id"], dialog.saved_list["id"])
        self.assertNotEqual(original["rows"][0]["id"], dialog.saved_list["rows"][0]["id"])
        self.assertEqual(self.errors, [])
        dialog.reject()

    def test_workflow_ui_word_uses_excel_source_without_bulk_dialog(self):
        self.open_test_project()
        bulk, _ = self.source_bulk_fixture(excel=True)
        bulk.reject()
        self.window.navigate("word_replace")
        self.wait_idle()
        page = self.window.pages["word_replace"]
        path = self.directory / "original.docx"
        _make_docx(path)
        page.source_path.setText(str(path))
        page.add_rule(find="{{NAME}}", field="contact.display_name")
        self.window.open_word_replace_source("import")
        self.wait_idle()
        source = self.window._bulk_source_dialog
        source.select_visible()
        source.apply()
        self.wait_idle()
        self.assertEqual(len(page.sources), 2)
        self.assertIsNone(self.window._bulk_dialog)
        page.request_preview()
        self.wait_idle()
        self.assertTrue(page.preview["ready"])
        self.assertIn("Hello Person 0", page.after.toPlainText())
        page.recipient.setCurrentIndex(1)
        self.assertIn("Hello Person 1", page.after.toPlainText())

    def test_workflow_ui_word_company_excel_without_email(self):
        from openpyxl import Workbook
        self.open_test_project()
        path = self.directory / "firms.xlsx"
        book = Workbook()
        book.active.append(["Firma"])
        book.active.append(["Example GmbH"])
        book.save(path)
        book.close()
        self.action_result("prepare_import", {"source_path": str(path), "entity": "companies", "mapping": {"name": "Firma"}})
        self.window.navigate("word_replace")
        self.wait_idle()
        page = self.window.pages["word_replace"]
        original = self.directory / "original.docx"
        _make_docx(original)
        page.source_path.setText(str(original))
        page.add_rule(find="{{COMPANY}}", field="company.name")
        self.window.open_word_replace_source("import")
        self.wait_idle()
        source = self.window._bulk_source_dialog
        self.assertIsNone(source.mapping["contact.email"].currentData())
        source.select_visible()
        self.assertTrue(source.apply_button.isEnabled())
        source.apply()
        self.wait_idle()
        page.request_preview()
        self.wait_idle()
        self.assertTrue(page.preview["ready"])
        self.assertIn("at Example GmbH", page.after.toPlainText())
        self.assertEqual(self.errors, [])

    def test_workflow_ui_mapping_profile_save_and_apply_defaults(self):
        from PySide6.QtWidgets import QInputDialog
        self.open_test_project()
        bulk, _ = self.source_bulk_fixture(excel=True)
        self.window.open_bulk_source("import")
        self.wait_idle()
        source = self.window._bulk_source_dialog
        source.defaults["contact.salutation"].setText("Hallo")
        source.fallback.setChecked(False)
        with patch.object(QInputDialog, "getText", return_value=("Meine Zuordnung", True)):
            source.save_profile()
        self.wait_idle()
        self.assertEqual(len(source.profiles), 1)
        source.defaults["contact.salutation"].clear()
        source.mapping["company.name"].setCurrentIndex(0)
        source.fallback.setChecked(True)
        source.apply_profile()
        self.assertEqual(source.defaults["contact.salutation"].text(), "Hallo")
        self.assertEqual(source.mapping["company.name"].currentData(), "Firma")
        self.assertFalse(source.fallback.isChecked())
        self.assertEqual(len(bulk.data["source_profiles"]), 1)
        source.reject()
        bulk.reject()

    def test_workflow_ui_export_selection_columns_and_settings_defaults(self):
        from PySide6.QtCore import Qt
        from PySide6.QtWidgets import QFileDialog
        import csv
        self.open_test_project()
        ada = self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": "Ada", "email": "ada@example.test"}})
        self.action_result("create_record", {"entity": "contacts", "fields": {"display_name": "Grace", "email": "grace@example.test"}})
        self.window.open_export(entity="contacts", selected_ids=[ada["id"]])
        self.wait_idle()
        dialog = self.window._export_dialog
        self.assertEqual(dialog.preview["total"], 2)
        dialog.scope.setCurrentIndex(1)
        self.assertIsNone(dialog.preview)
        dialog.request_preview()
        self.wait_idle()
        self.assertEqual(dialog.preview["total"], 1)
        for i in range(dialog.columns.count()):
            if dialog.columns.item(i).data(Qt.ItemDataRole.UserRole) not in ("display_name", "email"):
                dialog.columns.item(i).setCheckState(Qt.CheckState.Unchecked)
        screenshot = os.environ.get("EMAILCENTER_EXPORT_SCREENSHOT_PATH")
        if screenshot:
            dialog.grab().save(screenshot)
        path = self.directory / "selection.csv"
        with patch.object(QFileDialog, "getSaveFileName", return_value=(str(path), "CSV")):
            dialog.export()
        self.wait_idle()
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream, delimiter=";"))
        self.assertEqual(rows, [["Name", "E-Mail"], ["Ada", "ada@example.test"]])
        self.assertIn("1 Zeilen", dialog.status.text())
        dialog.reject()

    def test_workflow_ui_settings_persist_and_toolbar_retains_export_edits(self):
        self.window.navigate("settings")
        settings = self.window.pages["settings"]
        settings.controls["export_format"].setCurrentIndex(1)
        settings.controls["word_replace_missing"].setCurrentIndex(1)
        settings.controls["conversion_timeout"].setValue(45)
        self.window.toggle_theme()
        self.assertEqual(settings.controls["export_format"].currentData(), "xlsx")
        self.assertEqual(settings.controls["word_replace_missing"].currentData(), "keep")
        self.assertEqual(settings.controls["conversion_timeout"].value(), 45)
        settings.save_button.click()
        self.assertEqual(self.errors, [])
        loaded, warning = PreferencesStore(self.directory / "preferences.json").load()
        self.assertIsNone(warning)
        self.assertEqual(loaded.export_format, "xlsx")
        self.assertEqual(loaded.conversion_timeout, 45)
        self.open_test_project()
        self.window.open_word_replace_settings()
        self.assertEqual(settings.tabs.currentIndex(), 2)
        self.window.open_export()
        self.wait_idle()
        self.assertEqual(self.window._export_dialog.format.currentData(), "xlsx")
        self.window._export_dialog.reject()

    def test_workflow_ui_group_report_export_never_sends(self):
        bulk, delivery = self.delivery_fixture()
        delivery.export_button.click()
        self.wait_idle()
        dialog = self.window._export_dialog
        self.assertEqual(dialog.preview["entity"], "email_jobs")
        self.assertEqual(dialog.preview["batch_id"], delivery.data["batch"]["id"])
        self.assertEqual(dialog.preview["total"], 2)
        self.assertEqual(self.fake_provider.sent, [])
        dialog.reject()
        delivery.reject()
        self.wait_idle()
        bulk.reject()

    def test_workflow_ui_word_cancel_is_reachable_during_worker(self):
        from threading import Event
        from emailcenter.word_replace import WordReplaceService
        self.open_test_project()
        self.window.navigate("word_replace")
        self.wait_idle()
        page = self.window.pages["word_replace"]
        path = self.directory / "original.docx"
        _make_docx(path)
        page.source_path.setText(str(path))
        page.add_rule(find="{{NAME}}", replacement="Ada")
        entered, release = Event(), Event()
        original = WordReplaceService.preview
        def held(service, *args, **kwargs):
            entered.set()
            release.wait(5)
            return original(service, *args, **kwargs)
        with patch.object(WordReplaceService, "preview", held):
            try:
                page.request_preview()
                self.wait_for(entered.is_set)
                self.assertTrue(page.cancel_button.isEnabled())
                self.assertFalse(page.rules.isEnabled())
                page.cancel_button.click()
            finally:
                release.set()
            self.wait_idle()
        self.assertEqual(len(self.errors), 1)
        self.assertIn("abgebrochen", self.errors[0])
        self.assertIsNone(page.preview)


    def test_hotkeys_settings_apply_disable_reset_and_retain_unsaved_edits(self):
        self.window.navigate("settings")
        settings = self.window.pages["settings"]
        settings.shortcut_controls["refresh"].setText("strg+r")
        settings.shortcut_controls["export"].clear()
        self.window.toggle_theme()
        self.assertEqual(settings.shortcut_controls["refresh"].text(), "strg+r")
        settings.save_button.click()
        self.assertEqual(self.errors, [])
        self.assertEqual(self.window.actions["refresh"].shortcut().toString(), "Ctrl+R")
        self.assertTrue(self.window.actions["export"].shortcut().isEmpty())
        self.assertEqual(PreferencesStore(self.directory / "preferences.json").load()[0].shortcuts["export"], "")
        settings.reset_shortcuts()
        self.assertEqual(settings.shortcut_controls["refresh"].text(), "F5")
        self.assertEqual(self.window.actions["refresh"].shortcut().toString(), "Ctrl+R")
        settings.discard()
        self.assertEqual(settings.shortcut_controls["refresh"].text(), "Ctrl+R")

    def test_hotkeys_conflict_keeps_active_bindings_and_form(self):
        self.window.navigate("settings")
        settings = self.window.pages["settings"]
        settings.shortcut_controls["new"].setText("Ctrl+O")
        self.assertIn("doppelt", settings.state.text())
        settings.save_button.click()
        self.assertEqual(len(self.errors), 1)
        self.assertIn("doppelt", self.errors[0])
        self.assertEqual(self.window.actions["new"].shortcut().toString(), "Ctrl+N")
        self.assertEqual(settings.shortcut_controls["new"].text(), "Ctrl+O")
        settings.reset_shortcuts()
        settings.save_button.click()
        self.assertEqual(self.window.preferences.shortcuts["new"], "Ctrl+N")

    def test_hotkeys_actual_navigation_and_search_key_events(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        self.open_test_project()
        self.window.activateWindow()
        self.application.processEvents()
        QTest.keyClick(self.window, Qt.Key.Key_2, Qt.KeyboardModifier.ControlModifier)
        self.wait_idle()
        self.assertIs(self.window.current_page(), self.window.pages["word_replace"])
        QTest.keyClick(self.window, Qt.Key.Key_1, Qt.KeyboardModifier.ControlModifier)
        self.wait_idle()
        self.assertIs(self.window.current_page(), self.window.pages["workspace"])
        self.window.navigate("contacts")
        self.wait_idle()
        page = self.window.pages["contacts"]
        page.search.setText("Ada")
        page.table.setFocus()
        QTest.keyClick(self.window, Qt.Key.Key_F, Qt.KeyboardModifier.ControlModifier)
        self.application.processEvents()
        self.assertTrue(page.search.hasFocus())
        self.assertEqual(page.search.selectedText(), "Ada")

    def test_hotkeys_actual_record_save_key_event(self):
        from PySide6.QtCore import Qt
        from PySide6.QtTest import QTest
        self.open_test_project()
        self.window.create_record("companies")
        self.wait_idle()
        dialog = self.window._record_dialog
        dialog.fields["name"].setText("Keyboard Company")
        dialog.activateWindow()
        self.application.processEvents()
        QTest.keyClick(dialog, Qt.Key.Key_S, Qt.KeyboardModifier.ControlModifier)
        self.wait_idle()
        self.assertEqual(self.errors, [])
        self.assertIsNone(self.window._record_dialog)
        self.assertEqual(self.window.snapshot["counts"]["companies"], 1)

    def test_hotkeys_help_and_save_do_not_approve_or_send_jobs(self):
        bulk, delivery = self.delivery_fixture()
        job = delivery.dispatch["items"][0]["job"]
        self.window.review_mail(job)
        dialog = self.window._mail_dialog
        self.application.processEvents()
        self.window.trigger_hotkey("save")
        self.assertEqual(self.fake_provider.sent, [])
        self.assertIs(self.window._mail_dialog, dialog)
        with patch.object(QMessageBox, "information") as info:
            self.window.trigger_hotkey("help")
        self.assertFalse(info.called)
        self.assertIsNotNone(self.window._guide_dialog)
        self.assertIn('EmailCenter', self.window._guide_dialog.browser.toPlainText())
        self.assertEqual(self.fake_provider.sent, [])
        self.window._guide_dialog.close()
        self.assertEqual(delivery.dispatch["items"][0]["job"]["approval_status"], "pending")
        dialog.reject()
        delivery.reject()
        self.wait_idle()
        bulk.reject()


if __name__ == "__main__":
    unittest.main()
