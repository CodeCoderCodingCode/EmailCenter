"""Bounded CSV/XLSX readers and conservative column suggestions."""

from __future__ import annotations

from .language import translate as _ui_tr, format_text as _ui_format

import csv
from dataclasses import dataclass
from datetime import date, datetime
import hashlib
import math
from pathlib import Path
import unicodedata
import zipfile


class ImportSourceError(ValueError):
    """A source cannot be read without losing or ambiguously mapping data."""


MAX_ROWS = 100_000
MAX_COLUMNS = 200
MAX_CELLS = 1_000_000
MAX_FILE_BYTES = 50 * 1024 * 1024

FIELDS = {
    "data_rows": [],
    "companies": [("name", "Unternehmen", True), ("website", "Website", False),
                  ("industry", "Branche", False), ("notes", "Notizen", False)],
    "contacts": [("display_name", "Name", True), ("first_name", "Vorname", False),
                 ("last_name", "Nachname", False), ("email", "E-Mail", False),
                 ("phone", "Telefon", False), ("position", "Position", False), ("notes", "Notizen", False)],
    "applications": [("title", "Bewerbung / Stelle", True), ("company_name", "Unternehmen (vorhanden)", False),
                     ("applicant_email", "Bewerber-E-Mail (vorhandenes Profil)", False),
                     ("status_code", "Statuscode", False), ("application_date", "Bewerbungsdatum", False),
                     ("source", "Quelle", False), ("notes", "Notizen", False)],
}

ALIASES = {
    "name": {"name", "unternehmen", "company", "firma", "firmenname"},
    "display_name": {"name", "display_name", "kontakt", "contact", "voller name", "ansprechpartner", "ansprechpartner name"},
    "first_name": {"vorname", "first_name", "firstname"}, "last_name": {"nachname", "last_name", "lastname"},
    "email": {"email", "e-mail", "e_mail", "mail", "emailadresse"},
    "phone": {"phone", "telefon", "telephone", "telefonnummer"},
    "website": {"website", "webseite", "url"}, "industry": {"industry", "branche"},
    "notes": {"notes", "notizen", "bemerkung"}, "position": {"position", "funktion"},
    "title": {"title", "titel", "stelle", "stellenbezeichnung", "bewerbung"},
    "company_name": {"company", "unternehmen", "firma", "company_name"},
    "applicant_email": {"applicant_email", "bewerber-email", "bewerber-e-mail"},
    "status_code": {"status", "statuscode", "status_code"},
    "application_date": {"datum", "date", "bewerbungsdatum", "application_date"},
    "source": {"source", "quelle"},
}


@dataclass(frozen=True)
class SourceRow:
    number: int
    values: dict[str, str]
    errors: tuple[str, ...] = ()


@dataclass(frozen=True)
class SourceTable:
    headers: tuple[str, ...]
    rows: tuple[SourceRow, ...]
    options: dict[str, str]
    sheets: tuple[str, ...] = ()


def source_hash(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_source(path: str | Path, options: dict | None = None) -> SourceTable:
    source = Path(path)
    options = options or {}
    if not source.is_file() or source.suffix.lower() not in (".csv", ".xlsx"):
        raise ImportSourceError(_ui_tr('Bitte eine vorhandene CSV- oder XLSX-Datei auswählen.'))
    if source.stat().st_size > MAX_FILE_BYTES:
        raise ImportSourceError(_ui_tr('Die Importdatei überschreitet 50 MiB.'))
    try:
        return _read_csv(source, options) if source.suffix.lower() == ".csv" else _read_xlsx(source, options)
    except ImportSourceError:
        raise
    except Exception as error:
        raise ImportSourceError(_ui_format('Importdatei konnte nicht gelesen werden: {0}', f'{error}')) from error


def _headers(values) -> tuple[str, ...]:
    result = tuple(str(value or "").strip() for value in values)
    if not result or not any(result):
        raise ImportSourceError(_ui_tr('Die erste Zeile muss Spaltenüberschriften enthalten.'))
    if len(result) > MAX_COLUMNS:
        raise ImportSourceError(_ui_format('Maximal {0} Spalten werden unterstützt.', f'{MAX_COLUMNS}'))
    if any(not value for value in result) or len({value.casefold() for value in result}) != len(result):
        raise ImportSourceError(_ui_tr('Spaltenüberschriften müssen eindeutig und vollständig sein.'))
    return result


def _append_row(rows, number, headers, values, errors=()):
    if len(rows) >= MAX_ROWS or (len(rows) + 1) * len(headers) > MAX_CELLS:
        raise ImportSourceError(_ui_tr('Importlimit erreicht: 100.000 Zeilen bzw. 1.000.000 Zellen.'))
    errors = list(errors)
    if len(values) > len(headers) and any(str(value or "").strip() for value in values[len(headers):]):
        errors.append("Zusätzliche Zellen ohne Spaltenüberschrift.")
    normalized = [str(value or "").replace("\r\n", "\n").replace("\r", "\n") for value in values[:len(headers)]]
    if any(len(value) > 32768 for value in normalized):
        errors.append("Eine Zelle enthält mehr als 32.768 Zeichen.")
    data = {header: normalized[index] if index < len(normalized) else "" for index, header in enumerate(headers)}
    rows.append(SourceRow(number, data, tuple(errors)))


def _read_csv(path, options):
    requested = options.get("encoding", "auto")
    if requested not in ("auto", "utf-8-sig", "cp1252"):
        raise ImportSourceError(_ui_tr('CSV-Zeichensatz muss UTF-8 oder Windows-1252 sein.'))
    raw = path.read_bytes()
    if requested == "auto":
        try:
            text, encoding = raw.decode("utf-8-sig"), "utf-8-sig"
        except UnicodeDecodeError:
            text, encoding = raw.decode("cp1252"), "cp1252"
    else:
        text, encoding = raw.decode(requested), requested
    delimiter = options.get("delimiter", "auto")
    if delimiter == "auto":
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ";" if ";" in text.partition("\n")[0] else ","
    if delimiter not in (",", ";", "\t", "|"):
        raise ImportSourceError(_ui_tr('Ungültiges CSV-Trennzeichen.'))
    import io
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    headers = _headers(next(reader, []))
    rows = []
    for number, values in enumerate(reader, 2):
        _append_row(rows, number, headers, values)
    return SourceTable(headers, tuple(rows), {"encoding": encoding, "delimiter": delimiter})


def _read_xlsx(path, options):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise ImportSourceError(_ui_tr('Excel-Unterstützung fehlt. Bitte requirements.txt installieren.')) from error
    with zipfile.ZipFile(path) as archive:
        if sum(entry.file_size for entry in archive.infolist()) > 200 * 1024 * 1024:
            raise ImportSourceError(_ui_tr('Die entpackte Excel-Datei überschreitet 200 MiB.'))
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheets = tuple(workbook.sheetnames)
        selected = options.get("sheet") or sheets[0]
        if selected not in sheets:
            raise ImportSourceError(_ui_format('Tabellenblatt fehlt: {0}', f'{selected}'))
        sheet = workbook[selected]
        sheet.reset_dimensions()
        iterator = sheet.iter_rows()
        first = list(next(iterator, ()))
        while first and first[-1].value is None:
            first.pop()
        if any(cell.data_type == "f" for cell in first):
            raise ImportSourceError(_ui_tr('Spaltenüberschriften dürfen keine Formeln enthalten.'))
        headers = _headers([cell.value for cell in first])
        rows = []
        for number, cells in enumerate(iterator, 2):
            values, errors = [], []
            for cell in cells:
                value = cell.value
                if cell.data_type in ("f", "e"):
                    errors.append("Formel oder Excel-Fehler: Bitte einen festen Zellwert verwenden.")
                if isinstance(value, (datetime, date)):
                    value = value.isoformat()
                elif isinstance(value, float) and not math.isfinite(value):
                    errors.append("Nicht endlicher Zahlenwert.")
                values.append("" if value is None else str(value))
            _append_row(rows, number, headers, values, errors)
        return SourceTable(headers, tuple(rows), {"sheet": selected}, sheets)
    finally:
        workbook.close()


def _column_key(value):
    return "".join(character for character in unicodedata.normalize("NFKC", value).casefold() if character.isalnum())


def suggest_mapping(headers, entity):
    """Suggest only a unique exact alias; the user still approves the mapping."""
    proposals, used = {}, set()
    for field, _, _ in FIELDS[entity]:
        aliases = {_column_key(alias) for alias in ALIASES.get(field, {field})}
        matches = [header for header in headers if _column_key(header) in aliases]
        if len(matches) == 1 and matches[0] not in used:
            proposals[field] = matches[0]
            used.add(matches[0])
    return proposals
