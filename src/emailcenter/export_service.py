"""Whitelisted data exports with reviewed rows, formula-safe cells and atomic files."""

from .language import translate as _ui_tr, format_text as _ui_format

import csv
from datetime import datetime, timezone
import hashlib
import io
import os
from pathlib import Path
import tempfile

from .mail_bulk import digest
from .mail_models import MailError


# No token stores, account configuration, credential IDs or OAuth data are projected.
EXPORTS = {
    "companies": ("Unternehmen", "companies c", "c.deleted_at IS NULL", "c.name", [
        ("name", "Unternehmen", "c.name", "text"), ("industry", "Branche", "c.industry", "text"),
        ("website", "Website", "c.website", "text"), ("notes", "Notizen", "c.notes", "text")]),
    "contacts": ("Kontakte", "contacts c", "c.deleted_at IS NULL", "c.display_name", [
        ("display_name", "Name", "c.display_name", "text"), ("email", "E-Mail", "c.email", "text"),
        ("phone", "Telefon", "c.phone", "text"), ("position", "Position", "c.position", "text"), ("notes", "Notizen", "c.notes", "text")]),
    "applicants": ("Bewerberprofile", "applicant_profiles c", "c.deleted_at IS NULL", "c.display_name", [
        ("display_name", "Name", "c.display_name", "text"), ("email", "E-Mail", "c.email", "text"), ("phone", "Telefon", "c.phone", "text")]),
    "jobs": ("Stellen", "job_postings c JOIN companies b ON b.id=c.company_id", "c.deleted_at IS NULL", "c.title", [
        ("title", "Stelle", "c.title", "text"), ("company_name", "Unternehmen", "b.name", "text"),
        ("location", "Ort", "c.location", "text"), ("reference_number", "Referenz", "c.reference_number", "text"),
        ("deadline", "Frist", "c.deadline", "date"), ("source_url", "Quelle", "c.source_url", "text")]),
    "applications": ("Bewerbungen", "applications c LEFT JOIN companies b ON b.id=c.company_id JOIN application_statuses s ON s.id=c.status_id", "c.deleted_at IS NULL", "c.title", [
        ("title", "Bewerbung", "c.title", "text"), ("company_name", "Unternehmen", "b.name", "text"),
        ("status_label", "Status", "s.label", "text"), ("application_date", "Bewerbungsdatum", "c.application_date", "date"),
        ("next_follow_up_at", "Nachfassen (UTC bei Zeitzonenangabe)", "c.next_follow_up_at", "date"), ("notes", "Notizen", "c.notes", "text")]),
    "email_jobs": ("Versandbericht", "email_jobs c LEFT JOIN email_draft_batches b ON b.id=c.batch_id", "1=1", "c.subject_snapshot", [
        ("subject", "Betreff", "c.subject_snapshot", "text"), ("recipients", "Empfänger (JSON)", "c.recipients_snapshot_json", "text"),
        ("sender", "Absender", "c.sender_snapshot", "text"), ("status", "Status", "c.status", "text"),
        ("approval", "Freigabe", "c.approval_status", "text"), ("group_name", "Gruppe", "b.name", "text"),
        ("scheduled", "Geplant (UTC)", "c.scheduled_at_utc", "date"), ("sent", "Gesendet (UTC)", "c.sent_at_utc", "date"),
        ("attempts", "Versuche", "c.attempt_count", "integer"), ("error", "Letzter Fehler", "c.last_error_message", "text")]),
}


def safe_text(value):
    text = "" if value is None else str(value)
    return "'" + text if text.lstrip(" \t\r\n\v\f").startswith(("=", "+", "-", "@")) else text


def excel_date(value):
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        return parsed.astimezone(timezone.utc).replace(tzinfo=None) if parsed.tzinfo else parsed
    except (ValueError, TypeError, AttributeError):
        return str(value)


class ExportService:
    def __init__(self, project):
        self.project, self.db = project, project.connection

    def preview(self, entity="contacts", search="", selected_ids=None, batch_id=None):
        if entity not in EXPORTS or not isinstance(search, str) or len(search) > 1000:
            raise MailError(_ui_tr('Ungültige Exportauswahl.'))
        title, tables, condition, ordering, columns = EXPORTS[entity]
        parameters = []
        if search.strip():
            escaped = search.strip().replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
            searchable = {"companies": ("c.name", "c.industry", "c.website"), "contacts": ("c.display_name", "c.email", "c.position"),
                          "applications": ("c.title", "b.name", "s.label"), "jobs": ("c.title", "b.name", "c.location", "c.reference_number"),
                          "applicants": ("c.display_name", "c.email", "c.phone"), "email_jobs": ("c.subject_snapshot", "c.recipients_snapshot_json", "c.status")}[entity]
            condition += " AND (" + " OR ".join(f"{expression} LIKE ? ESCAPE '\\'" for expression in searchable) + ")"
            parameters.extend([f"%{escaped}%"] * len(searchable))
        if selected_ids is not None:
            if not isinstance(selected_ids, list) or not 1 <= len(selected_ids) <= 1000 or any(not isinstance(key, str) for key in selected_ids) or len(set(selected_ids)) != len(selected_ids):
                raise MailError(_ui_tr('1 bis 1000 verschiedene Datensätze für den Export auswählen.'))
            condition += " AND c.id IN (" + ",".join("?" for _ in selected_ids) + ")"
            parameters.extend(selected_ids)
        if batch_id is not None:
            if entity != "email_jobs" or not self.db.execute("SELECT 1 FROM email_draft_batches WHERE id=?", (batch_id,)).fetchone():
                raise MailError(_ui_tr('Ungültige Versandgruppe.'))
            condition += " AND c.batch_id=?"
            parameters.append(batch_id)
        projection = "c.id," + ",".join(f"{expression} AS {key}" for key, _, expression, _ in columns)
        rows = [dict(row) for row in self.db.execute(f"SELECT {projection} FROM {tables} WHERE {condition} ORDER BY {ordering} COLLATE NOCASE,c.id LIMIT 10001", parameters)]
        if len(rows) > 10000:
            raise MailError(_ui_tr('Export auf höchstens 10000 Zeilen begrenzen; Filter oder Auswahl verwenden.'))
        value = {"entity": entity, "title": title, "columns": [{"key": key, "caption": caption, "type": kind} for key, caption, _, kind in columns],
                 "rows": rows, "total": len(rows), "search": search, "selected_ids": selected_ids, "batch_id": batch_id}
        return {**value, "fingerprint": digest(value)}

    def write(self, *, preview, columns, target_path, file_format="csv", delimiter=";", encoding="utf-8-sig", overwrite=False, cancelled=lambda: False):
        available = {column["key"]: column for column in preview["columns"]}
        if not isinstance(columns, list) or not columns or any(not isinstance(key, str) or key not in available for key in columns) or len(set(columns)) != len(columns):
            raise MailError(_ui_tr('Mindestens eine gültige Exportspalte auswählen.'))
        if file_format not in ("csv", "xlsx") or delimiter not in (";", ",", "\t") or encoding not in ("utf-8-sig", "utf-8", "cp1252", "utf-16") or type(overwrite) is not bool:
            raise MailError(_ui_tr('Ungültiges Exportformat.'))
        target = Path(target_path).expanduser()
        if not target.is_absolute():
            raise MailError(_ui_tr('Einen absoluten Exportpfad auswählen.'))
        target = target.resolve()
        if target.suffix.lower() != "." + file_format:
            raise MailError(_ui_tr('Dateiendung muss zum Exportformat passen.'))
        if target.is_relative_to(self.project.path.resolve()) and not target.is_relative_to((self.project.path / "exports").resolve()):
            raise MailError(_ui_tr('Innerhalb des Projekts nur unter exports speichern; Projektdaten bleiben geschützt.'))
        existed = target.exists()
        if existed and (not overwrite or not target.is_file()):
            raise MailError(_ui_tr('Die Zieldatei existiert bereits. Überschreiben ausdrücklich bestätigen.'))
        def target_hash():
            result = hashlib.sha256()
            with target.open("rb") as existing:
                for chunk in iter(lambda: existing.read(1024 * 1024), b""):
                    result.update(chunk)
            return result.hexdigest()
        original_hash = target_hash() if existed else None
        target.parent.mkdir(parents=True, exist_ok=True)
        temporary = None
        try:
            with tempfile.NamedTemporaryFile(dir=target.parent, prefix=".emailcenter-export-", suffix=".tmp", delete=False) as stream:
                temporary = Path(stream.name)
                if file_format == "csv":
                    wrapper = io.TextIOWrapper(stream, encoding=encoding, newline="", write_through=True)
                    writer = csv.writer(wrapper, delimiter=delimiter)
                    writer.writerow([available[key]["caption"] for key in columns])
                    for row in preview["rows"]:
                        if cancelled():
                            raise MailError(_ui_tr('Export abgebrochen. Die Zieldatei wurde nicht verändert.'))
                        writer.writerow([safe_text(row[key]) for key in columns])
                    wrapper.flush()
                    wrapper.detach()
                    stream.flush()
                    os.fsync(stream.fileno())
            if file_format == "xlsx":
                from openpyxl import Workbook
                from openpyxl.styles import Font, PatternFill
                book = Workbook()
                sheet = book.active
                sheet.title = preview["title"][:31]
                sheet.append([available[key]["caption"] for key in columns])
                for cell in sheet[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="2862C4")
                for row in preview["rows"]:
                    if cancelled():
                        raise MailError(_ui_tr('Export abgebrochen. Die Zieldatei wurde nicht verändert.'))
                    values = [excel_date(row[key]) if available[key]["type"] == "date" else row[key] for key in columns]
                    if any(isinstance(value, str) and (len(value) > 32767 or any(ord(c) < 32 and c not in "\t\n\r" for c in value)) for value in values):
                        raise MailError(_ui_tr('Excel-Zelle ist zu lang oder enthält ungültige Zeichen. Spalte abwählen oder CSV nutzen.'))
                    sheet.append(values)
                    for key, cell in zip(columns, sheet[sheet.max_row]):
                        if isinstance(cell.value, str):
                            if len(cell.value) > 32767 or any(ord(c) < 32 and c not in "\t\n\r" for c in cell.value):
                                raise MailError(_ui_tr('Excel-Zelle ist zu lang oder enthält ungültige Zeichen. Spalte abwählen oder CSV nutzen.'))
                            cell.data_type = "s"  # Also disables formulas beginning with '='.
                        elif isinstance(cell.value, datetime):
                            cell.number_format = "yyyy-mm-dd hh:mm"
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for column in sheet.columns:
                    sheet.column_dimensions[column[0].column_letter].width = min(55, max(14, max(len(str(cell.value or "")) for cell in column[:50]) + 2))
                book.save(temporary)
                book.close()
            if cancelled():
                raise MailError(_ui_tr('Export abgebrochen. Die Zieldatei wurde nicht verändert.'))
            if target.exists() != existed or (existed and target_hash() != original_hash):
                raise MailError(_ui_tr('Zieldatei wurde zwischenzeitlich geändert. Export erneut starten.'))
            os.replace(temporary, target)
            return {"path": str(target), "count": preview["total"], "format": file_format, "columns": columns}
        finally:
            if temporary:
                temporary.unlink(missing_ok=True)
